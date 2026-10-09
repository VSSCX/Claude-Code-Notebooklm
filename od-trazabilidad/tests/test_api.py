import os
import tempfile
from pathlib import Path

import pytest

_tmp = tempfile.mkdtemp()
os.environ["DATABASE_URL"] = f"sqlite:///{_tmp}/test.db"
os.environ["VISORES_DIR"] = f"{_tmp}/visores"       # los visores de prueba no pisan los reales

from fastapi.testclient import TestClient  # noqa: E402

from app.db import engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Base  # noqa: E402

# Paquete real armado desde Vicente_S_Proy_Automatización.xlsm
LINEAS = [
    {"sku": "900081624", "desc": "MDWMT16W", "pendiente": 11, "enEntrega": 0},
    {"sku": "900276673", "desc": "COCINA FE5SXC", "pendiente": 40, "enEntrega": 0},
    {"sku": "900276672", "desc": "COCINA FE4SXC", "pendiente": 100, "enEntrega": 0},
    {"sku": "240086651", "desc": "EXPERIENCE CARE 14 SZ", "pendiente": 22, "enEntrega": 8},
    {"sku": "900276668", "desc": "COCINA FM4SSC", "pendiente": 40, "enEntrega": 0},
    {"sku": "900276666", "desc": "COCINA FM4TSC", "pendiente": 50, "enEntrega": 0},
]
ENTREGA = {"entrega": 8705699485, "pedido": 4005171502, "camion": 1, "vehiculo": "Rampla 53",
           "grupo": "1387800", "cita": "", "fecha": "", "hora": "", "carga": "MIX",
           "lineas": [{"sku": "900081624", "qty": 11}, {"sku": "900276672", "qty": 10},
                      {"sku": "240086651", "qty": 22}, {"sku": "900276668", "qty": 40},
                      {"sku": "900276666", "qty": 50}]}


def paquete(evento, pedidos=True, entregas=True, **extra):
    d = {"tipo": "od-traz", "v": 1, "evento": evento, "cliente": "PARIS", "un": "SDA",
         "modalidad": "Predistribuido",
         "pedidos": [{"pedido": "4005171502", "oc": "527512", "fechaOC": "2026-09-15",
                      "lineas": LINEAS}] if pedidos else [],
         "entregas": [dict(ENTREGA)] if entregas else [], "sapOk": []}
    d.update(extra)
    return d


@pytest.fixture()
def c():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    return TestClient(app)


def _estado(c):
    return c.get("/api/estado").json()


def test_flujo_script_completo(c):
    assert c.post("/api/paquetes", json=paquete("pedido", entregas=False)).status_code == 200
    ped = _estado(c)["pedidos"][0]
    assert sum(l["qty"] for l in ped["lineas"]) == 271
    assert {l["sku"]: l["externa"] for l in ped["lineas"]}["240086651"] == 8
    assert ped["oc"] == "527512" and ped["fechaOC"] == "2026-09-15"

    assert c.post("/api/paquetes", json=paquete("entregas")).json()["entregas"] == 1
    c.post("/api/paquetes", json=paquete("grupos"))
    ent = dict(ENTREGA, cita="111409", fecha="2026-09-21", hora="09:00")
    r = c.post("/api/paquetes", json=paquete("fecha_sap", pedidos=False, entregas=False,
                                             sapOk=["1387800"]) | {"entregas": [ent]})
    assert r.status_code == 200

    e = _estado(c)["entregas"][0]
    assert e["cita"] == {"numero": "111409", "fecha": "2026-09-21", "hora": "09:00"}
    assert all(e["pasos"][k]["ok"] for k in ("solicitada", "confirmada", "sap"))
    assert "etq" not in e["pasos"]
    assert sum(l["qty"] for l in e["lineas"]) == 133
    assert e["region"] == "RM" and e["carga"] == "MIX"
    assert e["un"] == "SDA" and e["tipo"] == "Predistribuido"
    assert len([x for x in e["log"] if x["origen"] == "script"]) == 3

    # Reenviar el pedido no baja cantidades ni duplica lo externo
    c.post("/api/paquetes", json=paquete("pedido", entregas=False))
    ped = _estado(c)["pedidos"][0]
    assert sum(l["qty"] for l in ped["lineas"]) == 271


def test_crud_entrega_y_reglas(c):
    r = c.put("/api/pedidos/4001", json={"pedido": "4001", "cliente": "paris", "canal": "RETAIL",
                                         "lineas": [{"sku": "0001", "desc": "X", "qty": 5}]})
    assert r.json()["cliente"] == "PARIS" and r.json()["lineas"][0]["sku"] == "1"

    base = {"entrega": "87", "pedido": "4001", "lineas": [{"sku": "1", "qty": 5}]}
    # confirmar sin fecha/hora -> 422
    r = c.put("/api/entregas/87", json=base | {"pasos": {"confirmada": {"ok": True}}})
    assert r.status_code == 422
    # paso inexistente -> 422
    assert c.put("/api/entregas/87", json=base | {"pasos": {"inventado": {"ok": True}}}).status_code == 422

    log = [{"at": "2026-09-17T12:00:00Z", "txt": "Entrega creada"}]
    r = c.put("/api/entregas/87", json=base | {"cita": {"fecha": "2026-09-20", "hora": "08:00"},
                                               "pasos": {"confirmada": {"ok": True}}, "log": log})
    assert r.status_code == 200 and r.json()["pasos"]["confirmada"]["ok"]
    # mismo log reenviado no se duplica
    r = c.put("/api/entregas/87", json=base | {"log": log + [{"at": "2026-09-17T12:05:00Z", "txt": "Otra"}]})
    # el historial trae la línea automática de creación más las que manda la web
    assert [x["txt"] for x in r.json()["log"]][0] == "Entrega creada"
    assert len(r.json()["log"]) == 3

    v1 = c.get("/api/version").json()["version"]
    assert c.delete("/api/pedidos/4001").status_code == 204
    assert c.get("/api/version").json()["version"] != v1
    assert _estado(c)["entregas"] == []   # borrado en cascada


def test_entrega_sin_pedido(c):
    r = c.put("/api/entregas/9", json={"entrega": "9", "pedido": "nope", "lineas": [{"sku": "1", "qty": 1}]})
    assert r.status_code == 422


def test_ya_no_hay_acciones_de_excel(c):
    """La plataforma no ejecuta macros: no hay catálogo ni libro, solo el seguimiento de los trabajos nativos."""
    assert c.get("/api/acciones").status_code == 404
    assert c.post("/api/acciones/cubicar").status_code in (404, 405)
    assert c.get("/api/acciones/trabajos/no-existe").status_code == 404
    from app.config import settings
    assert not hasattr(settings, "macros_workbook") and not hasattr(settings, "acciones")


def test_archivos(c):
    c.put("/api/pedidos/4001", json={"pedido": "4001", "cliente": "PARIS",
                                     "lineas": [{"sku": "1", "qty": 1}]})
    r = c.post("/api/archivos", data={"pedido": "4001", "grupo": "777", "tipo": "pdf"},
               files={"file": ("camion.pdf", b"%PDF-1.4 test", "application/pdf")})
    assert r.status_code == 201 and r.json()["grupo"] == "777"
    aid = r.json()["id"]
    assert c.get(f"/archivos/{aid}").content.startswith(b"%PDF")
    assert [a["nombre"] for a in _estado(c)["archivos"]] == ["camion.pdf"]
    # formato no permitido
    assert c.post("/api/archivos", data={"pedido": "4001"},
                  files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 422
    assert c.delete(f"/api/archivos/{aid}").status_code == 204
    assert _estado(c)["archivos"] == []


def test_plan_sop(c, monkeypatch):
    from app.integrations import bases
    c.put("/api/pedidos/4001", json={"pedido": "4001", "cliente": "PARIS",
          "lineas": [{"sku": "0900081624", "qty": 50}, {"sku": "111", "qty": 5}]})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {
        "plan": 100.0, "vendido": 30.0, "pdte_mes": 20.0, "tipo": "CONSENSO"}} if g == "PARIS" else {})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {"900081624": {"cantidad": 40, "fecha": "2026-09-20"}})
    r = c.get("/api/plan/4001").json()
    assert r["ok"] and r["grupo"] == "PARIS" and r["tipo_plan"] == "CONSENSO"
    assert r["productos"]["900081624"]["plan"]["vendido"] == 30
    assert r["productos"]["900081624"]["disponible"]["cantidad"] == 40
    assert r["productos"]["111"] == {"plan": None, "disponible": None}

    def falla(g):
        raise RuntimeError("Login failed")
    monkeypatch.setattr(bases, "plan_sop", falla)
    r = c.get("/api/plan/4001").json()
    assert r["ok"] is False and "Login failed" in r["error"]


def test_lectura_sap_directa(c, monkeypatch):
    import time
    from app.integrations import sap
    monkeypatch.setattr(sap, "leer_pedido", lambda ped, pu, f: sap.Lectura(pedido=ped, posiciones=[
        sap.Posicion("000900081624", 11, 11), sap.Posicion("900276672", 10, 100)]))
    r = c.post("/api/sap/leer_pedido", json={"pedido": "4005171502", "puesto": "pn01",
                                             "fecha": "2026-09-17", "cliente": "paris"})
    assert r.status_code == 202
    tid = r.json()["id"]
    for _ in range(50):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok" and t["datos"]["agregadas"] == 2
    ped = _estado(c)["pedidos"][0]
    assert ped["cliente"] == "PARIS"
    assert {l["sku"]: l["qty"] for l in ped["lineas"]} == {"900081624": 11, "900276672": 100}
    # validaciones
    assert c.post("/api/sap/leer_pedido", json={"pedido": "abc", "puesto": "PN01"}).status_code == 422
    assert c.post("/api/sap/leer_pedido", json={"pedido": "4001", "puesto": "!!"}).status_code == 422


def test_asignar_visor_a_camion(c):
    c.put("/api/pedidos/4001", json={"pedido": "4001", "cliente": "PARIS", "lineas": [{"sku": "1", "qty": 1}]})
    a = c.post("/api/archivos", data={"pedido": "4001", "tipo": "adjunto"},
               files={"file": ("visor.html", b"<html>visor</html>", "text/html")}).json()
    assert c.patch(f"/api/archivos/{a['id']}", json={"grupo": "777"}).json()["grupo"] == "777"
    r = c.get(f"/archivos/{a['id']}")
    assert r.headers["content-disposition"].startswith("inline")


def test_descripcion_desde_maestra(c, monkeypatch, tmp_path):
    from openpyxl import Workbook
    from app.integrations import maestra
    wb = Workbook(); ws = wb.active; ws.title = "MAESTRA HOMOLOGACION"
    ws.append(["RUT proveedor", "COD. PROD. PROVEEDOR (ELUX)", "Canal", "Descripción", "NOMBRE CLIENTE"])
    ws.append([1, 900276671, "RETAIL", "COCINA FE5SXC", "PARIS"])
    ws.append([1, "900276671", "RETAIL", "otra descripcion", "HITES"])   # duplicado: manda la primera
    ruta = tmp_path / "maestra.xlsx"; wb.save(ruta)
    monkeypatch.setattr(maestra.settings.__class__, "maestra_homologacion", str(ruta), raising=False)
    object.__setattr__(maestra.settings, "maestra_homologacion", str(ruta))
    maestra._estado.update(mtime=None, revisado=0.0, datos={})
    c.put("/api/pedidos/4001", json={"pedido": "4001", "cliente": "PARIS",
          "lineas": [{"sku": "900276671", "qty": 27}, {"sku": "555", "desc": "manual", "qty": 1}]})
    lineas = {l["sku"]: l["desc"] for l in _estado(c)["pedidos"][0]["lineas"]}
    assert lineas == {"900276671": "COCINA FE5SXC", "555": "manual"}
    assert c.get("/api/salud").json()["maestra"]["productos"] == 1


# ---------------- Análisis (fórmulas de 02_Posiciones) ----------------
def test_alertas_como_excel():
    from app.analisis import alerta
    assert alerta(0, 27, 48) == "Sin stock"          # caso real: J=0
    assert alerta(30, 20, 25) == "Completo"
    assert alerta(30, 20, 10) == "Limitado SOP"       # T<K y T<=J
    assert alerta(5, 20, 25) == "Stock parcial"


def test_calculo_caso_real():
    from app.analisis import calcular, en_entrega_por_modelo, grupo_sop, codigo_cliente
    assert grupo_sop("paris") == "PARIS" and codigo_cliente("PARIS") == "266566"
    assert grupo_sop("LA POLAR") == "ABC" and grupo_sop("Copelec") == "REGION 3"
    zsd = [{"Nombre Codigo de Material": "COCINA FM5SSC", "Qty. En Entrega": 5},
           {"Nombre Codigo de Material": "cocina  fm5ssc", "Qty. En Entrega": "3"},
           {"Nombre Codigo de Material": "OTRA", "Qty. En Entrega": 99}]
    ent = en_entrega_por_modelo(zsd)
    med = {"900276671": {"desc": "900276671 COCINA FM5SSC", "max_camion": 270},
           "111": {"desc": "111 LAVADORA X", "max_camion": 100}}
    plan = {"900276671": {"plan": 126, "vendido": 78}, "111": {"plan": 50, "vendido": 0}}
    pos = [{"sku": "111", "qty_entrega": 40, "qty_pendiente": 40},
           {"sku": "900276671", "qty_entrega": 0, "qty_pendiente": 27},
           {"sku": "999", "qty_entrega": 10, "qty_pendiente": 10}]
    r = calcular(pos, plan, ent, med, {"900276671": {"cantidad": 132, "fecha": "2026-09-23"}})
    f = {x["sku"]: x for x in r["filas"]}
    assert f["900276671"]["en_entrega"] == 8 and f["900276671"]["saldo"] == 126 - 78 - 8
    assert f["900276671"]["alerta"] == "Sin stock" and f["900276671"]["carga"] == 0
    assert f["900276671"]["disponibilidad"] == "132 - 2026-09-23"
    assert f["111"]["alerta"] == "Completo" and f["111"]["carga"] == 40 and f["111"]["ocupacion"] == 0.4
    assert f["999"]["descripcion"] == "Producto no Encontrado" and f["999"]["ocupacion"] == 0
    assert [x["sku"] for x in r["filas"]][0] == "900276671"        # ordenado por alerta
    assert r["alertadas"] == ["900276671", "999"] or set(r["alertadas"]) == {"900276671", "999"}
    # ajuste manual de carga
    r2 = calcular(pos, plan, ent, med, {}, ajustes={"111": 20})
    assert {x["sku"]: x for x in r2["filas"]}["111"]["ocupacion"] == 0.2


def test_analisis_en_un_clic(c, monkeypatch):
    import time
    from app.integrations import sap, bases, base_medidas
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(pedido=p, posiciones=[
        sap.Posicion("900276671", 0, 27), sap.Posicion("111", 40, 40)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda cli, mats, carpeta, nombre, **kw: [
        {"Nombre Codigo de Material": "COCINA FM5SSC", "Qty. En Entrega": 0}])
    llamados = {}
    def mmbe(skus, avance=None):
        llamados["skus"] = skus
        return {s: {"cd30": 5, "reserva_cd30": 1, "ec01": 0, "tp01": 2} for s in skus}
    monkeypatch.setattr(sap, "mmbe", mmbe)
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900276671": {"plan": 126, "vendido": 78, "pdte_mes": 76, "saldo": 0, "tipo": "CONSENSO"},
                                                      "111": {"plan": 50, "vendido": 0, "pdte_mes": 0, "saldo": 0, "tipo": "CONSENSO"}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(bases, "oc_de", lambda pedido: "7788990011")        # OC de Pedidos Ingresados
    monkeypatch.setattr(base_medidas, "medidas", lambda: {"900276671": {"desc": "900276671 COCINA FM5SSC", "max_camion": 270},
                                                          "111": {"desc": "111 LAVADORA X", "max_camion": 100}})
    r = c.post("/api/analisis/4005171502", json={"puesto": "PN01", "cliente": "PARIS", "fecha": "2026-09-22"})
    assert r.status_code == 202
    tid = r.json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok", t
    assert llamados["skus"] == ["900276671"]                     # MMBE solo para alertados
    a = c.get("/api/analisis/4005171502").json()
    f = {x["sku"]: x for x in a["resultado"]["filas"]}
    assert f["900276671"]["saldo"] == 48 and f["900276671"]["stock"]["cd30"] == 5
    assert _estado(c)["pedidos"][0]["cliente"] == "PARIS"
    assert _estado(c)["pedidos"][0]["oc"] == "7788990011"        # la OC sale de Pedidos Ingresados
    # ajustar carga y volver al calculado
    a = c.put("/api/analisis/4005171502/carga", json={"sku": "111", "carga": 20}).json()
    assert {x["sku"]: x for x in a["resultado"]["filas"]}["111"]["carga"] == 20
    a = c.put("/api/analisis/4005171502/carga", json={"sku": "111", "carga": ""}).json()
    assert {x["sku"]: x for x in a["resultado"]["filas"]}["111"]["carga"] == 40
    assert c.put("/api/analisis/4005171502/carga", json={"sku": "111", "carga": -3}).status_code == 422


def test_cubicaje_desde_analisis(c, monkeypatch, tmp_path):
    from app.integrations import base_medidas, bases, sap
    from app.config import settings
    # análisis previo (la carga del cubicaje sale de ahí)
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(pedido=p, posiciones=[
        sap.Posicion("900081624", 30, 30), sap.Posicion("900276671", 20, 20)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 500, "vendido": 0},
                                                      "900276671": {"plan": 500, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    medidas = {"900081624": {"desc": "900081624 MDWMT16W", "max_camion": 90},
               "900276671": {"desc": "900276671 COCINA FM5SSC", "max_camion": 80}}
    monkeypatch.setattr(base_medidas, "medidas", lambda: medidas)
    monkeypatch.setattr(base_medidas, "filas", lambda: [
        ["900081624", "900081624 MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4],
        ["900276671", "900276671 COCINA FM5SSC", 1, 62, 66, 90, 40, "Y", "N", "N", 80, 3]])
    import time
    r = c.post("/api/analisis/4001", json={"puesto": "PN01", "cliente": "PARIS"})
    tid = r.json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok", t

    plantilla = tmp_path / "Plantilla_Visor.html"
    plantilla.write_text("<html>VISOR __CUBICAJE_JSON__</html>", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    d = c.post("/api/cubicaje/4001", json={}).json()
    assert d["modo"] == "MDA" and d["unidades"] == 50 and d["camiones"]
    assert d["filas"][0]["tipo_carga"] == "Mono-pedido"
    assert d["visor"].startswith("/visor/pedido_4001_")
    assert "VISOR" in c.get(d["visor"]).text and '"camiones"' in c.get(d["visor"]).text
    assert c.get("/api/cubicaje/4001").json()["unidades"] == 50
    # los modos SDA también funcionan y entregan el detalle por pallet
    d2 = c.post("/api/cubicaje/4001", json={"modo": "SDA STOCK", "caja_master": "SIN CAJA MASTER"}).json()
    assert d2["modo"] == "SDA STOCK" and d2["filas04"] and d2["pallets_detalle"]
    assert sum(f["unidades"] for f in d2["filas04"]) == 50
    assert c.post("/api/cubicaje/9999", json={}).status_code == 422

    # Ajustar la carga de un pedido ya cubicado rehace el camión en el mismo paso
    a = c.put("/api/analisis/4001/carga", json={"sku": "900081624", "carga": 10})
    assert a.status_code == 200, a.text
    cub = a.json()["cubicaje"]
    assert cub and not cub.get("error"), cub
    assert cub["unidades"] == 30                 # 10 ajustadas + 20 de la otra posición
    assert cub["modo"] == "SDA STOCK"            # conserva el modo del cubicaje anterior


def test_visor_no_se_acumula(c, monkeypatch, tmp_path):
    """El visor se sirve desde /visor con sus librerías y no se acumulan versiones."""
    from app.config import settings
    from app.integrations import base_medidas, bases, sap
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(
        pedido=p, posiciones=[sap.Posicion("900081624", 10, 10)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 100, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    filas = [["900081624", "900081624 MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4]]
    monkeypatch.setattr(base_medidas, "filas", lambda: filas)
    monkeypatch.setattr(base_medidas, "medidas", lambda: {"900081624": {"desc": "MDWMT16W", "max_camion": 90}})
    origen = tmp_path / "Visor"
    origen.mkdir()
    (origen / "Plantilla_Visor.html").write_text(
        '<html><script src="visor/three.min.js"></script><script>const D=__CUBICAJE_JSON__;</script></html>',
        encoding="utf-8")
    for lib in ("three.min.js", "jspdf.min.js", "gltf_loader.js", "scania_data.js"):
        (origen / lib).write_text(f"// {lib}", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(origen / "Plantilla_Visor.html"))
    object.__setattr__(settings, "visor_assets", str(origen))

    import time
    tid = c.post("/api/analisis/4002", json={"puesto": "PN01", "cliente": "PARIS"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    v1 = c.post("/api/cubicaje/4002", json={}).json()["visor"]
    v2 = c.post("/api/cubicaje/4002", json={}).json()["visor"]
    assert v1.startswith("/visor/pedido_4002_") and v1 != v2
    carpeta = Path(settings.visores_dir)
    assert len(list(carpeta.glob("pedido_4002_*.html"))) == 1        # no se acumulan
    assert (carpeta / "visor" / "three.min.js").exists()             # librerías copiadas
    r = c.get(v2)
    assert r.status_code == 200 and "camiones" in r.text
    assert c.get("/visor/visor/three.min.js").status_code == 200     # las sirve la plataforma


def _archivo_medidas(tmp_path, filas):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Base para carga"
    ws.append(["EasyCargo plantilla de la importación"])
    ws.append(["Grupo", "Descripción", "Piezas", "Longitud", "Anchura", "Altura", "Peso total",
               "Apilar", "Inclinar", "Rotar", "Máx Camión", "Máx Pallet"])
    ws.append([None, None, None, "centímetros", None, None, "kilogramos", "Y/N", "Y/N", "Y/N"])
    for f in filas:
        ws.append(f)
    ruta = tmp_path / "Base de Medidas.xlsm"; wb.save(ruta)
    return ruta


def test_carga_masiva_de_medidas(c, tmp_path):
    f1 = _archivo_medidas(tmp_path, [
        ["0900081624", "MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4],
        ["900276671", "COCINA FM5SSC", 1, 62, 66, 90, 40, "y", "n", "n", 80, 3],
        ["SIN-MEDIDAS", "producto incompleto", 1, 0, 0, 0, 0, "N", "N", "N", 0, 0],
    ])
    with open(f1, "rb") as fh:
        r = c.post("/api/medidas/importar", files={"file": ("Base de Medidas.xlsm", fh.read())}).json()
    assert r["nuevos"] == 2 and r["ignorados"] == 1 and r["productos"] == 2
    assert r["ejemplos_ignorados"] == ["SIN-MEDIDAS"]

    filas = {m["sku"]: m for m in c.get("/api/medidas").json()["filas"]}
    assert set(filas) == {"900081624", "900276671"}          # ceros a la izquierda normalizados
    assert filas["900081624"]["largo"] == 60 and filas["900276671"]["apilar"] == "Y"

    # segunda carga: uno cambia, otro se agrega, el resto queda igual
    f2 = _archivo_medidas(tmp_path, [
        ["0900081624", "MDWMT16W", 1, 61, 65, 85, 38.5, "Y", "N", "N", 90, 4],
        ["900276671", "COCINA FM5SSC", 1, 62, 66, 90, 40, "Y", "N", "N", 80, 3],
        ["900999999", "PRODUCTO NUEVO", 1, 40, 40, 40, 10, "Y", "N", "N", 100, 10],
    ])
    with open(f2, "rb") as fh:
        r = c.post("/api/medidas/importar", files={"file": ("Base de Medidas.xlsm", fh.read())}).json()
    assert (r["nuevos"], r["actualizados"], r["sin_cambios"]) == (1, 1, 1)
    assert r["productos"] == 3 and r["ultima_carga"]

    # corrección manual de un producto
    r = c.put("/api/medidas/900999999", json={"descripcion": "CORREGIDO", "piezas": 1, "largo": 45,
                                              "ancho": 40, "alto": 40, "peso": 10, "apilar": "Y",
                                              "inclinar": "N", "rotar": "N", "max_camion": 100,
                                              "max_pallet": 10}).json()
    assert r["filas"][0]["largo"] == 45 and r["filas"][0]["descripcion"] == "CORREGIDO"
    assert c.put("/api/medidas/900999999", json={"largo": 0}).status_code == 422
    assert c.post("/api/medidas/importar", files={"file": ("x.txt", b"nada")}).status_code == 422
    assert c.get("/api/medidas?buscar=COCINA").json()["filas"][0]["sku"] == "900276671"
    assert c.get("/api/salud").json()["medidas_cargadas"]["productos"] == 3


def test_predistribuido_y_modo_por_sucursal(c, monkeypatch, tmp_path):
    import time
    from app.config import settings
    from app.integrations import base_medidas, bases, sap
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(pedido=p, posiciones=[
        sap.Posicion("900081624", 40, 40), sap.Posicion("900276671", 40, 40)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 500, "vendido": 0},
                                                      "900276671": {"plan": 500, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    filas = [["900081624", "900081624 MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4],
             ["900276671", "900276671 COCINA FM5SSC", 1, 62, 66, 90, 40, "Y", "N", "N", 80, 3]]
    monkeypatch.setattr(base_medidas, "filas", lambda: filas)
    monkeypatch.setattr(base_medidas, "medidas", lambda: {r[0]: {"desc": r[1], "max_camion": r[10]} for r in filas})
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    object.__setattr__(settings, "visor_assets", str(tmp_path))

    tid = c.post("/api/analisis/4010", json={"puesto": "PN01", "cliente": "PARIS"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)

    # sin tabla de reparto, el modo avisa
    r = c.post("/api/cubicaje/4010", json={"modo": "MDA PREDISTRIBUIDO", "caja_master": "SIN CAJA MASTER"})
    assert r.status_code == 422 and "Predistribuido" in r.json()["detail"]

    # se pega la tabla (acepta tabulaciones, punto y coma y encabezado)
    texto = "Sucursal\tSKU\tUnidades\nSUC-01\t0900081624\t20\nSUC-02;900276671;15\nSUC-01,900276671,mal"
    d = c.put("/api/predistribuido/4010", json={"texto": texto}).json()
    assert d["sucursales"] == ["SUC-01", "SUC-02"] and d["unidades"] == 35
    assert d["filas"][0]["sku"] == "900081624"        # ceros a la izquierda normalizados
    assert len(d["errores"]) == 1 and "no es una cantidad" in d["errores"][0]

    d = c.post("/api/cubicaje/4010", json={"modo": "MDA PREDISTRIBUIDO", "caja_master": "SIN CAJA MASTER"}).json()
    assert d["modo"] == "MDA PREDISTRIBUIDO" and d["unidades"] == 35
    assert {f["sucursal"] for f in d["filas"]} == {"SUC-01", "SUC-02"}
    assert c.get("/api/predistribuido/4010").json()["unidades"] == 35

    # el mismo reparto, ahora desde la plantilla Excel (reemplaza la tabla anterior)
    from openpyxl import Workbook
    libro = Workbook(); ws = libro.active; ws.title = "Predistribuido"
    ws.append(["Sucursal", "SKU", "Unidades", "Unidades por bulto (opcional)"])
    ws.append(["suc-01", "0900081624", 20, None])     # minúsculas y ceros a la izquierda
    ws.append(["SUC-03", 900276671, 25, 5])           # SKU como número, con bulto del cliente
    ws.append(["SUC-04", "900276671", "mal", None])   # cantidad inválida: se informa
    ws.append([None, None, None, None])               # fila vacía: se salta
    ws.append(["SUC-05", "900081624", 0, None])       # en cero: se ignora
    ruta = tmp_path / "reparto.xlsx"; libro.save(ruta)
    with open(ruta, "rb") as fh:
        r = c.post("/api/predistribuido/4010/importar", files={"file": ("reparto.xlsx", fh.read())})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["sucursales"] == ["SUC-01", "SUC-03"] and d["unidades"] == 45
    assert [f["por_bulto"] for f in d["filas"]] == [0, 5]
    assert len(d["errores"]) == 1 and d["errores"][0].startswith("fila 4")
    assert c.get("/api/predistribuido/4010").json()["unidades"] == 45

    # las unidades por bulto llegan al motor (las usa SDA Predistribuido)
    from app.cubicaje import motor
    recibido = {}
    original = motor.cubicaje_sda_predistribuido
    def espia(posiciones, predist, *a, **k):
        recibido["filas"] = predist
        return original(posiciones, predist, *a, **k)
    monkeypatch.setattr(motor, "cubicaje_sda_predistribuido", espia)
    r = c.post("/api/cubicaje/4010", json={"modo": "SDA PREDISTRIBUIDO", "caja_master": "SIN CAJA MASTER"})
    assert r.status_code == 200, r.text
    assert [(f.sucursal, f.por_bulto) for f in recibido["filas"]] == [("SUC-01", 0), ("SUC-03", 5)]


def test_plantilla_de_predistribuido(c, tmp_path):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    r = c.get("/api/predistribuido/plantilla")
    assert r.status_code == 200 and "plantilla_predistribuido.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Predistribuido", "Cómo se usa"]
    assert [x.value for x in wb["Predistribuido"][1]] == [
        "Sucursal", "SKU", "Unidades", "Unidades por bulto (opcional)"]

    # la plantilla tal como se descarga se puede importar sin tocarla
    d = c.post("/api/predistribuido/4030/importar",
               files={"file": ("plantilla_predistribuido.xlsx", r.content)}).json()
    assert d["sucursales"] == ["SUC-01", "SUC-02"] and d["unidades"] == 30 and d["errores"] == []

    # las columnas se ubican por nombre: sirve aunque el cliente cambie el orden
    from openpyxl import Workbook as _WB
    otro = _WB(); ws = otro.active
    ws.append(["Código", "Tienda", "Descripción", "Cantidad"])
    ws.append(["900081624", "suc-09", "MICROONDAS", 7])
    ruta = tmp_path / "otro.xlsx"; otro.save(ruta)
    with open(ruta, "rb") as fh:
        d = c.post("/api/predistribuido/4031/importar", files={"file": ("otro.xlsx", fh.read())}).json()
    assert [(f["sucursal"], f["sku"], f["unidades"]) for f in d["filas"]] == [("SUC-09", "900081624", 7)]

    # si falta una columna obligatoria, se dice cuál
    sin_cant = _WB(); sin_cant.active.append(["Sucursal", "SKU"]); sin_cant.active.append(["S1", "900081624"])
    ruta = tmp_path / "sin.xlsx"; sin_cant.save(ruta)
    with open(ruta, "rb") as fh:
        r = c.post("/api/predistribuido/4032/importar", files={"file": ("sin.xlsx", fh.read())})
    assert r.status_code == 422 and "Unidades" in r.json()["detail"]

    # sin encabezado se asume el orden de la plantilla
    crudo = _WB(); crudo.active.append(["SUC-01", "900081624", 4])
    ruta = tmp_path / "crudo.xlsx"; crudo.save(ruta)
    with open(ruta, "rb") as fh:
        d = c.post("/api/predistribuido/4033/importar", files={"file": ("crudo.xlsx", fh.read())}).json()
    assert d["unidades"] == 4

    # formato equivocado y archivo sin filas válidas: se rechazan sin tocar lo guardado
    assert c.post("/api/predistribuido/4020/importar",
                  files={"file": ("x.txt", b"nada")}).status_code == 422
    vacio = Workbook(); vacio.active.append(["Sucursal", "SKU", "Unidades"])
    vacio.active.append(["SUC-01", "900081624", "mal"])
    ruta = tmp_path / "vacio.xlsx"; vacio.save(ruta)
    with open(ruta, "rb") as fh:
        r = c.post("/api/predistribuido/4020/importar", files={"file": ("vacio.xlsx", fh.read())})
    assert r.status_code == 422 and "fila 2" in r.json()["detail"]
    assert c.get("/api/predistribuido/4020").json()["filas"] == []


def test_cubicador_libre_usa_la_columna_sucursal(c, tmp_path):
    """Antes la columna Sucursal de la plantilla se ignoraba y fundía los SKU repetidos."""
    from app.config import settings
    from openpyxl import Workbook
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    libro = Workbook(); ws = libro.active; ws.title = "Carga"
    ws.append(["SKU", "Unidades", "Camión fijo (opcional)", "Sucursal (opcional)"])
    ws.append(["900081624", 20, "", "SUC-01"])
    ws.append(["900081624", 10, "", "suc-02"])         # mismo SKU, otra sucursal
    ws.append(["900276671", 15, "", "SUC-01"])
    ws.append(["900276671", 5, "", ""])               # le falta la sucursal: se informa
    ruta = tmp_path / "carga.xlsx"; libro.save(ruta)
    with open(ruta, "rb") as fh:
        r = c.post("/api/cubicaje-libre/importar", files={"file": ("carga.xlsx", fh.read())})
    assert r.status_code == 200, r.text
    d = r.json()
    assert {l["sku"]: l["qty"] for l in d["lineas"]} == {"900081624": 30, "900276671": 15}
    assert [(p["sucursal"], p["sku"], p["qty"]) for p in d["predistribuido"]] == [
        ("SUC-01", "900081624", 20), ("SUC-02", "900081624", 10), ("SUC-01", "900276671", 15)]
    assert d["modo"] == "MDA PREDISTRIBUIDO"                # pasa solo al modo que usa sucursales
    assert any("predistribuido" in a.lower() for a in d["avisos"])
    assert any(e.startswith("fila 5") for e in d["errores_archivo"])
    assert d["unidades"] == 45
    assert {f["sucursal"] for f in d["filas"]} == {"SUC-01", "SUC-02"}


def test_borrado_en_sap_requiere_confirmacion(c, monkeypatch):
    import time
    from app.integrations import sap
    c.put("/api/pedidos/4100", json={"pedido": "4100", "cliente": "PARIS",
                                     "lineas": [{"sku": "1", "qty": 5}]})
    c.put("/api/entregas/8705709527", json={"entrega": "8705709527", "pedido": "4100",
                                            "grupo": "1392270", "lineas": [{"sku": "1", "qty": 5}]})
    # sin confirmar no se toca SAP
    r = c.post("/api/sap/borrar_entrega", json={"entrega": "8705709527"})
    assert r.status_code == 422 and "confirmar" in r.json()["detail"]
    assert c.post("/api/sap/borrar_entrega", json={"entrega": "abc", "confirmar": "abc"}).status_code == 422

    llamadas = []
    monkeypatch.setattr(sap, "borrar_entrega", lambda n, ses=None: (llamadas.append(n), (True, "ok"))[1])
    tid = c.post("/api/sap/borrar_entrega", json={"entrega": "8705709527",
                                                  "confirmar": "8705709527"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok" and llamadas == ["8705709527"]
    assert _estado(c)["entregas"] == []          # se borró en SAP y después en la plataforma


def test_borrado_falla_en_sap_y_la_entrega_se_conserva(c, monkeypatch):
    import time
    from app.integrations import sap
    c.put("/api/pedidos/4101", json={"pedido": "4101", "cliente": "PARIS", "lineas": [{"sku": "1", "qty": 5}]})
    c.put("/api/entregas/8705709528", json={"entrega": "8705709528", "pedido": "4101",
                                            "lineas": [{"sku": "1", "qty": 5}]})
    monkeypatch.setattr(sap, "borrar_entrega", lambda n, ses=None: (False, "Documento facturado"))
    tid = c.post("/api/sap/borrar_entrega", json={"entrega": "8705709528",
                                                  "confirmar": "8705709528"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "error" and "facturado" in t["error"]
    assert len(_estado(c)["entregas"]) == 1     # no se borró nada en la plataforma


def test_borrar_grupo_deja_las_entregas_sin_grupo(c, monkeypatch):
    import time
    from app.integrations import sap
    c.put("/api/pedidos/4102", json={"pedido": "4102", "cliente": "PARIS", "lineas": [{"sku": "1", "qty": 5}]})
    for n in ("8705709530", "8705709531"):
        c.put(f"/api/entregas/{n}", json={"entrega": n, "pedido": "4102", "grupo": "1392270",
                                          "lineas": [{"sku": "1", "qty": 2}]})
    monkeypatch.setattr(sap, "borrar_grupo", lambda g, ses=None: (True, "Grupo borrado"))
    tid = c.post("/api/sap/borrar_grupo", json={"grupo": "1392270",
                                                "confirmar": "1392270"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok"
    ents = _estado(c)["entregas"]
    assert len(ents) == 2 and all(e["grupo"] == "" for e in ents)      # siguen, pero sin grupo
    assert all(any("borrado en SAP" in x["txt"] for x in e["log"]) for e in ents)


def test_reglas_por_cliente_editables(c):
    d = c.get("/api/clientes").json()["filas"]
    por_nombre = {x["nombre"]: x for x in d}
    assert por_nombre["PARIS"]["pallet"] == [120, 100, 150]
    assert por_nombre["HITES"]["calefon_aparte"] is True
    assert por_nombre["SODIMAC"]["hibrido"] is True and por_nombre["PARIS"]["hibrido"] is False
    assert por_nombre["PARIS"]["codigo"] == "266566"

    # cambiar el pallet de un cliente
    r = c.put("/api/clientes/PARIS", json={"pallet": [120, 100, 170], "notas": "pallet alto"}).json()
    assert r["pallet"] == [120, 100, 170] and r["notas"] == "pallet alto"
    assert c.put("/api/clientes/PARIS", json={"pallet": [0, 100, 150]}).status_code == 422

    # cliente nuevo sin tocar código
    r = c.put("/api/clientes/CLIENTE NUEVO", json={"grupo_sop": "NUEVO", "codigo": "999999",
                                                   "pallet": [110, 90, 130], "hibrido": True}).json()
    assert r["hibrido"] is True and r["pallet"] == [110, 90, 130]
    assert "CLIENTE NUEVO" in [x["nombre"] for x in c.get("/api/clientes").json()["filas"]]


def test_cubicaje_usa_el_pallet_del_cliente(c, monkeypatch, tmp_path):
    import time
    from app.config import settings
    from app.integrations import base_medidas, bases, sap
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(
        pedido=p, posiciones=[sap.Posicion("900081624", 20, 20)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 100, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    filas = [["900081624", "900081624 MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4]]
    monkeypatch.setattr(base_medidas, "filas", lambda: filas)
    monkeypatch.setattr(base_medidas, "medidas", lambda: {"900081624": {"desc": "x", "max_camion": 90}})
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    c.put("/api/clientes/PARIS", json={"pallet": [120, 100, 165]})
    tid = c.post("/api/analisis/4200", json={"puesto": "PN01", "cliente": "PARIS"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    d = c.post("/api/cubicaje/4200", json={"modo": "SDA STOCK", "caja_master": "SIN CAJA MASTER"}).json()
    assert d["pallet"] == [120, 100, 165]        # toma el pallet de la tabla de clientes


def _cargar_medidas_basicas(c, tmp_path):
    f = _archivo_medidas(tmp_path, [
        ["900081624", "900081624 MDWMT16W", 1, 60, 65, 85, 38.5, "Y", "N", "N", 90, 4],
        ["C900081624", "CAJA MASTER MDWMT16W", 4, 120, 65, 85, 154, "Y", "N", "N", 40, 2],
        ["900276671", "900276671 COCINA FM5SSC", 1, 62, 66, 90, 40, "Y", "N", "N", 80, 3],
    ])
    with open(f, "rb") as fh:
        c.post("/api/medidas/importar", files={"file": ("Base de Medidas.xlsm", fh.read())})


def test_autocompletar_sku(c, tmp_path):
    _cargar_medidas_basicas(c, tmp_path)
    s = c.get("/api/medidas/sugerir?q=MDWMT").json()["sugerencias"]
    skus = [x["sku"] for x in s]
    assert "900081624" in skus and "C900081624" in skus          # sugiere también la caja master
    caja = next(x for x in s if x["sku"] == "C900081624")
    assert caja["caja_master"] is True and caja["piezas"] == 4
    assert c.get("/api/medidas/sugerir?q=9").json()["sugerencias"] == []   # muy corto, no busca
    assert [x["sku"] for x in c.get("/api/medidas/sugerir?q=COCINA").json()["sugerencias"]] == ["900276671"]


def test_cubicador_libre(c, tmp_path, monkeypatch):
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    # carga a mano, sin pedido y sin cliente
    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "900081624", "qty": 30}],
                                            "vista": "rampla"}).json()
    assert d["unidades"] == 30 and len(d["camiones"]) == 1 and d["visor"]
    assert d["camiones"][0]["tipo"] == "Rampla 53"

    # cambiar de vista a un solo pallet: se arma en pallets y se muestra uno
    d = c.post("/api/cubicaje-libre", json={**d, "vista": "pallet",
                                            "caja_master": "SIN CAJA MASTER"}).json()
    assert d["modo_usado"] == "SDA STOCK"          # en MDA no hay pallets: usa el motor de pallets
    assert d["pallet_visto"] >= 1 and len(d["pallets_disponibles"]) >= 1
    assert d["pallets_detalle"] and d["unidades"] == 30

    # con cliente toma su pallet y sus reglas
    c.put("/api/clientes/PARIS", json={"pallet": [120, 100, 150]})
    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "900081624", "qty": 10}],
                                            "cliente": "PARIS", "vista": "pallet"}).json()
    assert d["pallet"] == [120, 100, 150]

    # SKU sin medidas: se avisa y no rompe
    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "999999", "qty": 5},
                                                       {"sku": "900276671", "qty": 4}]}).json()
    assert d["desconocidos"] == ["999999"] and d["unidades"] == 4

    # carga vacía: igual entrega el visor del contenedor elegido
    d = c.post("/api/cubicaje-libre", json={"lineas": [], "vista": "camion50"}).json()
    assert d["unidades"] == 0 and d["visor"] and d["camiones"] == []
    assert c.get("/api/cubicaje-libre").json()["vista"] == "camion50"   # queda guardado

    # modo SDA con caja master
    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "900081624", "qty": 9}],
                                            "modo": "SDA STOCK", "caja_master": "CON CAJA MASTER",
                                            "cliente": "PARIS"}).json()
    assert sum(f["unidades"] for f in d["filas04"]) == 9
    assert any(f["cajas"] * 4 == f["unidades"] for f in d["filas04"])   # 2 cajas de 4 + 1 suelta


def test_ajustes_de_cubicaje(c, tmp_path):
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    aj = c.get("/api/ajustes-cubicaje").json()
    assert {k: aj[k] for k in ("orientacion_pallet", "celda_cm", "capacidad_pallet")} == {
        "orientacion_pallet": "largo", "celda_cm": 1, "capacidad_pallet": "geometria"}     # por defecto 1 cm: más fiel a la carga real
    assert aj["camiones"]["rampla"] == {"largo": 1540.0, "ancho": 245.0, "alto": 235.0}     # la rampla 53 mide 2350 mm de alto
    assert aj["camiones"]["camion50"] == {"largo": 620.0, "ancho": 244.0, "alto": 230.0}
    assert c.put("/api/ajustes-cubicaje", json={"orientacion_pallet": "otro"}).status_code == 422
    assert c.put("/api/ajustes-cubicaje", json={"celda_cm": 5}).status_code == 422

    carga = {"lineas": [{"sku": "900276671", "qty": 60}], "modo": "SDA STOCK",
             "caja_master": "SIN CAJA MASTER", "cliente": "PARIS", "vista": "pallet"}
    c.put("/api/clientes/PARIS", json={"pallet": [120, 100, 150]})
    d1 = c.post("/api/cubicaje-libre", json=carga).json()
    pallets_excel = len(d1["pallets_detalle"])

    c.put("/api/ajustes-cubicaje", json={"orientacion_pallet": "excel", "celda_cm": 2})
    d2 = c.post("/api/cubicaje-libre", json=carga).json()
    assert d2["ajustes"]["orientacion_pallet"] == "excel"
    assert len(d2["pallets_detalle"]) >= pallets_excel      # la del Excel usa peor el pallet
    assert sum(f["unidades"] for f in d2["filas04"]) == 60  # y no se pierde carga
    c.put("/api/ajustes-cubicaje", json={"orientacion_pallet": "excel", "celda_cm": 2})


def test_cubicador_desde_pedido_y_excel(c, tmp_path, monkeypatch):
    import time
    from app.config import settings
    from app.integrations import base_medidas, bases, sap
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    # sin análisis previo, avisa
    r = c.post("/api/cubicaje-libre/desde-pedido", json={"pedido": "4300"})
    assert r.status_code == 422 and "no está analizado" in r.json()["detail"]

    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(pedido=p, posiciones=[
        sap.Posicion("900081624", 30, 30), sap.Posicion("900276671", 12, 12)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 500, "vendido": 0},
                                                      "900276671": {"plan": 500, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(base_medidas, "medidas", lambda: {})
    tid = c.post("/api/analisis/4300", json={"puesto": "PN01", "cliente": "PARIS"}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)

    d = c.post("/api/cubicaje-libre/desde-pedido", json={"pedido": "4300"}).json()
    assert d["pedido"] == "4300" and d["cliente"] == "PARIS"
    assert {l["sku"] for l in d["lineas"]} == {"900081624", "900276671"}
    assert d["unidades"] == 42

    # Excel del cubicaje libre
    r = c.get("/api/cubicaje-libre/excel")
    assert r.status_code == 200 and r.content[:2] == b"PK"        # es un xlsx
    assert "cubicaje.xlsx" in r.headers["content-disposition"]
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames[0] == "Camiones"
    filas = list(wb["Camiones"].iter_rows(min_row=2, values_only=True))
    assert sum(f[5] for f in filas) == 42                         # unidades cuadran


def test_fijar_productos_a_un_camion(c, tmp_path):
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    libre = c.post("/api/cubicaje-libre", json={"lineas": [
        {"sku": "900081624", "qty": 20}, {"sku": "900276671", "qty": 20}]}).json()
    assert len(libre["camiones"]) == 1                            # los dos van juntos

    # El camión fijo se sacó del cubicador: una carga antigua que lo traiga se cubica normal
    antigua = c.post("/api/cubicaje-libre", json={"lineas": [
        {"sku": "900081624", "qty": 20, "camion": 1},
        {"sku": "900276671", "qty": 20}]}).json()
    assert len(antigua["camiones"]) == 1 and antigua["unidades"] == 40


def _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4400", cliente="PARIS"):
    import time
    from app.config import settings
    from app.integrations import base_medidas, bases, sap
    _cargar_medidas_basicas(c, tmp_path)
    monkeypatch.setattr(sap, "leer_pedido", lambda p, pu, f: sap.Lectura(pedido=p, posiciones=[
        sap.Posicion("900081624", 30, 30), sap.Posicion("900276671", 20, 20)]))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **kw: [])
    monkeypatch.setattr(sap, "mmbe", lambda skus, avance=None: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"900081624": {"plan": 500, "vendido": 0},
                                                      "900276671": {"plan": 500, "vendido": 0}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(base_medidas, "medidas", lambda: {
        "900081624": {"desc": "900081624 MDWMT16W", "max_camion": 90},
        "900276671": {"desc": "900276671 COCINA FM5SSC", "max_camion": 80}})
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    tid = c.post(f"/api/analisis/{pedido}", json={"puesto": "PN01", "cliente": cliente}).json()["id"]
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    return c.post(f"/api/cubicaje/{pedido}", json={}).json()


def _esperar_job(c, tid):
    import time
    for _ in range(200):
        t = c.get(f"/api/acciones/trabajos/{tid}").json()
        if t["estado"] != "en_curso":
            return t
        time.sleep(0.05)
    return t


def test_crear_entregas_ensayo_y_real(c, tmp_path, monkeypatch):
    from app.integrations import sap_crear
    cub = _preparar_cubicaje(c, tmp_path, monkeypatch)
    assert cub["camiones"]

    llamadas = []

    def falso(pedido, puesto, fecha, materiales, fecha_cita="", hora_cita="", ensayo=False, ses=None):
        llamadas.append({"pedido": pedido, "puesto": puesto, "ensayo": ensayo,
                         "materiales": dict(materiales)})
        r = sap_crear.Resultado(ok=True, ensayo=ensayo)
        if not ensayo:
            r.entrega = f"87057{len(llamadas):05d}"
        r.mensaje = "ok"
        return r

    monkeypatch.setattr(sap_crear, "crear_entrega", falso)

    # ensayo: no pide confirmación y no guarda entregas en la plataforma
    t = _esperar_job(c, c.post("/api/sap/crear_entregas",
                               json={"pedido": "4400", "ensayo": True}).json()["id"])
    assert t["estado"] == "ok" and all(x["ensayo"] for x in llamadas)
    assert llamadas[0]["puesto"] == "PN01"                    # lo toma del análisis
    assert sum(llamadas[0]["materiales"].values()) > 0
    assert _estado(c)["entregas"] == []

    # de verdad: crear ya no pide confirmación (solo borrar la pide)
    llamadas.clear()
    t = _esperar_job(c, c.post("/api/sap/crear_entregas",
                               json={"pedido": "4400", "ensayo": False}).json()["id"])
    assert t["estado"] == "ok", str(t)[:400]
    ents = _estado(c)["entregas"]
    assert len(ents) == len(cub["camiones"])                  # una entrega por camión
    assert all(e["pedido"] == "4400" for e in ents)


def test_crear_entregas_se_detiene_ante_un_error(c, tmp_path, monkeypatch):
    from app.integrations import sap_crear
    _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4401")
    monkeypatch.setattr(sap_crear, "crear_entrega",
                        lambda *a, **k: sap_crear.Resultado(ok=False, mensaje="Pedido bloqueado"))
    t = _esperar_job(c, c.post("/api/sap/crear_entregas",
                               json={"pedido": "4401", "ensayo": True}).json()["id"])
    assert t["estado"] == "error" and "bloqueado" in t["error"]
    assert _estado(c)["entregas"] == []
    assert c.post("/api/sap/crear_entregas", json={"pedido": "9999"}).status_code == 422


def test_crear_grupo_y_cita_por_grupo(c, monkeypatch):
    """El flujo es el del Excel: un grupo por camión y la cita se aplica al grupo."""
    from app.integrations import sap_crear
    c.put("/api/pedidos/4402", json={"pedido": "4402", "cliente": "PARIS",
                                     "lineas": [{"sku": "1", "qty": 5}]})
    for n in ("8705700001", "8705700002"):
        c.put(f"/api/entregas/{n}", json={"entrega": n, "pedido": "4402",
                                          "lineas": [{"sku": "1", "qty": 2}]})
    vistos = {}

    def falso_grupo(entregas, referencia="", ensayo=False, ses=None, cliente="", camion="1"):
        vistos.update({"entregas": list(entregas), "referencia": referencia, "ensayo": ensayo,
                       "cliente": cliente, "camion": camion})
        return sap_crear.Resultado(ok=True, ensayo=ensayo, grupo="" if ensayo else "1392270",
                                   mensaje="ok")

    monkeypatch.setattr(sap_crear, "crear_grupo", falso_grupo)
    t = _esperar_job(c, c.post("/api/sap/crear_grupo",
                               json={"entregas": ["8705700001", "8705700002"], "camion": "2",
                                     "cliente": "PARIS", "ensayo": True}).json()["id"])
    assert t["estado"] == "ok" and vistos["ensayo"] and vistos["camion"] == "2"
    assert all(e["grupo"] == "" for e in _estado(c)["entregas"])      # el ensayo no agrupa

    t = _esperar_job(c, c.post("/api/sap/crear_grupo",
                               json={"entregas": ["8705700001", "8705700002"], "ensayo": False,
                                     "camion": "2",
                                     "referencia": "CITA-123"}).json()["id"])
    assert t["estado"] == "ok" and vistos["referencia"] == "CITA-123"
    assert all(e["grupo"] == "1392270" for e in _estado(c)["entregas"])

    # la cita se aplica al grupo completo (VG02), no entrega por entrega
    monkeypatch.setattr(sap_crear, "actualizar_grupo",
                        lambda g, f="", h="", ref="", ensayo=False, ses=None:
                        sap_crear.Resultado(ok=True, ensayo=ensayo, grupo=g, mensaje="ok"))
    t = _esperar_job(c, c.post("/api/sap/fecha_grupo",
                               json={"grupo": "1392270", "fecha": "26.09.2026", "hora": "10:00:00",
                                     "ensayo": True}).json()["id"])
    assert t["estado"] == "ok"
    assert all(not e["cita"]["fecha"] for e in _estado(c)["entregas"])   # el ensayo no escribe

    t = _esperar_job(c, c.post("/api/sap/fecha_grupo",
                               json={"grupo": "1392270", "fecha": "26.09.2026", "hora": "10:00:00",
                                     "referencia": "CITA-123", "ensayo": False}).json()["id"])
    assert t["estado"] == "ok"
    ents = _estado(c)["entregas"]
    assert all(e["cita"]["fecha"] == "2026-09-26" and e["cita"]["hora"] == "10:00" and
               e["cita"]["numero"] == "CITA-123" for e in ents)
    assert c.post("/api/sap/fecha_grupo", json={"grupo": "1392270"}).status_code == 422


def test_referencia_del_grupo_como_el_excel():
    from app.integrations.sap_crear import referencia_grupo
    assert referencia_grupo("1", ["8705709527"], "PARIS") == "CAM1_8705709527_PARIS"
    assert len(referencia_grupo("12", ["8705709527", "8705709528", "8705709529"], "SODIMAC")) <= 30


def test_respaldo_automatico(c, tmp_path, monkeypatch):
    from app import respaldo
    monkeypatch.setattr(respaldo, "CARPETA", tmp_path / "respaldos")
    r = c.post("/api/respaldos").json()
    assert r["ok"] and r["archivo"].startswith("trazabilidad-") and r["tamano_kb"] >= 0
    assert c.get("/api/respaldos").json()["copias"][0]["archivo"] == r["archivo"]

    # repetir en el mismo día reemplaza, no acumula
    c.post("/api/respaldos")
    assert len(c.get("/api/respaldos").json()["copias"]) == 1

    # solo se conservan las últimas 15
    carpeta = tmp_path / "respaldos"
    for d in range(1, 25):
        (carpeta / f"trazabilidad-2026-01-{d:02d}.db").write_bytes(b"x")
    assert respaldo.limpiar() > 0
    assert len(respaldo.listar()) == 15


def test_usuario_queda_en_el_historial(c):
    c.put("/api/pedidos/4500", json={"pedido": "4500", "cliente": "PARIS",
                                     "lineas": [{"sku": "1", "qty": 5}]})
    c.put("/api/entregas/8705700010", json={"entrega": "8705700010", "pedido": "4500",
                                            "lineas": [{"sku": "1", "qty": 5}]},
          headers={"X-Usuario": "Vicente Soto"})
    e = _estado(c)["entregas"][0]
    assert any("Vicente Soto" in x["txt"] for x in e["log"])

    # sin nombre, el historial no inventa uno
    c.put("/api/entregas/8705700011", json={"entrega": "8705700011", "pedido": "4500",
                                            "lineas": [{"sku": "1", "qty": 2}]})
    otra = [x for x in _estado(c)["entregas"] if x["entrega"] == "8705700011"][0]
    assert not any("Vicente" in x["txt"] for x in otra["log"])


def test_sesion_sin_clave_no_molesta(c):
    r = c.get("/api/sesion").json()
    assert r["requiere_clave"] is False and r["abierta"] is True
    assert c.get("/api/estado").status_code == 200        # entra directo, como hasta ahora


def test_sesion_con_clave(monkeypatch):
    from fastapi.testclient import TestClient
    from app.config import settings
    from app.main import app
    object.__setattr__(settings, "clave_acceso", "secreta")
    cli = TestClient(app)
    try:
        assert cli.get("/api/estado").status_code == 401           # sin entrar, no pasa
        assert cli.post("/api/sesion", json={"clave": "mala", "usuario": "Vicente"}).status_code == 401
        assert cli.post("/api/sesion", json={"clave": "secreta", "usuario": ""}).status_code == 422
        r = cli.post("/api/sesion", json={"clave": "secreta", "usuario": "Vicente"})
        assert r.status_code == 200 and r.json()["usuario"] == "Vicente"
        assert cli.get("/api/estado").status_code == 200           # ya con sesión abierta
    finally:
        object.__setattr__(settings, "clave_acceso", "")


def test_plantilla_y_carga_masiva(c, tmp_path):
    from io import BytesIO
    from app.config import settings
    from openpyxl import Workbook, load_workbook
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))

    # la plantilla trae las columnas y una hoja de ayuda
    r = c.get("/api/cubicaje-libre/plantilla")
    assert r.status_code == 200 and "plantilla_cubicaje.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Carga", "Cómo se usa"]
    assert [x.value for x in wb["Carga"][1]] == ["SKU", "Unidades", "Grupo (opcional)", "Sucursal (opcional)"]

    # armar un archivo como lo haría el usuario
    libro = Workbook(); ws = libro.active; ws.title = "Carga"
    ws.append(["SKU", "Unidades", "Camión fijo (opcional)", "Sucursal (opcional)"])
    ws.append(["0900081624", 30, 1, ""])          # con ceros a la izquierda y camión fijo
    ws.append(["900276671", "12,5", "", ""])      # con coma decimal
    ws.append(["999999", 5, "", ""])              # sin medidas
    ws.append(["900276671", 0, "", ""])           # cantidad en cero: se ignora
    ws.append([None, None, None, None])
    ruta = tmp_path / "carga.xlsx"; libro.save(ruta)

    with open(ruta, "rb") as fh:
        d = c.post("/api/cubicaje-libre/importar", files={"file": ("carga.xlsx", fh.read())}).json()
    # el producto sin medidas no se cubica (no se inventan medidas) pero tampoco se pierde: queda en la lista y se alerta
    assert d["importadas"] == 3 and d["sin_medidas_archivo"] == ["999999"]
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("999999", 5)] and d["desconocidos"] == ["999999"]
    skus = {l["sku"]: l for l in d["lineas"]}
    assert set(skus) == {"900081624", "900276671", "999999"}
    assert "camion" not in skus["900081624"]                   # columna de la plantilla antigua: se ignora
    assert skus["900276671"]["qty"] == 12.5
    assert d["unidades"] > 0 and d["camiones"]                 # ya viene cubicado

    # archivo sin nada cubicable
    vacio = Workbook(); vacio.active.append(["SKU", "Unidades"]); vacio.active.append(["999999", 3])
    ruta2 = tmp_path / "vacio.xlsx"; vacio.save(ruta2)
    with open(ruta2, "rb") as fh:
        r = c.post("/api/cubicaje-libre/importar", files={"file": ("vacio.xlsx", fh.read())})
    assert r.status_code == 200 and r.json()["unidades"] == 0          # nada que cubicar, pero se avisa qué falta
    assert [(f["sku"], f["unidades"]) for f in r.json()["faltantes"]] == [("999999", 3)]
    assert c.post("/api/cubicaje-libre/importar",
                  files={"file": ("x.txt", b"nada")}).status_code == 422


def test_vista_de_un_solo_pallet(c, tmp_path):
    """Las cajas quedan apoyadas sobre la tarima y dentro de la huella del pallet."""
    import json as _json
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"
    plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    c.put("/api/clientes/PARIS", json={"pallet": [120, 100, 150]})

    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "900081624", "qty": 12}],
                                            "modo": "SDA STOCK", "caja_master": "SIN CAJA MASTER",
                                            "cliente": "PARIS", "vista": "pallet"}).json()
    assert d["pallet_visto"] == 1 and d["pallets_disponibles"]
    html = (Path(settings.visores_dir) / d["visor"].split("/")[-1]).read_text(encoding="utf-8")
    i = html.index('{"titulo"')
    prof, fin = 0, i
    for j, ch in enumerate(html[i:], start=i):
        prof += (ch == "{") - (ch == "}")
        if prof == 0:
            fin = j
            break
    D = _json.loads(html[i:fin + 1])
    cam = D["camiones"][0]
    pal = cam["pallets"][0]
    assert cam["tipo"].startswith("Pallet") and pal["tar"] > 0
    assert min(x["z"] for x in cam["cajas"]) >= pal["tar"] - 0.01     # apoyadas, no hundidas
    assert all(x["x"] >= -0.01 and x["x"] + x["dx"] <= pal["dl"] + 0.01 and
               x["y"] >= -0.01 and x["y"] + x["dy"] <= pal["dw"] + 0.01 for x in cam["cajas"])

    # se puede recorrer los pallets
    if len(d["pallets_disponibles"]) > 1:
        d2 = c.post("/api/cubicaje-libre", json={**d, "pallet_n": d["pallets_disponibles"][1]}).json()
        assert d2["pallet_visto"] == d["pallets_disponibles"][1]


def test_conexion_sql_desde_la_plataforma(c, monkeypatch):
    from app.integrations import bases
    # sin nada configurado no revela claves
    r = c.get("/api/conexion").json()
    assert r["tiene_clave"] is False

    c.put("/api/conexion", json={"url": "mssql+pyodbc://clws0156/161221_TS_ODS?driver=x",
                                 "usuario": "od_lectura", "clave": "secreta"})
    r = c.get("/api/conexion").json()
    assert r["usuario"] == "od_lectura" and r["tiene_clave"] is True
    assert "secreta" not in str(r)                     # la clave nunca vuelve a la web

    # guardar de nuevo con la clave vacía no la borra
    c.put("/api/conexion", json={"url": "mssql+pyodbc://clws0156/161221_TS_ODS?driver=x",
                                 "usuario": "od_lectura", "clave": ""})
    assert c.get("/api/conexion").json()["tiene_clave"] is True
    assert bases.url_bases().username == "od_lectura"
    assert bases.url_bases().password == "secreta"

    # sin usuario se entra con la cuenta de Windows y no queda clave guardada
    c.put("/api/conexion", json={"url": "mssql+pyodbc://clws0156/161221_TS_ODS?driver=x",
                                 "usuario": "", "clave": ""})
    assert c.get("/api/conexion").json()["tiene_clave"] is False
    assert bases.url_bases().username is None

    monkeypatch.setattr(bases, "motor_bases", lambda: (_ for _ in ()).throw(RuntimeError("login failed")))
    r = c.post("/api/conexion/probar").json()
    assert r["ok"] is False and "login failed" in r["mensaje"]
    bases.usar_conexion({})


def test_el_usuario_del_historial_no_se_mezcla_entre_peticiones():
    """Dos peticiones en paralelo deben quedar registradas cada una con su autor."""
    import threading
    from app import domain
    vistos, barrera = {}, threading.Barrier(2)

    def peticion(nombre):
        domain.usar_usuario(nombre)
        barrera.wait()                      # ambas peticiones "en vuelo" a la vez
        vistos[nombre] = domain.usuario_actual()

    hilos = [threading.Thread(target=peticion, args=(n,)) for n in ("Vicente", "Camila")]
    for h in hilos: h.start()
    for h in hilos: h.join()
    assert vistos == {"Vicente": "Vicente", "Camila": "Camila"}


def test_la_base_aguanta_dos_escrituras_a_la_vez(c):
    """Los análisis corren en hilos de fondo: una escritura simultánea debe esperar,
    no fallar con 'database is locked'."""
    import threading
    import time
    from sqlalchemy import text
    from app.db import engine

    arrancado, soltar = threading.Event(), threading.Event()

    def ocupar_la_base():
        with engine.connect() as cx:
            cx.execute(text("BEGIN IMMEDIATE"))
            cx.execute(text("INSERT OR REPLACE INTO config (clave, valor) VALUES ('t1','a')"))
            arrancado.set()
            soltar.wait(5)
            cx.rollback()

    h = threading.Thread(target=ocupar_la_base)
    h.start()
    arrancado.wait(5)
    threading.Timer(0.4, soltar.set).start()      # el otro hilo suelta a los 0,4 s
    inicio = time.monotonic()
    with engine.begin() as cx:                    # debe esperar y terminar bien
        cx.execute(text("INSERT OR REPLACE INTO config (clave, valor) VALUES ('t2','b')"))
    assert time.monotonic() - inicio >= 0.3, "no llegó a esperar: la prueba no probó nada"
    h.join()


def test_la_pagina_versiona_sus_archivos_para_no_mezclar_cache_vieja(c):
    html = c.get("/").text
    import re
    refs = re.findall(r'(?:src|href)="(/(?:js|css|vendor)/[^"]+)"', html)
    assert refs and all("?v=" in r for r in refs)
    r = c.get("/js/base.js")
    assert r.status_code == 200 and r.headers["cache-control"] == "no-cache"


def test_medidas_de_los_camiones_se_editan_y_llegan_al_cubicaje(c, tmp_path):
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    carga = {"lineas": [{"sku": "900081624", "qty": 30}], "vista": "rampla"}
    d = c.post("/api/cubicaje-libre", json=carga).json()
    assert (d["camiones"][0]["L"], d["camiones"][0]["w"], d["camiones"][0]["h"]) == (1540.0, 245.0, 235.0)

    nuevas = {"rampla": {"largo": 1500, "ancho": 240, "alto": 240}, "camion50": {"largo": 600, "ancho": 240, "alto": 225}}
    r = c.put("/api/ajustes-cubicaje", json={"camiones": nuevas})
    assert r.status_code == 200 and r.json()["camiones"]["rampla"]["alto"] == 240
    d = c.post("/api/cubicaje-libre", json=carga).json()
    assert (d["camiones"][0]["L"], d["camiones"][0]["w"], d["camiones"][0]["h"]) == (1500.0, 240.0, 240.0)
    d50 = c.post("/api/cubicaje-libre", json={**carga, "vista": "camion50"}).json()
    assert d50["camiones"][0]["h"] == 225.0

    # cambiar otro ajuste no pisa las medidas
    c.put("/api/ajustes-cubicaje", json={"celda_cm": 2})
    assert c.get("/api/ajustes-cubicaje").json()["camiones"]["rampla"]["alto"] == 240
    # valores fuera de rango o vacíos se rechazan, con el motivo en mm
    mal = c.put("/api/ajustes-cubicaje", json={"camiones": {**nuevas, "rampla": {"largo": 1500, "ancho": 240, "alto": 0}}})
    assert mal.status_code == 422 and "alto de la Rampla 53" in mal.json()["detail"] and "mm" in mal.json()["detail"]
    assert c.put("/api/ajustes-cubicaje", json={"camiones": {"rampla": {"largo": 1500}}}).status_code == 422
    # volver a lo de fábrica
    c.put("/api/ajustes-cubicaje", json={"camiones": {"rampla": {"largo": 1540, "ancho": 245, "alto": 235},
                                                      "camion50": {"largo": 620, "ancho": 244, "alto": 230}}, "celda_cm": 1})


def test_el_cubicaje_sda_asigna_vehiculos_con_las_medidas_editadas():
    from app.cubicaje import sda
    original = (sda.RAMPLA, sda.CAMION50)
    try:
        sda.usar_camiones(("Rampla", 1500, 240, 240), ("Camion 50", 600, 240, 225))
        v = sda.asignar_vehiculos(31)
        assert (v[0].L, v[0].w, v[0].h) == (1500, 240, 240) and v[1].h == 225
    finally:
        sda.usar_camiones(*original)


def test_plantilla_de_carga_trae_grupos_y_el_importador_los_lee(c, tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    _cargar_medidas_basicas(c, tmp_path)
    r = c.get("/api/cubicaje-libre/plantilla")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Carga"]
    assert [x.value for x in ws[1]] == ["SKU", "Unidades", "Grupo (opcional)", "Sucursal (opcional)"]
    assert len({ws.cell(row=i, column=3).value for i in range(2, ws.max_row + 1)}) >= 2       # el ejemplo ya muestra grupos
    assert "Cómo se usa" in wb.sheetnames
    # lo que trae la plantilla se importa y las tandas llegan como grupos; «Grupo 2» / «G1» también se entienden
    ws["C2"], ws["C3"], ws["C4"] = "Grupo 1", "G1", 2
    buf = BytesIO(); wb.save(buf)
    d = c.post("/api/cubicaje-libre/importar", files={"file": ("p.xlsx", buf.getvalue(), "application/vnd.ms-excel")},
               data={"reemplazar": "si"}).json()
    assert sorted({g["grupo"] for g in d["grupos"]}) == [1, 2] or {g["grupo"] for g in d["grupos"]} >= {1, 2}


def test_cubicaje_conjunto_acepta_clientes_distintos(c, tmp_path, monkeypatch):
    """Los clientes regionales se juntan: pedidos de clientes distintos se cubican juntos con las reglas de uno."""
    _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4410", cliente="PARIS")
    _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4411", cliente="RIPLEY")
    r = c.post("/api/cubicaje-conjunto", json={"pedidos": ["4410", "4411"]})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["conjunto"]["clientes"] == {"4410": "PARIS", "4411": "RIPLEY"}
    assert d["conjunto"]["distintos"] == ["PARIS", "RIPLEY"] and d["cliente"] == "PARIS"
    assert any("reglas distintas" in a for a in d["avisos"])
    r2 = c.post("/api/cubicaje-conjunto", json={"pedidos": ["4410", "4411"], "cliente": "RIPLEY"})
    assert r2.status_code == 200 and r2.json()["cliente"] == "RIPLEY"
    assert r2.json()["conjunto"]["reglas_de"] == "RIPLEY"


def test_cubicaje_conjunto_entregas_por_camion_y_pedido_y_camion_compartido(c, tmp_path, monkeypatch):
    """Multipedido como el Excel: los pedidos se cubican juntos, cada camión lleva cargas de varios pedidos, se crea una
    entrega por (camión, pedido) y las entregas de un mismo camión quedan vinculadas para agruparse."""
    from app.integrations import sap_crear
    _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4400")
    _preparar_cubicaje(c, tmp_path, monkeypatch, pedido="4401")

    assert c.post("/api/cubicaje-conjunto", json={"pedidos": ["4400"]}).status_code == 422          # con uno solo no es conjunto
    assert c.post("/api/cubicaje-conjunto", json={"pedidos": ["4400", "9999"]}).status_code == 422   # uno sin análisis
    r = c.post("/api/cubicaje-conjunto", json={"pedidos": ["4400", "4401"]})
    assert r.status_code == 200, r.text
    d = r.json()
    cid = d["conjunto"]["id"]
    assert d["conjunto"]["pedidos"] == ["4400", "4401"] and cid.startswith("J")
    assert {f["pedido"] for f in d["filas"]} == {"4400", "4401"}
    assert set(d["por_pedido"]) == {"4400", "4401"} and d["por_pedido"]["4400"]["unidades"] == 50
    assert d["unidades"] == 100
    cams_comunes = set(d["por_pedido"]["4400"]["camiones"]) & set(d["por_pedido"]["4401"]["camiones"])
    assert cams_comunes                                          # al menos un camión lleva carga de los dos pedidos
    llamadas = []

    def falso(pedido, puesto, fecha, materiales, fecha_cita="", hora_cita="", ensayo=False, ses=None):
        llamadas.append((pedido, dict(materiales)))
        r = sap_crear.Resultado(ok=True, ensayo=ensayo)
        r.entrega = "" if ensayo else f"87060{len(llamadas):05d}"
        r.mensaje = "ok"
        return r
    monkeypatch.setattr(sap_crear, "crear_entrega", falso)
    t = _esperar_job(c, c.post("/api/sap/crear_entregas", json={"pedidos": ["4400", "4401"], "ensayo": False}).json()["id"])
    assert t["estado"] == "ok", str(t)[:300]
    assert {x["pedido"] for x in t["datos"]["resultados"]} == {"4400", "4401"}
    hechas = [e for e in _estado(c)["entregas"] if e["entrega"].startswith("87060")]
    assert {e["pedido"] for e in hechas} == {"4400", "4401"}
    assert all(e["camion_ref"].startswith(cid + "-") for e in hechas)
    for cam in cams_comunes:                                     # un camión compartido = entregas de los dos pedidos con la misma referencia
        mismas = [e for e in hechas if e["camion_ref"] == f"{cid}-{cam}"]
        assert {e["pedido"] for e in mismas} == {"4400", "4401"}
    # cada pedido solo crea lo suyo: nada de mezclar cantidades de otro pedido
    assert sum(sum(m.values()) for p, m in llamadas if p == "4400") == 50
    assert sum(sum(m.values()) for p, m in llamadas if p == "4401") == 50


def _excel_kits(tmp_path, filas, nombre="kits.xlsx"):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Kit código SAP", "Descripción del kit", "Componente código SAP", "Cantidad por kit"])
    for f in filas:
        ws.append(f)
    ruta = tmp_path / nombre
    wb.save(ruta)
    return ruta.read_bytes()


def test_kits_se_cargan_desde_excel_se_reemplazan_y_se_borran(c, tmp_path):
    datos = _excel_kits(tmp_path, [["KIT-1", "Kit cocina", "900081624", 1], ["KIT-1", "", "900276671", 2],
                                   ["KIT-2", "Otro", "900081624", 1], ["KIT-2", "", "", 1], ["KIT-3", "x", "900081624", 0]])
    r = c.post("/api/kits/importar", files={"file": ("kits.xlsx", datos)})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["nuevos"] == 2 and d["ignoradas"] == 2 and len(d["problemas"]) == 2
    kit1 = next(k for k in d["filas"] if k["sku"] == "KIT-1")
    assert kit1["descripcion"] == "Kit cocina" and kit1["componentes"] == [
        {"sku": "900081624", "cantidad": 1}, {"sku": "900276671", "cantidad": 2}]
    # cargar de nuevo: lo igual no cambia, lo distinto se reemplaza
    r2 = c.post("/api/kits/importar", files={"file": ("kits.xlsx", _excel_kits(tmp_path, [
        ["KIT-1", "Kit cocina", "900081624", 1], ["KIT-1", "", "900276671", 2],
        ["KIT-2", "Otro", "900276671", 3]], "k2.xlsx"))}).json()
    assert r2["sin_cambios"] == 1 and r2["actualizados"] == 1 and r2["nuevos"] == 0
    assert c.get("/api/kits?buscar=KIT-2").json()["filas"][0]["componentes"] == [{"sku": "900276671", "cantidad": 3}]
    assert c.post("/api/kits/importar", files={"file": ("kits.txt", b"x")}).status_code == 422
    assert c.delete("/api/kits/KIT-2").status_code == 200 and c.delete("/api/kits/KIT-2").status_code == 404
    assert c.get("/api/kits").json()["total"] == 1
    pl = c.get("/api/kits/plantilla")
    assert pl.status_code == 200 and pl.content[:2] == b"PK"
    # la plantilla se puede cargar tal cual
    assert c.post("/api/kits/importar", files={"file": ("p.xlsx", pl.content)}).json()["nuevos"] == 2


def test_un_kit_se_cubica_con_las_medidas_de_sus_componentes(c, tmp_path, monkeypatch):
    from app.config import settings
    _cargar_medidas_basicas(c, tmp_path)
    plantilla = tmp_path / "p.html"; plantilla.write_text("X __CUBICAJE_JSON__", encoding="utf-8")
    object.__setattr__(settings, "plantilla_visor", str(plantilla))
    with open(_archivo_medidas(tmp_path, [["1001", "1001 HORNO CHICO", 1, 30, 30, 30, 5, "Y", "N", "N", 500, 40],
                                          ["1002", "1002 ENCIMERA", 1, 40, 20, 20, 5, "Y", "N", "N", 500, 40]]), "rb") as fh:
        c.post("/api/medidas/importar", files={"file": ("Base de Medidas.xlsm", fh.read())})
    c.post("/api/kits/importar", files={"file": ("k.xlsx", _excel_kits(tmp_path, [
        ["KITX", "Kit horno + encimera", "1001", 1], ["KITX", "", "1002", 1]]))})
    d = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "KITX", "qty": 20}], "modo": "SDA STOCK",
                                            "caja_master": "SIN CAJA MASTER", "cliente": "SODIMAC"})
    assert d.status_code == 200, d.text
    d = d.json()
    assert d["desconocidos"] == [] and d["faltantes"] == []
    assert d["unidades"] == 40 and {p["tipo"] for p in d["pallets_detalle"]} == {"Kit"}
    por_sku = {}
    for f in d["filas04"]:
        por_sku[f["sku"]] = por_sku.get(f["sku"], 0) + f["unidades"]
    assert por_sku == {"1001": 20, "1002": 20}
    # si a un componente le faltan medidas, se pide ese componente (no el kit)
    c.post("/api/kits/importar", files={"file": ("k2.xlsx", _excel_kits(tmp_path, [
        ["KITY", "Con faltante", "900081624", 1], ["KITY", "", "900999999", 1]], "k2.xlsx"))})
    d2 = c.post("/api/cubicaje-libre", json={"lineas": [{"sku": "KITY", "qty": 5}], "modo": "SDA STOCK",
                                             "caja_master": "SIN CAJA MASTER"}).json()
    assert [f["sku"] for f in d2["faltantes"]] == ["900999999"] and d2["faltantes"][0]["unidades"] == 5
