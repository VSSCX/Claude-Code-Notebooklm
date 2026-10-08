"""Revisión del backend: estado rápido e idéntico, motor SQL único, lecturas de red sin
bloqueos, reserva atómica de acciones y clave de acceso."""
import os
import tempfile
import threading
import time
from types import SimpleNamespace

import pytest

_tmp = tempfile.mkdtemp()
os.environ.setdefault("DATABASE_URL", f"sqlite:///{_tmp}/test_backend.db")
os.environ.setdefault("VISORES_DIR", f"{_tmp}/visores")

from sqlalchemy import select  # noqa: E402
from sqlalchemy.orm import selectinload  # noqa: E402

from app import domain, usuarios  # noqa: E402
from app.db import SessionLocal, engine  # noqa: E402
from app.integrations import acciones, base_medidas, bases, maestra  # noqa: E402
from app.models import (Archivo, Base, Config, Entrega, Evento, LineaEntrega, LineaPedido,  # noqa: E402
                        PasoEntrega, Pedido, ahora)


@pytest.fixture()
def datos():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as s:
        for i in range(6):
            p = Pedido(pedido=f"40051{i:05d}", cliente="RIPLEY", oc=f"OC{i}", obs="ñandú" if i == 0 else "")
            p.lineas = [LineaPedido(sku=f"9000{j}", descripcion="" if j % 2 else f"Prod {j}", qty=10 + j,
                                    externa=j) for j in range(4)]
            for k in range(2):
                e = Entrega(entrega=f"87{i}{k}", grupo=f"G{i}", vehiculo="Rampla 53")
                e.lineas = [LineaEntrega(sku=f"9000{j}", qty=3 + j) for j in range(3)]
                e.pasos = [PasoEntrega(paso="solicitada", ok=True), PasoEntrega(paso="sap", ok=False)]
                # más eventos que el máximo: solo deben salir los últimos
                e.eventos = [Evento(texto=f"evento {n}", origen="web", at=ahora().replace(microsecond=n * 1000))
                             for n in range(domain.LOG_MAX + 15)]
                p.entregas.append(e)
            s.add(p)
        s.flush()
        s.add(Archivo(pedido_id=s.scalar(select(Pedido.id).limit(1)), tipo="pdf", nombre="a.pdf", archivo="x.pdf"))
        s.add(Archivo(pedido_id=None, tipo="adjunto", nombre="b.pdf", archivo="y.pdf"))
        s.add(Config(clave="k", valor='{"a": 1}'))
        s.commit()
    yield


def _estado_orm(s):
    """La versión anterior de domain.estado (objetos del ORM): referencia para comparar."""
    peds = s.scalars(select(Pedido).options(selectinload(Pedido.lineas))).all()
    ents = s.scalars(select(Entrega).options(
        selectinload(Entrega.lineas), selectinload(Entrega.pasos),
        selectinload(Entrega.eventos), selectinload(Entrega.pedido_ref))).all()
    desc = maestra.descripciones()
    import json
    return {
        "pedidos": [domain.pedido_doc(p, desc) for p in peds],
        "entregas": [domain.entrega_doc(e) for e in ents],
        "config": {c.clave: json.loads(c.valor or "{}") for c in s.scalars(select(Config)).all()},
        "archivos": [domain.archivo_doc(a, p.pedido if p else "")
                     for a, p in s.execute(select(Archivo, Pedido).join(Pedido, isouter=True)
                                           .order_by(Archivo.at.desc())).all()],
        "version": domain.version(s),
    }


def test_estado_igual_al_de_objetos(datos):
    with SessionLocal() as s:
        nuevo, viejo = domain.estado(s), _estado_orm(s)
    assert nuevo["pedidos"] == viejo["pedidos"]
    assert nuevo["config"] == viejo["config"] and nuevo["version"] == viejo["version"]
    assert nuevo["archivos"] == viejo["archivos"]
    assert len(nuevo["entregas"]) == len(viejo["entregas"]) == 12
    for n, v in zip(nuevo["entregas"], viejo["entregas"]):
        assert n == v


def test_estado_solo_trae_el_historial_reciente(datos):
    with SessionLocal() as s:
        e = domain.estado(s)["entregas"][0]
    assert len(e["log"]) == domain.LOG_MAX
    assert e["log"][0]["txt"] == f"evento {domain.LOG_MAX + 14}"     # lo más nuevo primero


def test_estado_vacio():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as s:
        e = domain.estado(s)
    assert e["pedidos"] == [] and e["entregas"] == [] and e["archivos"] == []


# ---------- motor SQL ----------
def test_un_solo_motor_para_las_consultas(monkeypatch):
    monkeypatch.setattr(bases, "_guardada", {"url": "sqlite://"})
    bases._soltar_motor()
    a, b = bases.motor_bases(), bases.motor_bases()
    assert a is b
    bases.usar_conexion({"url": "sqlite:///otro.db"})          # cambiar la conexión rehace el motor
    assert bases.motor_bases() is not a
    bases.usar_conexion({})


# ---------- lecturas de red ----------
class _RutaCaida:
    llamadas = 0

    def __init__(self, *_):
        pass

    def stat(self):
        _RutaCaida.llamadas += 1
        raise OSError("la red no responde")


@pytest.mark.parametrize("modulo, config, funcion", [
    (maestra, "maestra_homologacion", "descripciones"),
    (base_medidas, "base_medidas", "medidas"),
])
def test_red_caida_no_se_reintenta_en_cada_peticion(monkeypatch, modulo, config, funcion):
    monkeypatch.setattr(modulo, "settings", SimpleNamespace(**{config: r"\\servidor\x.xlsx"}))
    monkeypatch.setattr(modulo, "Path", _RutaCaida)
    monkeypatch.setitem(modulo._estado, "revisado", 0.0)
    monkeypatch.setitem(modulo._estado, "mtime", None)
    _RutaCaida.llamadas = 0
    for _ in range(5):
        getattr(modulo, funcion)()
    assert _RutaCaida.llamadas == 1


def test_red_lenta_no_bloquea_a_las_demas_peticiones(monkeypatch):
    monkeypatch.setattr(maestra, "settings", SimpleNamespace(maestra_homologacion="x.xlsx"))
    monkeypatch.setitem(maestra._estado, "revisado", 0.0)
    monkeypatch.setitem(maestra._estado, "datos", {"1": "algo"})
    suelta = threading.Event()

    class Lenta:
        def __init__(self, *_):
            pass

        def stat(self):
            suelta.wait(5)
            raise OSError("tarde")

    monkeypatch.setattr(maestra, "Path", Lenta)
    t = threading.Thread(target=maestra.descripciones)
    t.start()
    time.sleep(0.2)
    inicio = time.time()
    assert maestra.descripciones() == {"1": "algo"}              # no espera la red
    assert time.time() - inicio < 0.5
    suelta.set()
    t.join()


# ---------- acciones ----------
def test_una_sola_accion_a_la_vez_y_lista_acotada():
    soltar = threading.Event()
    empezo = threading.Event()

    def larga(_avance):
        empezo.set()
        soltar.wait(5)
        return "ok"

    primero = acciones.lanzar_python("t1", "Larga", [], larga)
    empezo.wait(2)
    with pytest.raises(RuntimeError, match="Ya hay una acción"):
        acciones.lanzar_python("t2", "Otra", [], lambda _a: None)
    soltar.set()
    for _ in range(100):
        if acciones.trabajo(primero["id"])["estado"] != "en_curso":
            break
        time.sleep(0.02)
    for _ in range(100):                                           # la reserva se libera al terminar
        if not acciones._ocupado[0]:
            break
        time.sleep(0.02)
    for i in range(acciones.MAX_TRABAJOS + 10):
        acciones.lanzar_python("t", f"n{i}", [], lambda _a: None)
        for _ in range(100):
            if not acciones._ocupado[0]:
                break
            time.sleep(0.01)
    assert len(acciones._trabajos) <= acciones.MAX_TRABAJOS


def test_dos_acciones_simultaneas_solo_deja_pasar_una():
    resultados, barrera = [], threading.Barrier(8)

    def intentar():
        barrera.wait()
        try:
            acciones.lanzar_python("p", "P", [], lambda _a: time.sleep(0.3))
            resultados.append("ok")
        except RuntimeError:
            resultados.append("ocupado")

    hilos = [threading.Thread(target=intentar) for _ in range(8)]
    [h.start() for h in hilos]
    [h.join() for h in hilos]
    assert resultados.count("ok") == 1 and resultados.count("ocupado") == 7
    time.sleep(0.5)


# ---------- clave de acceso ----------
def test_con_clave_todo_queda_detras_de_la_sesion(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.main import app
    monkeypatch.setattr(usuarios, "settings", SimpleNamespace(clave_acceso="secreta"))
    monkeypatch.setattr(usuarios, "_sesiones", set())
    c = TestClient(app)
    assert c.get("/api/estado").status_code == 401
    assert c.get("/archivos/1").status_code == 401                 # antes se podía abrir sin sesión
    assert c.get("/visor/visor_vivo.html").status_code == 401
    assert c.get("/").status_code == 200                           # la página carga para poder pedir la clave
    assert c.post("/api/sesion", json={"clave": "mala", "usuario": "Ana"}).status_code == 401
    assert c.post("/api/sesion", json={"clave": "secreta", "usuario": "Ana"}).status_code == 200
    assert c.get("/api/estado").status_code == 200
    assert c.get("/archivos/999").status_code == 404                # con sesión llega a la ruta


def test_sin_clave_nada_cambia(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    assert c.get("/api/estado").status_code == 200


# ---------- SAP: errores claros ----------
def test_accion_en_curso_responde_409_en_todos_los_endpoints_sap(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        domain.guardar_config(s, "cubicaje:4005100000", {"filas": [{"camion": 1, "sku": "9000", "unidades": 5}]})
        s.commit()
    llamadas = [
        ("/api/sap/crear_entregas", {"pedido": "4005100000", "puesto": "PN01", "ensayo": True}),
        ("/api/sap/crear_grupo", {"entregas": ["8700"], "ensayo": True}),
        ("/api/sap/fecha_grupo", {"grupo": "123", "fecha": "01.10.2026", "ensayo": True}),
    ]
    acciones._ocupado[0] = True
    try:
        for ruta, cuerpo in llamadas:
            r = c.post(ruta, json=cuerpo)
            assert r.status_code == 409, (ruta, r.status_code, r.text)
            assert "Ya hay una acción" in r.json()["detail"]
    finally:
        acciones._ocupado[0] = False


def test_ajuste_de_carga_rechaza_nan_e_infinito(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        domain.guardar_config(s, "analisis:4005100000", {"posiciones": [{"sku": "9000"}], "ajustes": {}})
        s.commit()
    for malo in ("nan", "inf", "-1"):
        r = c.put("/api/analisis/4005100000/carga", json={"sku": "9000", "carga": malo})
        assert r.status_code == 422, (malo, r.text)


def test_entrega_creada_en_sap_se_informa_aunque_falle_el_registro(monkeypatch, datos):
    """Si no se puede guardar en la plataforma, no se aborta ni se pierde el número de la entrega."""
    import json
    from fastapi.testclient import TestClient
    from app.integrations import sap_crear
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        domain.guardar_config(s, "cubicaje:4009999999", {"filas": [
            {"camion": 1, "sku": "9000", "unidades": 5}, {"camion": 2, "sku": "9000", "unidades": 7}]})
        s.commit()
    creadas = iter(["8800001", "8800002"])
    monkeypatch.setattr(sap_crear, "crear_entrega", lambda *a, **k: SimpleNamespace(
        ok=True, entrega=next(creadas), mensaje="", pasos=[], borradas=[], ajustadas=[], incidencias=[]))
    # el pedido no existe en la plataforma: guardar_entrega falla, como si la base no respondiera
    r = c.post("/api/sap/crear_entregas", json={"pedido": "4009999999", "puesto": "PN01", "ensayo": False,
                                                "confirmar": "4009999999"})
    assert r.status_code == 202
    for _ in range(100):
        t = c.get(f"/api/acciones/trabajos/{r.json()['id']}").json()
        if t["estado"] != "en_curso":
            break
        time.sleep(0.05)
    assert t["estado"] == "ok", t
    res = t["datos"]["resultados"]
    assert [x["entrega"] for x in res] == ["8800001", "8800002"]          # siguió con el segundo camión
    assert all("no se pudo registrar" in x["incidencias"][0] for x in res)
    time.sleep(0.2)


def test_consulta_sql_fallida_no_se_repite_en_cada_peticion(monkeypatch):
    bases.limpiar_cache()
    intentos = []

    def cae():
        intentos.append(1)
        raise ConnectionError("servidor sin respuesta")

    for _ in range(6):
        with pytest.raises(ConnectionError):
            bases._cache("prueba", cae)
    assert len(intentos) == 1                  # antes cada petición esperaba el timeout completo
    bases.limpiar_cache()                      # cambiar la conexión o refrescar lo permite de nuevo
    with pytest.raises(ConnectionError):
        bases._cache("prueba", cae)
    assert len(intentos) == 2
    assert bases._cache("ok", lambda: 5) == 5
    bases.limpiar_cache()


# ---------- flujo del pedido: análisis → ajuste → cubicaje ----------
def _analisis_de_prueba(numero="4005100000"):
    from app.analisis import calcular
    pos = [{"sku": "9000", "qty_entrega": 100, "qty_pendiente": 100},       # el plan deja 40: limitado
           {"sku": "9001", "qty_entrega": 20, "qty_pendiente": 20}]         # el plan alcanza: completo
    plan = {"9000": {"plan": 100, "vendido": 60}, "9001": {"plan": 500, "vendido": 0}}
    med = {"9000": {"desc": "9000 PRODUCTO A", "max_camion": 200}, "9001": {"desc": "9001 PRODUCTO B", "max_camion": 100}}
    res = calcular(pos, plan, {}, med, {})
    return {"pedido": numero, "cliente": "PARIS", "grupo_sop": "PARIS", "puesto": "PN01", "fecha": "2026-10-01",
            "posiciones": pos, "en_entrega": {}, "plan": plan, "medidas": med, "disponibilidad": {},
            "stock": {}, "ajustes": {}, "resultado": res, "generado": "2026-10-01"}


def test_el_analisis_limita_la_carga_al_saldo_del_plan():
    filas = {f["sku"]: f for f in _analisis_de_prueba()["resultado"]["filas"]}
    assert filas["9000"]["saldo"] == 40 and filas["9000"]["carga"] == 40       # 100 pedidas, el plan deja 40
    assert filas["9000"]["alerta"] == "Limitado SOP"
    assert filas["9001"]["carga"] == 20 and filas["9001"]["alerta"] == "Completo"


def test_ajuste_por_lote_registra_quien_autoriza_exceder_el_plan(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        doc = _analisis_de_prueba()
        domain.guardar_config(s, "analisis:4005100000", doc)
        domain.anotar_flujo(s, "4005100000", analisis=domain.resumen_analisis(doc))
        s.commit()
    r = c.put("/api/analisis/4005100000/carga", json={"ajustes": {"9000": 100}}, headers={"x-usuario": "Ana"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["ajustes"] == {"9000": 100.0}
    assert d["autorizaciones"]["9000"]["por"] == "Ana" and d["autorizaciones"]["9000"]["calculada"] == 40
    fila = next(f for f in d["resultado"]["filas"] if f["sku"] == "9000")
    assert fila["carga"] == 100 and fila["ajustada"]
    # bajar por debajo de lo calculado no es exceder: no hay autorización
    d = c.put("/api/analisis/4005100000/carga", json={"ajustes": {"9000": 30}}).json()
    assert "9000" not in d["autorizaciones"]
    # volver al valor calculado
    d = c.put("/api/analisis/4005100000/carga", json={"ajustes": {"9000": None}}).json()
    assert d["ajustes"] == {} and d["autorizaciones"] == {}
    # un SKU que no es del pedido se rechaza
    assert c.put("/api/analisis/4005100000/carga", json={"ajustes": {"7777": 5}}).status_code == 422
    # la forma de un solo SKU sigue funcionando
    assert c.put("/api/analisis/4005100000/carga", json={"sku": "9001", "carga": 15}).status_code == 200


def test_estado_lleva_el_resumen_del_flujo_y_no_los_documentos_pesados(datos):
    with SessionLocal() as s:
        doc = _analisis_de_prueba()
        domain.guardar_config(s, "analisis:4005100000", doc)
        domain.guardar_config(s, "cubicaje:4005100000", {"camiones": [{}, {}], "modo": "MDA", "generado": "x", "filas": []})
        domain.anotar_flujo(s, "4005100000", analisis=domain.resumen_analisis(doc))
        s.commit()
        antes = domain.version(s)
        e = domain.estado(s)
    assert "analisis:4005100000" not in e["config"] and "cubicaje:4005100000" not in e["config"]
    fl = e["config"]["flujo:4005100000"]["analisis"]
    assert fl["limitadas"] == 1 and fl["filas"] == 2 and fl["pedida"] == 120 and fl["carga"] == 60
    with SessionLocal() as s:                          # cambiar un documento pesado sin resumen no mueve la versión
        domain.guardar_config(s, "analisis:4005100000", {**doc, "generado": "otro"})
        s.commit()
        assert domain.version(s) == antes


def test_los_pedidos_analizados_antes_reciben_su_resumen(datos):
    with SessionLocal() as s:
        domain.guardar_config(s, "analisis:4005100000", _analisis_de_prueba())
        domain.guardar_config(s, "cubicaje:4005100000", {"camiones": [{}], "modo": "SDA STOCK", "generado": "g"})
        s.commit()
        assert domain.rellenar_flujo(s) == 2
        s.commit()
        assert domain.rellenar_flujo(s) == 0           # una segunda vez no repite
        fl = domain.estado(s)["config"]["flujo:4005100000"]
    assert fl["analisis"]["limitadas"] == 1 and fl["cubicaje"] == {"generado": "g", "camiones": 1, "modo": "SDA STOCK"}


def test_el_cliente_editado_en_configuracion_manda_sobre_el_de_fabrica(datos):
    from app.routers.comun import _sop_de
    with SessionLocal() as s:
        assert _sop_de("FALABELLA", s)[1] == "237141"                      # de fábrica
        from app.integrations import clientes as cli
        cli.guardar(s, {"nombre": "FALABELLA", "codigo": "999111"})
        cli.guardar(s, {"nombre": "CLIENTE NUEVO", "grupo_sop": "REGION 2", "codigo": "555", "pallet": [120, 100, 150]})
        s.commit()
        assert _sop_de("FALABELLA", s) == ("FALABELLA", "999111")
        assert _sop_de("CLIENTE NUEVO", s) == ("REGION 2", "555")
        assert _sop_de("LA POLAR", s)[0] == "ABC"                          # las reglas del Excel se mantienen


def _esperar(c, jid):
    for _ in range(200):
        t = c.get(f"/api/acciones/trabajos/{jid}").json()
        if t["estado"] != "en_curso":
            return t
        time.sleep(0.05)
    raise AssertionError("el trabajo no terminó")


def test_reanalizar_con_otro_cliente_trae_su_plan_su_codigo_y_rehace_el_cubicaje(monkeypatch, datos):
    """Cambiar el cliente de un pedido: el análisis usa el grupo SOP y el código del cliente NUEVO,
    el pedido queda con ese cliente y el cubicaje se rehace con la carga nueva."""
    from fastapi.testclient import TestClient
    from app.integrations import base_medidas, sap
    from app.main import app
    from app.models import Medida
    c = TestClient(app)
    with SessionLocal() as s:
        for sku in ("9000", "9001"):
            s.add(Medida(sku=sku, descripcion=f"{sku} PRODUCTO {sku}", piezas=1, largo=60, ancho=50, alto=40,
                         peso=10, apilar="Y", inclinar="N", rotar="N", max_camion=300, max_pallet=20))
        p = s.scalar(select(Pedido).where(Pedido.pedido == "4005100000"))
        p.cliente = "FALABELLA"                                        # lo que quedó tras "Editar pedido"
        domain.guardar_config(s, "analisis:4005100000", _analisis_de_prueba())            # análisis viejo, de PARIS
        domain.guardar_config(s, "cubicaje:4005100000", {"modo": "MDA", "camiones": [], "generado": "x", "filas": []})
        s.commit()
    llamadas = {}
    pos = [SimpleNamespace(sku="9000", qty_entrega=100, qty_pendiente=100),
           SimpleNamespace(sku="9001", qty_entrega=20, qty_pendiente=20)]
    monkeypatch.setattr(sap, "leer_pedido", lambda *a: SimpleNamespace(posiciones=pos, aviso=""))

    def zsd(codigo, skus, carpeta, nombre, **kw):
        llamadas["codigo"] = codigo
        return [{"Nombre Codigo de Material": "PRODUCTO 9000", "Qty. En Entrega": 10}]
    monkeypatch.setattr(sap, "zsd001_03", zsd)
    monkeypatch.setattr(sap, "mmbe", lambda skus, cb: {})
    monkeypatch.setattr(bases, "plan_sop", lambda grupo: llamadas.setdefault("grupo", grupo) and
                        {"9000": {"plan": 300, "vendido": 100, "pdte_mes": 0, "tipo": "CONSENSO"},
                         "9001": {"plan": 500, "vendido": 0, "pdte_mes": 0, "tipo": "CONSENSO"}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(base_medidas, "medidas", lambda: {})
    r = c.post("/api/analisis/4005100000", json={"puesto": "PN01", "cliente": "FALABELLA", "fecha": "2026-10-01"})
    assert r.status_code == 202, r.text
    t = _esperar(c, r.json()["id"])
    assert t["estado"] == "ok", t
    assert llamadas["codigo"] == "237141"                              # el código de FALABELLA, no el de PARIS
    assert llamadas["grupo"] == "FALABELLA"
    assert t["datos"]["limitadas"] == 0 and t["datos"]["recubicado"] is True
    doc = c.get("/api/analisis/4005100000").json()
    assert doc["cliente"] == "FALABELLA" and doc["grupo_sop"] == "FALABELLA"
    nueva = {f["sku"]: f for f in doc["resultado"]["filas"]}
    assert nueva["9000"]["saldo"] == 190 and nueva["9000"]["carga"] == 100       # 300 − 100 − 10 en entrega
    est = c.get("/api/estado").json()
    assert next(x for x in est["pedidos"] if x["pedido"] == "4005100000")["cliente"] == "FALABELLA"
    assert est["config"]["flujo:4005100000"]["analisis"]["cliente"] == "FALABELLA"
    assert est["config"]["flujo:4005100000"]["cubicaje"]["camiones"] >= 1        # el cubicaje se rehizo
    time.sleep(0.2)


def test_cliente_sin_codigo_no_llama_a_sap_con_uno_inventado(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.integrations import sap
    from app.main import app
    c = TestClient(app)
    monkeypatch.setattr(sap, "leer_pedido", lambda *a: SimpleNamespace(
        posiciones=[SimpleNamespace(sku="9000", qty_entrega=1, qty_pendiente=1)], aviso=""))
    r = c.post("/api/analisis/4005100000", json={"puesto": "PN01", "cliente": "CLIENTE DESCONOCIDO"})
    t = _esperar(c, r.json()["id"])
    assert t["estado"] == "error" and "código de solicitante" in t["error"]


# ---------- cruce de la Qty en entrega por nombre de material ----------
# Filas reales de la hoja "Entregas" del Excel (reporte ZSD001_03) y descripciones de la Base de Medidas.
ZSD_REAL = [
    {"Código de Material": "240098445", "Nombre Codigo de Material": "LAVADORA MADEMSA MDWMT16W", "Qty. En Entrega": 11},
    {"Código de Material": "240086033", "Nombre Codigo de Material": "LAVADORA MADEMSA 9,5 BZG", "Qty. En Entrega": 30},
    {"Código de Material": "240094099", "Nombre Codigo de Material": "COCINA FM4ES", "Qty. En Entrega": 0},
]


def test_el_cruce_por_nombre_resuelve_los_casos_que_el_excel_dejaba_a_mano():
    from app.analisis import calcular, en_entrega_por_modelo
    pos = [{"sku": s, "qty_entrega": 100, "qty_pendiente": 100} for s in ("900081624", "240086033", "240094099", "926565426")]
    med = {"900081624": {"desc": "900081624 MDWMT16W", "max_camion": 100},                 # el reporte trae "LAVADORA MADEMSA MDWMT16W"
           "240086033": {"desc": "240086033 LAVADORA MADEMSA 9 5 BZG", "max_camion": 100},  # el reporte trae "9,5"
           "240094099": {"desc": "240094099 COCINA FM4ES", "max_camion": 100},             # igual
           "926565426": {"desc": "926565426 COCINA F5EPB", "max_camion": 100}}
    plan = {s: {"plan": 1000, "vendido": 0} for s in med}
    res = calcular(pos, plan, en_entrega_por_modelo(ZSD_REAL), med, {})
    ent = {f["sku"]: f["en_entrega"] for f in res["filas"]}
    assert ent == {"900081624": 11, "240086033": 30, "240094099": 0, "926565426": 0}
    assert next(f for f in res["filas"] if f["sku"] == "900081624")["cruce_zsd"] == "LAVADORA MADEMSA MDWMT16W"
    assert res["zsd_sin_cruce"] == []
    assert {f["sku"]: f["saldo"] for f in res["filas"]}["900081624"] == 989          # el saldo ya descuenta lo en entrega


def test_el_cruce_no_inventa_coincidencias_dudosas():
    from app.analisis import cruzar_en_entrega
    # dos nombres del reporte terminan igual: es ambiguo, no se reparte nada
    q, usado, sobran = cruzar_en_entrega({"A": "PRO 7D BZG"}, {"SECADORA MADEMSA PRO 7D BZG": 5, "SECADORA OTRA PRO 7D BZG": 9})
    assert q == {} and sorted(sobran) == ["SECADORA MADEMSA PRO 7D BZG", "SECADORA OTRA PRO 7D BZG"]
    # el nombre exacto de un SKU no se le quita con el cruce por final de otro SKU
    q, usado, sobran = cruzar_en_entrega({"A": "COCINA FM4ES", "B": "FM4ES"}, {"COCINA FM4ES": 7})
    assert q == {"A": 7} and sobran == []
    # tildes, mayúsculas y signos no importan; un modelo muy corto no se cruza por el final
    q, _, _ = cruzar_en_entrega({"A": "Lavadora Mademsa 9,5 BZG"}, {"LAVADORA MADEMSA 9 5 BZG": 3})
    assert q == {"A": 3}
    assert cruzar_en_entrega({"A": "X1"}, {"COCINA X1": 3})[0] == {}
    # los datos guardados antes con otro formato de nombre siguen sirviendo
    assert cruzar_en_entrega({"A": "COCINA FM4ES"}, {"COCINA  FM4ES": 4.0})[0] == {"A": 4.0}


def test_un_producto_del_reporte_sin_sku_se_avisa():
    from app.analisis import calcular
    res = calcular([{"sku": "1", "qty_entrega": 5, "qty_pendiente": 5}], {"1": {"plan": 10, "vendido": 0}},
                   {"COCINA FM4ES": 3, "NEVERA RARA": 8}, {"1": {"desc": "1 COCINA FM4ES", "max_camion": 10}}, {})
    assert res["zsd_sin_cruce"] == ["NEVERA RARA"]


# ---------- cubicaje: la grilla decide si coincide con el Excel ----------
MEDIDAS_HOJA_03 = [   # Base de Medidas del Excel: sku, descripción, largo, ancho, alto, peso, máx camión
    ("240086033", "240086033 LAVADORA MADEMSA 9 5 BZG", 61.0, 58.0, 105.0, 36.0, 200),
    ("240094099", "240094099 COCINA FM4ES", 72.0, 59.0, 97.5, 36.0, 168),
    ("240094551", "240094551 COCINA FM4LP", 70.0, 79.0, 99.0, 35.0, 132),
    ("240096076", "240096076 SECADORA MADEMSA PRO 7D BZG", 62.5, 55.5, 84.5, 26.4, 192),
    ("900081624", "900081624 MDWMT16W", 70.1, 66.4, 107.8, 47.0, 126),
    ("926565426", "926565426 COCINA F5EPB", 69.0, 79.0, 100.0, 55.0, 132),
]
CARGA_HOJA_03 = {"900081624": 7, "926565426": 25, "240086033": 200, "240094099": 60, "240094551": 50, "240096076": 125}


@pytest.mark.parametrize("celda, esperado", [
    # (ocupación de cada camión, unidades por camión). 2 cm es la grilla del Excel (CELL = 2) y da su resultado.
    (2, ([0.7898, 0.8229, 0.465], [140, 192, 135])),
    (1, ([0.7898, 0.8572, 0.4308], [140, 200, 127])),         # más fino: acomoda distinto, por eso no coincide
])
def test_cubicaje_del_excel_coincide_con_la_grilla_de_2_cm(datos, celda, esperado):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import Medida
    c = TestClient(app)
    with SessionLocal() as s:
        for sku, desc, largo, ancho, alto, peso, mx in MEDIDAS_HOJA_03:
            s.add(Medida(sku=sku, descripcion=desc, piezas=1, largo=largo, ancho=ancho, alto=alto, peso=peso,
                         apilar="Y", inclinar="N", rotar="N", max_camion=mx, max_pallet=1))
        pos = [{"sku": k, "qty_entrega": q, "qty_pendiente": q} for k, q in CARGA_HOJA_03.items()]
        plan = {k: {"plan": 10000, "vendido": 0} for k in CARGA_HOJA_03}
        med = {m[0]: {"desc": m[1], "max_camion": m[6]} for m in MEDIDAS_HOJA_03}
        from app.analisis import calcular
        doc = {"pedido": "4005100000", "cliente": "PARIS", "grupo_sop": "PARIS", "puesto": "PN01", "fecha": "2026-10-01",
               "posiciones": pos, "en_entrega": {}, "plan": plan, "medidas": med, "disponibilidad": {}, "stock": {},
               "ajustes": {}, "resultado": calcular(pos, plan, {}, med, {}), "generado": "2026-10-01"}
        domain.guardar_config(s, "analisis:4005100000", doc)
        s.commit()
    assert c.put("/api/ajustes-cubicaje", json={"celda_cm": celda}).status_code == 200
    r = c.post("/api/cubicaje/4005100000", json={"modo": "MDA", "caja_master": "SIN CAJA MASTER"})
    assert r.status_code == 200, r.text
    d = r.json()
    ocup = {f["camion"]: f["ocup_acum"] for f in d["filas"]}
    unid = {n: sum(f["unidades"] for f in d["filas"] if f["camion"] == n) for n in ocup}
    assert [round(ocup[n], 4) for n in sorted(ocup)] == esperado[0]
    assert [unid[n] for n in sorted(unid)] == esperado[1]
    assert [x["tipo"] for x in d["camiones"]] == ["Rampla 53"] * 3
    c.put("/api/ajustes-cubicaje", json={"celda_cm": 2})


# ---------- carga masiva con productos que no están en la Base de Medidas ----------
def _excel_carga(filas):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Carga"
    ws.append(["SKU", "Unidades", "Sucursal"])
    for f in filas:
        ws.append(list(f))
    b = BytesIO()
    wb.save(b)
    b.seek(0)
    return b


def _medida(s, sku, desc):
    from app.models import Medida
    s.add(Medida(sku=sku, descripcion=desc, piezas=1, largo=60, ancho=50, alto=40, peso=10,
                 apilar="Y", inclinar="N", rotar="N", max_camion=300, max_pallet=20))


def test_carga_masiva_con_productos_sin_medidas_alerta_con_sku_y_unidades(datos):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.integrations import medidas as med_mod
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    archivo = _excel_carga([("9000", 10, ""), ("7777", 40, ""), ("7777", 5, ""), ("8888", 3, "")])
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("carga.xlsx", archivo)})
    assert r.status_code == 200, r.text
    d = r.json()
    # los que faltan quedan en la carga (no se pierden) y se avisan con sus unidades, los de más unidades primero
    assert {l["sku"] for l in d["lineas"]} == {"9000", "7777", "8888"}
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("7777", 45), ("8888", 3)]
    assert d["desconocidos"] == ["7777", "8888"]
    assert d["unidades"] == 10                                       # solo lo que tiene medidas va al camión
    # la plantilla trae esos SKU en el formato de la Base de Medidas
    x = c.get("/api/cubicaje-libre/faltantes.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(BytesIO(x.content)).worksheets[0]
    assert ws.title == "Base para carga"
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][0] == "Grupo" and filas[0][10] == "Máx Camión" and filas[0][12] == "Unidades en la carga"
    assert [(f[0], f[12]) for f in filas[1:]] == [("7777", 45), ("8888", 3)]
    # sin completar las medidas la importación los ignora (informando), y completándolas desaparece la alerta
    with SessionLocal() as s:
        assert med_mod.importar(s, med_mod.leer_archivo(BytesIO(x.content))).ignorados == 2
    ws["D2"], ws["E2"], ws["F2"], ws["G2"] = 30, 20, 10, 5
    ws["D3"], ws["E3"], ws["F3"], ws["G3"] = 40, 30, 20, 8
    b = BytesIO()
    ws.parent.save(b)
    b.seek(0)
    assert c.post("/api/medidas/importar", files={"file": ("m.xlsx", b)}).status_code == 200
    d = c.post("/api/cubicaje-libre", json={"lineas": d["lineas"], "modo": "MDA", "vista": "rampla"}).json()
    assert d["faltantes"] == [] and d["desconocidos"] == [] and d["unidades"] == 58
    assert c.get("/api/cubicaje-libre/faltantes.xlsx").status_code == 404      # ya no falta ninguno


def test_carga_masiva_solo_con_productos_sin_medidas_no_se_pierde(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("c.xlsx", _excel_carga([("7777", 12, "")]))})
    assert r.status_code == 200, r.text
    assert [(f["sku"], f["unidades"]) for f in r.json()["faltantes"]] == [("7777", 12)]


def test_reparto_por_sucursal_con_producto_sin_medidas_no_rompe(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("c.xlsx", _excel_carga(
        [("9000", 10, "SUC A"), ("7777", 4, "SUC A"), ("9000", 6, "SUC B")]))})
    assert r.status_code == 200, r.text
    d = r.json()
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("7777", 4)]
    assert d["unidades"] == 16


def test_pedido_con_productos_sin_medidas_tambien_alerta(datos):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        doc = _analisis_de_prueba()            # SKU 9000 y 9001: solo el primero tiene medidas
        domain.guardar_config(s, "analisis:4005100000", doc)
        s.commit()
    r = c.post("/api/cubicaje/4005100000", json={"modo": "MDA", "caja_master": "SIN CAJA MASTER"})
    assert r.status_code == 200, r.text
    assert [(f["sku"], f["unidades"]) for f in r.json()["faltantes"]] == [("9001", 20)]
    x = c.get("/api/cubicaje/4005100000/faltantes.xlsx")
    assert [f[0] for f in list(load_workbook(BytesIO(x.content)).worksheets[0].iter_rows(values_only=True))[1:]] == ["9001"]


# ---------- carga masiva con productos que no están en la Base de Medidas ----------
def _excel_carga(filas):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Carga"
    ws.append(["SKU", "Unidades", "Sucursal"])
    for f in filas:
        ws.append(list(f))
    b = BytesIO()
    wb.save(b)
    b.seek(0)
    return b


def _medida(s, sku, desc):
    from app.models import Medida
    s.add(Medida(sku=sku, descripcion=desc, piezas=1, largo=60, ancho=50, alto=40, peso=10,
                 apilar="Y", inclinar="N", rotar="N", max_camion=300, max_pallet=20))


def test_los_productos_sin_medidas_no_se_cubican_y_se_alertan_con_sus_unidades(datos):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.integrations import medidas as med_mod
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    archivo = _excel_carga([("9000", 10, ""), ("7777", 40, ""), ("7777", 5, ""), ("8888", 3, "")])
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("carga.xlsx", archivo)})
    assert r.status_code == 200, r.text
    d = r.json()
    # no se inventa nada: solo lo que tiene medidas va al camión; el resto se avisa con sus unidades (más grandes primero)
    assert d["unidades"] == 10
    assert {l["sku"] for l in d["lineas"]} == {"9000", "7777", "8888"}      # siguen en la lista para poder verlos
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("7777", 45), ("8888", 3)]
    assert d["desconocidos"] == ["7777", "8888"]
    assert {f["sku"] for f in d["filas"]} == {"9000"}
    # la plantilla trae esos SKU en el formato de la Base de Medidas
    x = c.get("/api/cubicaje-libre/faltantes.xlsx")
    assert x.status_code == 200 and "spreadsheetml" in x.headers["content-type"]
    ws = load_workbook(BytesIO(x.content)).worksheets[0]
    assert ws.title == "Base para carga"
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][0] == "Grupo" and filas[0][10] == "Máx Camión" and filas[0][12] == "Unidades en la carga"
    assert [(f[0], f[12]) for f in filas[1:]] == [("7777", 45), ("8888", 3)]
    # sin completar las medidas la importación los ignora; completándolas, la alerta desaparece
    with SessionLocal() as s:
        assert med_mod.importar(s, med_mod.leer_archivo(BytesIO(x.content))).ignorados == 2
    ws["D2"], ws["E2"], ws["F2"], ws["G2"] = 30, 20, 10, 5
    ws["D3"], ws["E3"], ws["F3"], ws["G3"] = 40, 30, 20, 8
    b = BytesIO()
    ws.parent.save(b)
    b.seek(0)
    assert c.post("/api/medidas/importar", files={"file": ("m.xlsx", b)}).status_code == 200
    d = c.post("/api/cubicaje-libre", json={"lineas": d["lineas"], "modo": "MDA", "vista": "rampla"}).json()
    assert d["faltantes"] == [] and d["desconocidos"] == [] and d["unidades"] == 58
    assert c.get("/api/cubicaje-libre/faltantes.xlsx").status_code == 404


def test_carga_masiva_solo_con_productos_sin_medidas_no_se_pierde(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("c.xlsx", _excel_carga([("7777", 12, "")]))})
    assert r.status_code == 200, r.text
    d = r.json()
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("7777", 12)] and d["unidades"] == 0


def test_reparto_por_sucursal_con_producto_sin_medidas_no_rompe(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("c.xlsx", _excel_carga(
        [("9000", 10, "SUC A"), ("7777", 4, "SUC A"), ("9000", 6, "SUC B")]))})
    assert r.status_code == 200, r.text
    d = r.json()
    assert [(f["sku"], f["unidades"]) for f in d["faltantes"]] == [("7777", 4)]
    assert d["unidades"] == 16


def test_pedido_con_productos_sin_medidas_tambien_alerta(datos):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        domain.guardar_config(s, "analisis:4005100000", _analisis_de_prueba())     # 9000 y 9001: solo el primero tiene medidas
        s.commit()
    r = c.post("/api/cubicaje/4005100000", json={"modo": "MDA", "caja_master": "SIN CAJA MASTER"})
    assert r.status_code == 200, r.text
    assert [(f["sku"], f["unidades"]) for f in r.json()["faltantes"]] == [("9001", 20)]
    x = c.get("/api/cubicaje/4005100000/faltantes.xlsx")
    assert [f[0] for f in list(load_workbook(BytesIO(x.content)).worksheets[0].iter_rows(values_only=True))[1:]] == ["9001"]


# ---------- cruce pedido ↔ orden de compra ----------
class _Sql:
    """Base de 'Pedidos Ingresados' simulada: registra la consulta y sus parámetros."""
    def __init__(self, filas):
        self.filas, self.llamadas = filas, []

    def __call__(self, sql, params=()):
        self.llamadas.append((" ".join(sql.split()), params))
        return self.filas


FILAS_SQL = [   # (fechaCreacion, pedidoVenta, ordenCompra, fechaVencimiento), como las devuelve SQL Server
    ("2026-09-22", "4005175955", "", "2026-10-04"),                  # una posición sin OC: no sirve
    ("2026-09-22", "0004005175955", " 4500123 ", "2026-10-04"),     # con ceros delante y espacios
    ("2026-09-22", "4005175955", "4500123", "2026-10-04"),           # la misma OC en otra posición
    ("2026-09-20", "4005170000", "4500-999", "2026-10-01"),
]


def test_la_oc_de_un_pedido_se_consulta_dirigida_y_normalizada(monkeypatch):
    sql = _Sql(FILAS_SQL)
    monkeypatch.setattr(bases, "_filas", sql)
    assert bases.oc_de("004005175955") == "4500123"                  # ceros a la izquierda y espacios
    consulta, params = sql.llamadas[0]
    assert "pedidoVenta IN (?, ?)" in consulta and "WHERE" in consulta
    assert "YEAR(" not in consulta and "RETAIL" not in consulta       # ya no se baja el año completo ni solo RETAIL
    assert params == ("4005175955", "4005175955")
    assert bases.oc_de("4005179999") == ""                           # pedido que aún no aparece
    d = bases.datos_pedido("4005175955")
    assert d == {"oc": "4500123", "fecha": "2026-09-22", "vence": "2026-10-04"}


def test_sin_conexion_la_oc_queda_vacia_y_no_falla(monkeypatch):
    def cae(*a, **k):
        raise RuntimeError("Falta la conexión a SQL Server")
    monkeypatch.setattr(bases, "_filas", cae)
    assert bases.oc_de("4005175955") == ""
    assert bases.datos_pedido("") == {"oc": "", "fecha": "", "vence": ""}


def test_un_pedido_con_dos_oc_toma_la_que_mas_se_repite(monkeypatch):
    monkeypatch.setattr(bases, "_filas", _Sql([
        ("2026-09-22", "4005175955", "AAA1", ""), ("2026-09-22", "4005175955", "BBB2", ""),
        ("2026-09-22", "4005175955", "bbb-2", "")]))              # BBB2 y bbb-2 son la misma OC
    assert bases.oc_de("4005175955").upper().replace("-", "") == "BBB2"


def test_buscar_por_oc_en_varios_pedidos_devuelve_todos_para_elegir(monkeypatch):
    sql = _Sql([("2026-09-22", "4005175955", "4500123", "2026-10-04"),
                ("2026-09-25", "4005176001", "4500123", "2026-10-08"),
                ("2026-09-25", "4005176001", "4500123", "2026-10-08")])      # duplicado: una sola vez
    monkeypatch.setattr(bases, "_filas", sql)
    r = bases.buscar_pedido("4500-123")                               # con otro formato de la misma OC
    assert [x["pedido"] for x in r] == ["4005175955", "4005176001"]
    assert r[0]["fecha"] == "2026-09-22" and r[0]["vence"] == "2026-10-04"
    assert len(sql.llamadas) == 1                                    # lo exacto alcanza: sin búsqueda por contenido


def test_buscar_por_numero_de_pedido_y_por_texto_de_la_oc(monkeypatch):
    sql = _Sql([("2026-09-22", "4005175955", "4500123", "")])
    monkeypatch.setattr(bases, "_filas", sql)
    assert [x["oc"] for x in bases.buscar_pedido("4005175955")] == ["4500123"]
    sql.filas = []                                                    # sin exactos: se busca la OC que contenga el texto
    assert bases.buscar_pedido("45001") == []
    assert "ordenCompra LIKE ?" in sql.llamadas[-1][0] and sql.llamadas[-1][1] == ("%45001%",)
    n = len(sql.llamadas)
    assert bases.buscar_pedido("45") == [] and len(sql.llamadas) == n + 1      # muy corto: no se busca por contenido
    assert bases.buscar_pedido("   ") == []


def test_completar_oc_no_pisa_una_escrita_a_mano():
    p = Pedido(pedido="1", oc="")
    assert domain.completar_oc(p, " 4500123 ") and p.oc == "4500123"
    p = Pedido(pedido="2", oc="ESCRITA-A-MANO")
    assert not domain.completar_oc(p, "4500123") and p.oc == "ESCRITA-A-MANO"
    p = Pedido(pedido="3", oc="")
    assert not domain.completar_oc(p, "") and p.oc == ""


def test_la_oc_tambien_sale_del_reporte_zsd001_03():
    from app.analisis import oc_de_zsd
    zsd = [{"Documento de Ventas": "4005171502", "Orden de compra": "527512", "Nombre Codigo de Material": "X"},
           {"Documento de Ventas": 4005175956, "Orden de compra": "529853"},
           {"Documento de Ventas": "4005175956", "Orden de compra": "529853"},
           {"Documento de Ventas": "4005175956", "Orden de compra": ""}]
    assert oc_de_zsd(zsd, "4005175956") == "529853"
    assert oc_de_zsd(zsd, "004005171502") == "527512"
    assert oc_de_zsd(zsd, "4009999999") == ""
    assert oc_de_zsd([{"Doc. Ventas": "x"}], "4005175956") == ""


def _correr_analisis(c, monkeypatch, oc_sql, zsd_filas, numero="4005100000", cliente="PARIS"):
    from app.integrations import base_medidas, sap
    monkeypatch.setattr(sap, "leer_pedido", lambda *a: SimpleNamespace(
        posiciones=[SimpleNamespace(sku="9000", qty_entrega=10, qty_pendiente=10)], aviso=""))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **k: zsd_filas)
    monkeypatch.setattr(sap, "mmbe", lambda skus, cb: {})
    monkeypatch.setattr(bases, "plan_sop", lambda g: {"9000": {"plan": 100, "vendido": 0, "pdte_mes": 0, "tipo": "CONSENSO"}})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(bases, "oc_de", lambda p: oc_sql)
    monkeypatch.setattr(base_medidas, "medidas", lambda: {"9000": {"desc": "9000 PRODUCTO A", "max_camion": 100}})
    r = c.post(f"/api/analisis/{numero}", json={"puesto": "PN01", "cliente": cliente, "fecha": "2026-10-01"})
    assert r.status_code == 202, r.text
    t = _esperar(c, r.json()["id"])
    assert t["estado"] == "ok", t
    return t


def test_el_analisis_toma_la_oc_de_sql_y_si_no_esta_del_reporte(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        s.get(Pedido, 1).oc = ""
        s.commit()
    zsd = [{"Documento de Ventas": "4005100000", "Orden de compra": "ZSD-77", "Nombre Codigo de Material": "PRODUCTO A", "Qty. En Entrega": 0}]
    t = _correr_analisis(c, monkeypatch, "SQL-55", zsd)
    assert t["datos"]["oc"] == "SQL-55" and t["datos"]["oc_origen"] == "Pedidos Ingresados"
    assert c.get("/api/estado").json()["pedidos"][0]["oc"] == "SQL-55"
    with SessionLocal() as s:
        s.get(Pedido, 1).oc = ""
        s.commit()
    t = _correr_analisis(c, monkeypatch, "", zsd)                    # SQL no la tiene: el reporte sí
    assert t["datos"]["oc"] == "ZSD-77" and t["datos"]["oc_origen"] == "ZSD001_03"
    assert c.get("/api/analisis/4005100000").json()["oc_sap"] == "ZSD-77"


def test_una_oc_escrita_a_mano_no_se_pisa_y_el_conflicto_queda_en_el_analisis(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        s.get(Pedido, 1).oc = "MANUAL-1"
        s.commit()
    _correr_analisis(c, monkeypatch, "SQL-55", [])
    assert c.get("/api/estado").json()["pedidos"][0]["oc"] == "MANUAL-1"       # no se reemplaza sola
    a = c.get("/api/analisis/4005100000").json()
    assert a["oc_sap"] == "SQL-55" and a["oc_origen"] == "Pedidos Ingresados"    # la web ofrece "Usar la OC de SAP"


def test_el_endpoint_de_busqueda_responde_con_los_resultados_o_el_error(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    monkeypatch.setattr(bases, "_filas", _Sql([("2026-09-22", "4005175955", "4500123", "2026-10-04")]))
    r = c.get("/api/pedidos-sap", params={"q": "4500123"}).json()
    assert r["ok"] and r["resultados"][0] == {"fecha": "2026-09-22", "pedido": "4005175955", "oc": "4500123", "vence": "2026-10-04"}

    def cae(*a, **k):
        raise RuntimeError("Falta la conexión a SQL Server")
    monkeypatch.setattr(bases, "_filas", cae)
    r = c.get("/api/pedidos-sap", params={"q": "4500123"}).json()
    assert r["ok"] is False and "conexión" in r["error"]


# ---------- vista de un pallet: solo el pallet o todos los pallets dentro del camión ----------
def test_la_vista_de_un_pallet_puede_mostrar_los_pallets_dentro_del_camion(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 PRODUCTO A")
        s.commit()
    cuerpo = {"lineas": [{"sku": "9000", "qty": 400}], "modo": "SDA STOCK", "caja_master": "SIN CAJA MASTER",
              "vista": "pallet", "cliente": "PARIS"}
    solo = c.post("/api/cubicaje-libre", json=cuerpo).json()
    assert len(solo["pallets_disponibles"]) >= 2 and solo["pallet_camion"] is False
    assert [x["tipo"] for x in solo["camiones"]] == ["Pallet"]                 # un pallet aislado, sin camión

    # "En el camión": el camión completo con todos sus pallets y el elegido marcado
    otro = solo["pallets_disponibles"][1]
    d = c.post("/api/cubicaje-libre", json={**cuerpo, "pallet_camion": True, "pallet_n": otro}).json()
    assert d["pallet_camion"] is True and d["pallet_visto"] == otro and d["pallets_disponibles"] == solo["pallets_disponibles"]
    assert all(x["tipo"] != "Pallet" for x in d["camiones"])
    assert d["pallet_vehiculo"] in [x["numero"] for x in d["camiones"]]
    assert d["unidades"] == 400 and len(d["pallets_detalle"]) == len(solo["pallets_disponibles"])
    # fuera de la vista de pallet no hay pallet elegido
    assert "pallet_visto" not in c.post("/api/cubicaje-libre", json={**cuerpo, "vista": "rampla"}).json()


# --- Cargas masivas con el formato de cada cliente --------------------------------------------

def _archivo_cliente(filas, encabezado=None, titulo=True):
    """Un predistribuido como el de Paris: título arriba, una fila por bulto, dos códigos de producto."""
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "OC"
    if titulo:
        ws.append(["Predistribución 1390888"])
        ws.append([])
    ws.append(encabezado or ["ID Bulto", "Código Local", "Local", "SKU", "Cód. Prov.", "Descripción", "Cantidad", "Grupo"])
    for f in filas:
        ws.append(list(f))
    b = BytesIO()
    wb.save(b)
    return b.getvalue()


def test_lector_ubica_columnas_del_cliente_y_suma_por_sucursal():
    from app import cargas
    filas = [("B1", 10, "PLAZA OESTE", "518578999", "240097509", "FREIDORA", 1, "1390888")] * 3 + \
            [("B4", 20, "ALTO LAS CONDES", "518578999", "240097509", "FREIDORA", 1, "1390888")]
    arch = _archivo_cliente(filas)
    # con la Base de Medidas, el código Electrolux es el que coincide (no el del cliente)
    lec = cargas.leer_carga(arch, "p.xlsx", {"240097509"}, obligatorios=("sucursal", "sku", "unidades"))
    assert [(f.sucursal, f.sku, f.qty) for f in lec.filas] == [("PLAZA OESTE", "240097509", 3.0), ("ALTO LAS CONDES", "240097509", 1.0)]
    assert lec.columnas == {"sku": "Cód. Prov.", "unidades": "Cantidad", "sucursal": "Local"}
    assert lec.leidas == 4 and lec.resumen()["sucursales"] == 2
    # sin coincidencias se usa la prioridad de nombres y se avisa que hay que revisarlo
    lec = cargas.leer_carga(arch, "p.xlsx", set(), obligatorios=("sucursal", "sku", "unidades"))
    assert lec.columnas["sku"] == "Cód. Prov." and any("Revisa" in n for n in lec.notas)


def test_lector_csv_y_sin_encabezado_y_errores():
    import pytest
    from app import cargas
    csv = "Tienda;Material;Unidades\nCentro;0900;1.250\nSur;0900;12,5\nSur;;4\n".encode("latin-1")
    lec = cargas.leer_carga(csv, "x.csv", set(), obligatorios=("sucursal", "sku", "unidades"))
    assert [(f.sucursal, f.sku, f.qty) for f in lec.filas] == [("CENTRO", "900", 1250.0), ("SUR", "900", 12.5)]
    assert lec.errores == ["fila 4: falta el SKU"]
    with pytest.raises(ValueError, match="Falta la columna Unidades"):
        cargas.leer_carga(_archivo_cliente([], ["Local", "SKU"], titulo=False), "x.xlsx", set())
    with pytest.raises(ValueError, match=".xlsx"):
        cargas.leer_carga(b"x", "x.pdf", set())


def test_lector_xls_real_de_paris():
    import pytest
    pytest.importorskip("xlrd")
    from pathlib import Path
    from app import cargas
    ruta = Path(__file__).parent / "datos" / "predistribucion_paris.xls"
    lec = cargas.leer_carga(ruta.read_bytes(), ruta.name, {"240097509"}, obligatorios=("sucursal", "sku", "unidades"))
    assert lec.hoja == "OC" and lec.leidas == 50 and len(lec.filas) == 5
    assert {f.sku for f in lec.filas} == {"240097509"} and sum(f.qty for f in lec.filas) == 50
    assert {f.sucursal for f in lec.filas} >= {"PLAZA OESTE", "ALTO LAS CONDES"}
    assert lec.columnas["sku"] == "Cód. Prov."                 # la tilde mal exportada se limpia en la etiqueta


def test_importar_archivo_de_cliente_respeta_el_modo_en_pantalla(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "240097509", "240097509 FREIDORA")
        s.commit()
    filas = [("B1", 10, "PLAZA OESTE", "518578999", "240097509", "F", 1, "g")] * 6 + \
            [("B2", 20, "ALTO LAS CONDES", "518578999", "240097509", "F", 1, "g")] * 4
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("p.xlsx", _archivo_cliente(filas))},
               data={"modo": "MDA PREDISTRIBUIDO"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["modo"] == "MDA PREDISTRIBUIDO" and d["lectura"]["sucursales"] == 2
    assert [(l["sku"], l["qty"]) for l in d["lineas"]] == [("240097509", 10.0)]
    assert not any("se cubicó en" in a for a in d["avisos"])       # ya estaba en ese modo: nada que avisar
    assert sorted((f["sucursal"], f["unidades"]) for f in d["filas"]) == [("ALTO LAS CONDES", 4), ("PLAZA OESTE", 6)]


def test_reglas_del_reparto_avisan_lo_que_no_se_carga(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 A")
        _medida(s, "9001", "9001 B")
        s.commit()
    r = c.post("/api/cubicaje-libre", json={
        "modo": "MDA PREDISTRIBUIDO",
        "lineas": [{"sku": "9000", "qty": 10}, {"sku": "9001", "qty": 5}],
        "predistribuido": [{"sucursal": "A", "sku": "9000", "qty": 8}, {"sucursal": "B", "sku": "9000", "qty": 8}]})
    assert r.status_code == 200, r.text
    av = " ".join(r.json()["avisos"])
    assert "9000 (+6)" in av and "9001 (5 un.)" in av


# --- Carga (piso / pallet) y destino mandan sobre el modo, como H2 en el Excel -------------------

def _libre(c, **body):
    r = c.post("/api/cubicaje-libre", json={"cliente": "", "caja_master": "SIN CAJA MASTER", **body})
    assert r.status_code == 200, r.text
    return r.json()


def test_carga_manda_sobre_el_modo(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 A")
        s.commit()
    lineas = [{"sku": "9000", "qty": 60}]
    sda = _libre(c, modo="SDA STOCK", lineas=lineas)
    assert sda["pallets_detalle"] and sda["unidades"] == 60                   # SDA: en pallets

    # SDA + "a piso": no se arman pallets y se carga todo directo al camión
    piso = _libre(c, modo="SDA STOCK", lineas=lineas, piso_pallet="PISO")
    assert not piso["pallets_detalle"] and piso["unidades"] == 60 and piso["camiones"]
    assert piso["modo"] == "SDA STOCK" and piso["piso_pallet"] == "PISO"      # lo elegido se conserva

    # MDA + "en pallets": manda la carga y se cubica en pallets
    mda = _libre(c, modo="MDA", lineas=lineas)
    assert not mda["pallets_detalle"]
    pal = _libre(c, modo="MDA", lineas=lineas, piso_pallet="PALLET")
    assert pal["pallets_detalle"] and pal["unidades"] == 60


def test_sda_predistribuido_a_piso_va_por_sucursal_y_destino_manda(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 A")
        s.commit()
    reparto = [{"sucursal": "NORTE", "sku": "9000", "qty": 20}, {"sucursal": "SUR", "sku": "9000", "qty": 15}]
    lineas = [{"sku": "9000", "qty": 35}]
    d = _libre(c, modo="SDA PREDISTRIBUIDO", lineas=lineas, predistribuido=reparto, piso_pallet="PISO")
    assert not d["pallets_detalle"] and d["unidades"] == 35
    assert {f["sucursal"] for f in d["filas"]} == {"NORTE", "SUR"}
    # Destino = Stock manda sobre un modo predistribuido: toda la carga junta, sin separar por sucursal
    s2 = _libre(c, modo="SDA PREDISTRIBUIDO", lineas=lineas, predistribuido=reparto, destino="STOCK")
    assert s2["unidades"] == 35 and not any(f.get("sucursal") for f in s2["filas"])
    # Destino = Por sucursal sobre MDA usa el reparto
    p = _libre(c, modo="MDA", lineas=lineas, predistribuido=reparto, destino="SUCURSAL")
    assert {f["sucursal"] for f in p["filas"]} == {"NORTE", "SUR"}


# --- Grupos de carga: el 1 al fondo, cada grupo empieza donde terminó el anterior -----------------

def test_lector_lee_grupos_e_ignora_un_grupo_unico():
    from app import cargas
    arch = _archivo_cliente([("a", 1, "X", "9000", "9000", "A", 5, 1), ("b", 1, "X", "9001", "9001", "B", 3, 2),
                             ("c", 1, "X", "9002", "9002", "C", 2, "")],
                            ["ID", "Código Local", "Local", "SKU", "Cód. Prov.", "Descripción", "Cantidad", "Grupo de carga"],
                            titulo=False)
    lec = cargas.leer_carga(arch, "g.xlsx", {"9000", "9001", "9002"})
    assert [(f.sku, f.grupo) for f in lec.filas] == [("9000", 1), ("9001", 2), ("9002", 0)]
    # la columna Grupo de un cliente que numera su pedido (todas iguales) no ordena nada
    igual = _archivo_cliente([("a", 1, "X", "9000", "9000", "A", 5, "1390888"), ("b", 1, "X", "9001", "9001", "B", 3, "1390888")],
                             titulo=False)
    lec = cargas.leer_carga(igual, "g.xlsx", {"9000", "9001"})
    assert all(f.grupo == 0 for f in lec.filas) and "grupo" not in lec.columnas


def test_grupos_se_cargan_en_bloques_en_orden(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        for sku in ("9000", "9001", "9002"):
            _medida(s, sku, f"{sku} P")
        s.commit()
    r = c.post("/api/cubicaje-libre/importar", files={"file": ("g.xlsx", _excel_carga([("9000", 20, ""), ("9001", 20, ""), ("9002", 10, "")]))})
    assert r.status_code == 200
    # con la columna Grupo (4.ª de la plantilla)
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Carga"
    ws.append(["SKU", "Unidades", "Sucursal (opcional)", "Grupo (opcional)"])
    ws.append(["9002", 10, "", 2]); ws.append(["9000", 20, "", 1]); ws.append(["9001", 20, "", 1])
    b = BytesIO(); wb.save(b)
    d = c.post("/api/cubicaje-libre/importar", files={"file": ("g.xlsx", b.getvalue())}).json()
    assert d["modo"] == "MDA" and len(d["grupos"]) == 3                       # el modo elegido no cambia
    assert d["unidades"] == 50
    assert any("2 grupos" in a for a in d["avisos"])
    # el grupo 1 queda en el fondo (X menor) y el 2 después
    x = {}
    for p in d["visor_json"]["placed"] if "placed" in d.get("visor_json", {}) else []:
        x.setdefault(p.get("suc"), []).append(p)
    # Destino = Stock ignora los grupos: toda la carga junta
    from app.routers.cubicaje import _reparto_por_grupos
    assert [f.sucursal for f in _reparto_por_grupos(d["grupos"])] == ["GRUPO 1".replace("GRUPO 1", "GRUPO 01")] * 2 + ["GRUPO 02"]
    juntos = c.post("/api/cubicaje-libre", json={"cliente": "", "modo": "MDA", "lineas": d["lineas"], "grupos": d["grupos"], "destino": "STOCK"}).json()
    assert juntos["unidades"] == 50 and not any("grupos" in a for a in juntos["avisos"])


def test_detalle_trae_piezas_de_la_caja_master_para_contar_bultos(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import Medida
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "9000", "9000 LICUADORA")
        s.add(Medida(sku="C9000", descripcion="CAJA MASTER", piezas=2, largo=60, ancho=50, alto=40, peso=20,
                     apilar="Y", inclinar="N", rotar="N", max_camion=300, max_pallet=20))
        s.commit()
    d = c.post("/api/cubicaje-libre", json={"cliente": "", "modo": "MDA", "lineas": [{"sku": "C9000", "qty": 50}, {"sku": "9000", "qty": 4}]}).json()
    assert d["detalle_lineas"]["C9000"]["piezas"] == 2 and d["detalle_lineas"]["9000"]["caja"] == 2


def test_exportar_el_analisis_para_los_kam(datos):
    from io import BytesIO
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    c = TestClient(app)
    with SessionLocal() as s:
        doc = _analisis_de_prueba()
        doc["oc_sap"] = "OC-77"
        fila = next(f for f in doc["resultado"]["filas"] if f["sku"] == "9000")
        fila["stock"] = {"cd30": 5, "reserva_cd30": 1, "ec01": 2, "tp01": 0}
        fila["disponibilidad"] = "300 - 15-10-2026"
        domain.guardar_config(s, "analisis:4005100000", doc)
        s.commit()
    r = c.get("/api/analisis/4005100000/excel")
    assert r.status_code == 200 and "Analisis_4005100000" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Análisis", "Entregas"]
    filas = list(wb["Análisis"].iter_rows(min_row=2, values_only=True))
    assert [x[1] for x in filas] and all(x[0] for x in filas)                   # la OC va en cada fila
    f9000 = next(x for x in filas if x[1] == "9000")
    assert f9000[3] == 100 and f9000[4] == 100 and f9000[5:9] == (5, 1, 2, 0) and f9000[9:11] == ("15-10-2026", "300")
    ent = list(wb["Entregas"].iter_rows(min_row=2, values_only=True))
    assert wb["Entregas"][1][0].value == "N° entrega" and all(len(x) == 4 for x in ent)
    assert c.get("/api/analisis/9999999999/excel").status_code == 404


def test_las_cajas_master_cuentan_unidades_de_producto_y_bultos(datos):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import Medida
    c = TestClient(app)
    with SessionLocal() as s:
        _medida(s, "TOST", "TOSTADOR")
        _medida(s, "LIC", "LICUADORA")
        for sku, desc, piezas in (("CTOST", "CAJA MASTER TOSTADOR", 4), ("CLIC", "CAJA MASTER LICUADORA", 2)):
            s.add(Medida(sku=sku, descripcion=desc, piezas=piezas, largo=60, ancho=50, alto=40, peso=20,
                         apilar="Y", inclinar="N", rotar="N", max_camion=300, max_pallet=20))
        s.commit()
    # 10 cajas de 4 tostadores + 5 cajas de 2 licuadoras + 3 licuadoras sueltas
    d = c.post("/api/cubicaje-libre", json={"cliente": "", "modo": "MDA", "lineas": [
        {"sku": "CTOST", "qty": 10}, {"sku": "CLIC", "qty": 5}, {"sku": "LIC", "qty": 3}]}).json()
    assert d["detalle_lineas"]["CTOST"]["master"] is True and d["detalle_lineas"]["LIC"]["master"] is False
    assert d["por_caja_master"] == {"CTOST": 4, "CLIC": 2}
    por_sku = {}
    for f in d["filas"]:
        por_sku[f["sku"]] = por_sku.get(f["sku"], 0) + f["unidades"]
    assert por_sku == {"CTOST": 40, "CLIC": 10, "LIC": 3}              # unidades de producto: 10×4, 5×2 y 3 sueltas
    assert d["unidades"] == 53
    assert all(f.get("bultos") for f in d["filas"] if f["sku"] in ("CTOST", "CLIC"))


def test_el_pallet_de_sodimac_deja_105_cm_de_carga(datos):
    from app.analisis import pallet_cliente
    from app.integrations import clientes as cli_mod
    from app.models import Cliente
    assert pallet_cliente("SODIMAC") == (120, 100, 105)
    with SessionLocal() as s:
        assert cli_mod.buscar(s, "SODIMAC").pallet_alto == 105
        s.get(Cliente, "SODIMAC").pallet_alto = 120.0              # instalación con el valor de fábrica anterior
        s.commit()
        assert cli_mod.buscar(s, "SODIMAC").pallet_alto == 105      # se corrige sola
        s.get(Cliente, "SODIMAC").pallet_alto = 110.0              # cambiado a mano: se respeta
        s.commit()
        assert cli_mod.buscar(s, "SODIMAC").pallet_alto == 110


# ---------- el análisis no debe morir justo después de exportar ZSD001_03 ----------
def test_sin_sql_server_el_analisis_falla_antes_de_tocar_sap_con_un_motivo_claro(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.integrations import bases, sap
    from app.main import app
    c = TestClient(app)
    llamadas = []
    monkeypatch.setattr(sap, "leer_pedido", lambda *a: llamadas.append("vl01n"))
    monkeypatch.setattr(sap, "zsd001_03", lambda *a, **k: llamadas.append("zsd"))

    def sin_red(grupo):
        raise RuntimeError("('08001', '[08001] Named Pipes Provider: Could not open a connection to SQL Server [53]')")
    monkeypatch.setattr(bases, "plan_sop", sin_red)
    r = c.post("/api/analisis/4005100000", json={"puesto": "PN01", "cliente": "PARIS", "fecha": "2026-10-01"})
    t = _esperar(c, r.json()["id"])
    assert t["estado"] == "error"
    assert "SQL Server" in t["error"] and "no se empezó con SAP" in t["error"] and "VPN" in t["error"]
    assert t["paso"].startswith("0/4")
    assert llamadas == []                                    # no se abrió SAP ni se exportó nada


def test_el_error_de_un_paso_dice_en_cual_se_detuvo(monkeypatch, datos):
    from fastapi.testclient import TestClient
    from app.integrations import base_medidas, bases, sap
    from app.main import app
    c = TestClient(app)
    monkeypatch.setattr(sap, "leer_pedido", lambda *a: SimpleNamespace(
        posiciones=[SimpleNamespace(sku="9000", qty_entrega=1, qty_pendiente=1)], aviso=""))

    def zsd(*a, **k):
        raise sap.ErrorSap("SAP no guardó el archivo de ZSD001_03 (SAP muestra una ventana: «Seguridad de SAP GUI»).")
    monkeypatch.setattr(sap, "zsd001_03", zsd)
    monkeypatch.setattr(bases, "plan_sop", lambda g: {})
    monkeypatch.setattr(bases, "disponibilidad", lambda: {})
    monkeypatch.setattr(base_medidas, "medidas", lambda: {})
    r = c.post("/api/analisis/4005100000", json={"puesto": "PN01", "cliente": "PARIS", "fecha": "2026-10-01"})
    t = _esperar(c, r.json()["id"])
    assert t["estado"] == "error" and "Seguridad de SAP GUI" in t["error"]
    assert t["paso"].startswith("2/4")


def test_se_puede_retomar_el_seguimiento_de_un_trabajo_en_curso(datos):
    from fastapi.testclient import TestClient
    from app.integrations import acciones
    from app.main import app
    c = TestClient(app)
    acciones._trabajos.clear()
    acciones._registrar({"id": "viejo", "estado": "ok", "accion": "x", "label": "x", "args": [], "inicio": "", "fin": ""})
    acciones._registrar({"id": "vivo", "estado": "en_curso", "accion": "analizar_pedido", "label": "Analizar pedido",
                         "args": ["1"], "inicio": "2026-10-08T10:00:00", "fin": "", "progreso": "2/4"})
    acciones._registrar({"id": "nuevo", "estado": "ok", "accion": "x", "label": "x", "args": [], "inicio": "", "fin": ""})
    lista = c.get("/api/acciones/trabajos").json()
    assert lista[0]["id"] == "vivo"                           # el que sigue en curso va primero
    acciones._trabajos.clear()


def test_exportacion_de_zsd_se_lee_solo_cuando_esta_completa_y_no_devuelve_vacio_en_silencio(tmp_path):
    import threading
    import time as _t
    from openpyxl import Workbook
    from app.integrations import sap

    class SesionFalsa:                                          # sin ventanas abiertas
        def findById(self, ident):
            raise KeyError(ident)

    destino = tmp_path / "Qty En Entrega 1.xlsx"

    def escribir_de_a_poco():                                  # SAP crea el archivo y lo termina de llenar después
        destino.write_bytes(b"PK")
        _t.sleep(0.6)
        wb = Workbook()
        ws = wb.active
        ws.append(["Material", "Nombre Codigo de Material", "Qty. En Entrega"])
        ws.append(["9000", "PRODUCTO", 7])
        wb.save(destino)
    hilo = threading.Thread(target=escribir_de_a_poco)
    hilo.start()
    sap._esperar_archivo(SesionFalsa(), destino, maximo=20)
    filas = sap.leer_export(destino)
    hilo.join()
    assert filas == [{"Material": "9000", "Nombre Codigo de Material": "PRODUCTO", "Qty. En Entrega": 7}]

    vacio = tmp_path / "vacio.xlsx"
    vacio.write_bytes(b"")
    with pytest.raises(sap.ErrorSap, match="vacío"):
        sap.leer_export(vacio)
    raro = tmp_path / "raro.xlsx"
    raro.write_text("hola\tmundo\n1\t2\n", encoding="utf-8")
    with pytest.raises(sap.ErrorSap, match="columna de material"):
        sap.leer_export(raro)


def test_si_sap_no_guarda_el_archivo_se_nombra_la_ventana_en_vez_de_esperar_en_vano(tmp_path):
    from app.integrations import sap

    class Ventana:
        Text = "Seguridad de SAP GUI"

    class SesionConAviso:
        def findById(self, ident):
            if ident == "wnd[1]":
                return Ventana()
            raise KeyError(ident)

    t0 = time.time()
    with pytest.raises(sap.ErrorSap, match="Seguridad de SAP GUI"):
        sap._esperar_archivo(SesionConAviso(), tmp_path / "no_llega.xlsx", maximo=30)
    assert time.time() - t0 < 12


def test_carpeta_de_exportacion_es_fija_y_se_comprueba_antes_de_abrir_sap(tmp_path, monkeypatch):
    from app.config import settings
    from app.integrations import sap
    buena = tmp_path / "sap"
    assert sap.carpeta_export(buena) == buena and buena.is_dir() and not (buena / ".escritura").exists()
    assert sap.carpeta_export(buena) == buena                                # siempre la misma
    larga = tmp_path / ("x" * 120)                                           # ruta que SAP maneja mal: cae a la de respaldo
    assert sap.carpeta_export(larga) != larga
    object.__setattr__(settings, "sap_export_dir", str(tmp_path / "fija"))
    try:
        assert sap.carpeta_export(buena) == tmp_path / "fija"                # la configurada manda
    finally:
        object.__setattr__(settings, "sap_export_dir", "")


def test_el_export_se_publica_en_la_ruta_fija_con_el_nombre_fijo(tmp_path):
    from app.config import settings
    from app.integrations import sap
    pub = tmp_path / "Script Pendiente"
    pub.mkdir()
    local = tmp_path / "local.xlsx"
    local.write_bytes(b"contenido-del-export")
    object.__setattr__(settings, "zsd_publicar_dir", str(pub))
    try:
        assert sap.revisar_publicacion() == ""
        r = sap.publicar_export(local)
        assert r["ok"] and r["ruta"] == str(pub / "Qty En Entrega.xlsx")
        assert (pub / "Qty En Entrega.xlsx").read_bytes() == b"contenido-del-export"
        assert [f.name for f in pub.iterdir()] == ["Qty En Entrega.xlsx"]        # sin temporales ni pruebas sueltas
        local.write_bytes(b"segundo")                                           # un análisis nuevo reemplaza al anterior
        assert sap.publicar_export(local)["ok"] and (pub / "Qty En Entrega.xlsx").read_bytes() == b"segundo"
        object.__setattr__(settings, "zsd_publicar_dir", str(tmp_path / "no_existe"))
        assert "No se llega" in sap.revisar_publicacion()
        object.__setattr__(settings, "zsd_publicar_dir", "")
        assert sap.revisar_publicacion() == "" and sap.publicar_export(local)["omitido"]   # vacío = no publicar
    finally:
        object.__setattr__(settings, "zsd_publicar_dir", "")


def test_si_no_se_puede_publicar_el_motivo_es_claro_y_no_se_pierde_el_analisis(tmp_path, monkeypatch):
    from app.config import settings
    from app.integrations import sap
    monkeypatch.setattr(sap.time, "sleep", lambda s: None)
    local = tmp_path / "local.xlsx"
    local.write_bytes(b"x")
    object.__setattr__(settings, "zsd_publicar_dir", str(tmp_path / "red_caida"))
    try:
        r = sap.publicar_export(local)
        assert not r["ok"] and "No se pudo publicar" in r["motivo"] and "Qty En Entrega.xlsx" in r["motivo"]
    finally:
        object.__setattr__(settings, "zsd_publicar_dir", "")


# ---------- lo que cambia de un analista a otro en SAP ----------
class _Layout:
    def __init__(self, textos):
        self.t, self.RowCount = textos, len(textos)

    def GetCellValue(self, i, col):
        return self.t[i]


def test_layout_se_busca_por_nombre_o_se_avisa_si_la_fila_43_no_existe():
    from app.config import settings
    from app.integrations import sap
    pocos = _Layout([f"Layout {i}" for i in range(10)] + ["/QTY ENTREGA"])
    with pytest.raises(sap.ErrorSap, match="fila 43.*11 layouts"):                 # el analista con pocos layouts
        sap._fila_layout(pocos)
    object.__setattr__(settings, "zsd_layout", "qty entrega")
    try:
        assert sap._fila_layout(pocos) == 10                                      # por nombre, sin depender de la posición
        object.__setattr__(settings, "zsd_layout", "no existe")
        with pytest.raises(sap.ErrorSap, match="No hay un layout"):
            sap._fila_layout(pocos)
    finally:
        object.__setattr__(settings, "zsd_layout", "")
    assert sap._fila_layout(_Layout(["x"] * 60)) == 43                            # con la lista completa, como la macro


def test_un_export_sin_las_columnas_necesarias_no_deja_la_qty_en_cero_en_silencio():
    from app.integrations import sap
    sap.validar_columnas([{"Material": "9", "Nombre Código de Material": "X", "Qty. En Entrega": 1}])     # ok (con tilde)
    sap.validar_columnas([])                                                                              # sin filas: nada que validar
    with pytest.raises(sap.ErrorSap, match="Qty. En Entrega.*Layout elegido: «MI LAYOUT»"):
        sap.validar_columnas([{"Material": "9", "Nombre Codigo de Material": "X", "Cantidad": 1}], "MI LAYOUT")


def test_conectar_elige_la_sesion_con_usuario_y_respeta_sap_sistema(monkeypatch):
    import sys
    from types import SimpleNamespace as NS
    from app.config import settings
    from app.integrations import sap

    def ses(sis, usu):
        return NS(Info=NS(SystemName=sis, Client="300", User=usu))

    class Lista:
        def __init__(self, items):
            self.items, self.Count = items, len(items)

        def __call__(self, i):
            return self.items[i]
    sin_login, qas, prd = ses("PRD", ""), ses("QAS", "ANA"), ses("PRD", "ANA")
    con1 = NS(Children=Lista([sin_login, qas]))
    con2 = NS(Children=Lista([prd]))
    gui = NS(Children=Lista([con1, con2]))
    fake = NS(client=NS(GetObject=lambda n: NS(GetScriptingEngine=gui)))
    monkeypatch.setitem(sys.modules, "win32com", fake)
    monkeypatch.setitem(sys.modules, "win32com.client", fake.client)
    assert sap.conectar() is qas                                   # la primera con usuario iniciado, no la de la conexión sin login
    object.__setattr__(settings, "sap_sistema", "prd")
    try:
        assert sap.conectar() is prd and sap.ultima_sesion == "PRD/300 · ANA"
        object.__setattr__(settings, "sap_sistema", "DEV")
        with pytest.raises(sap.ErrorSap, match="SAP_SISTEMA=DEV"):
            sap.conectar()
    finally:
        object.__setattr__(settings, "sap_sistema", "")


def test_la_limpieza_no_toca_la_ultima_exportacion_ni_las_recientes(tmp_path):
    import os
    from app.integrations import sap
    vieja, reciente, fija = (tmp_path / n for n in ("Qty En Entrega 1 viejo.xlsx", "Qty En Entrega 2 nuevo.xlsx", "Qty En Entrega.xlsx"))
    for f in (vieja, reciente, fija):
        f.write_bytes(b"x")
    antes = time.time() - 3 * 86400
    os.utime(vieja, (antes, antes))
    os.utime(fija, (antes, antes))
    sap._limpiar_exportaciones(tmp_path)
    assert not vieja.exists() and reciente.exists() and fija.exists()
