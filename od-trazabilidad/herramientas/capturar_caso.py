"""Captura un caso de control del cubicaje desde el libro de Excel.

Guarda en casos/<nombre>.json la ENTRADA (parámetros, posiciones, medidas,
pallet del cliente, camiones) y el RESULTADO ESPERADO (hojas 03 y 04 tal como
las dejó el VBA). Sirve para verificar cualquier motor futuro contra lo que
hoy produce el Excel, que es la referencia.

Uso:
    py herramientas/capturar_caso.py "V:\\ruta\\libro.xlsm" nombre_del_caso

Hazlo con el cubicaje ya ejecutado y las hojas 03/04 escritas.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
DIR_CASOS = BASE / "casos"


def txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def num(v) -> float:
    try:
        return round(float(v), 4)
    except (TypeError, ValueError):
        return 0.0


def filas(ws, fila_encabezado: int, columnas: int) -> list[dict]:
    """Devuelve las filas con datos, usando la fila de encabezado dada."""
    cab = [txt(c.value) for c in ws[fila_encabezado][:columnas]]
    out = []
    for fila in ws.iter_rows(min_row=fila_encabezado + 1, max_col=columnas, values_only=True):
        if all(c is None or txt(c) == "" for c in fila):
            continue
        out.append({k: txt(v) for k, v in zip(cab, fila) if k})
    return out


def capturar(ruta_libro: str, nombre: str) -> Path:
    wb = load_workbook(ruta_libro, data_only=True, read_only=True)
    ent = wb["01_Entrada"]

    caso = {"nombre": nombre, "libro": Path(ruta_libro).name}
    caso["parametros"] = {
        "pedido": txt(ent["A2"].value), "puesto": txt(ent["B2"].value),
        "cliente": txt(ent["D2"].value), "tipo_pedido": txt(ent["E2"].value),
        "caja_master": txt(ent["F2"].value), "qty_pedidos": txt(ent["G2"].value),
        "tipo_carga": txt(ent["H2"].value),
    }

    # Posiciones a cubicar (02_Posiciones, encabezado en la fila 4)
    caso["posiciones"] = [
        {"sku": f.get("SKU", ""), "descripcion": f.get("DESCRIPCION", ""),
         "caja_master": f.get("Caja Master", ""), "qty_entrega": num(f.get("CARGA") or 0),
         "pedido": f.get("Pedido2", "")}
        for f in filas(wb["02_Posiciones"], 4, 14) if f.get("SKU")
    ]
    skus = {p["sku"] for p in caso["posiciones"]}

    # Medidas (Base Medidas, encabezado en la fila 5). Se guardan las del caso
    # y también las cajas master C+SKU, que empiezan con C.
    caso["medidas"] = [
        f for f in filas(wb["Base Medidas"], 5, 12)
        if f.get("Grupo") and (f["Grupo"] in skus or f["Grupo"].lstrip("C") in skus)
    ]

    cliente = caso["parametros"]["cliente"].upper()
    caso["pallet"] = next((f for f in filas(wb["Clientes"], 1, 6)
                           if f.get("Cliente", "").upper() == cliente), {})
    caso["camiones"] = filas(wb["Config"], 1, 4)
    caso["predistribuido"] = filas(wb["Predistribuido"], 1, 5)

    # Resultado que hay que reproducir
    caso["esperado"] = {
        "03_PedidoCubicado": filas(wb["03_PedidoCubicado"], 1, 13),
        "04_CubicajeSDA": filas(wb["04_CubicajeSDA"], 1, 10),
    }
    wb.close()

    if not caso["esperado"]["03_PedidoCubicado"] and not caso["esperado"]["04_CubicajeSDA"]:
        raise SystemExit("Las hojas 03 y 04 están vacías: ejecuta el cubicaje antes de capturar.")

    DIR_CASOS.mkdir(exist_ok=True)
    destino = DIR_CASOS / f"{nombre}.json"
    destino.write_text(json.dumps(caso, ensure_ascii=False, indent=1), encoding="utf-8")
    return destino


def main():
    if len(sys.argv) != 3:
        raise SystemExit('Uso: py herramientas/capturar_caso.py "ruta\\libro.xlsm" nombre_caso')
    destino = capturar(sys.argv[1], sys.argv[2])
    caso = json.loads(destino.read_text(encoding="utf-8"))
    print(f"Caso guardado en {destino}")
    print(f"  Pedido {caso['parametros']['pedido']} · {caso['parametros']['cliente']} · "
          f"{caso['parametros']['tipo_pedido']}")
    print(f"  {len(caso['posiciones'])} posiciones · {len(caso['medidas'])} medidas")
    print(f"  Esperado: {len(caso['esperado']['03_PedidoCubicado'])} filas en 03, "
          f"{len(caso['esperado']['04_CubicajeSDA'])} en 04")


if __name__ == "__main__":
    main()
