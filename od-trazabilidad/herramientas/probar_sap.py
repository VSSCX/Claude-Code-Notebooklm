"""Lee un pedido desde SAP con Python y lo compara con lo que dejó el Excel.

Paso 1: en Excel, corre tu botón "01 Leer Pedido SAP" con el pedido y guarda el libro.
Paso 2:

    py herramientas\\probar_sap.py 4005171502 PN01 17.09.2026 "V:\\ruta\\libro.xlsm"

Sin la ruta del libro solo muestra lo que leyó Python. No escribe nada en SAP.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.integrations.sap import ErrorSap, leer_pedido  # noqa: E402


def del_excel(ruta: str, pedido: str) -> dict[str, tuple[float, float]]:
    from openpyxl import load_workbook
    ws = load_workbook(ruta, data_only=True, read_only=True)["02_Posiciones"]
    cab = {str(c.value).strip().lower(): c.column for c in ws[4] if c.value}
    out = {}
    for fila in ws.iter_rows(min_row=5, values_only=True):
        sku = fila[cab["sku"] - 1]
        if not sku or str(fila[cab["pedido2"] - 1] or "").strip() != pedido:
            continue
        out[str(sku).strip()] = (float(fila[cab["qty entrega"] - 1] or 0),
                                 float(fila[cab["qty pendiente"] - 1] or 0))
    return out


def main():
    if len(sys.argv) < 4:
        raise SystemExit('Uso: py herramientas\\probar_sap.py PEDIDO PUESTO dd.MM.yyyy ["libro.xlsm"]')
    pedido, puesto, fecha = sys.argv[1], sys.argv[2], sys.argv[3]
    try:
        lectura = leer_pedido(pedido, puesto, fecha)
    except ErrorSap as e:
        raise SystemExit(f"ERROR: {e}")

    print(f"\nPython leyó {len(lectura.posiciones)} posiciones del pedido {pedido}")
    if lectura.aviso:
        print(f"Aviso: {lectura.aviso}")
    for p in lectura.posiciones:
        print(f"  {p.sku:<18} entrega {p.qty_entrega:>10.2f}   pendiente {p.qty_pendiente:>10.2f}")

    if len(sys.argv) < 5:
        return
    excel = del_excel(sys.argv[4], pedido)
    py = {p.sku: (p.qty_entrega, p.qty_pendiente) for p in lectura.posiciones}
    print(f"\nExcel tiene {len(excel)} posiciones de ese pedido")

    difs = []
    for sku in sorted(set(py) | set(excel)):
        a, b = py.get(sku), excel.get(sku)
        if a is None:
            difs.append(f"  {sku}: solo en Excel {b}")
        elif b is None:
            difs.append(f"  {sku}: solo en Python {a}")
        elif abs(a[0] - b[0]) > 0.001 or abs(a[1] - b[1]) > 0.001:
            difs.append(f"  {sku}: Python {a} vs Excel {b}")
    if difs:
        print("\nDIFERENCIAS:")
        print("\n".join(difs))
    else:
        print("\nOK: Python y Excel leyeron exactamente lo mismo.")


if __name__ == "__main__":
    main()
