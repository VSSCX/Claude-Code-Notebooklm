"""Lectura de cargas masivas: la plantilla de la plataforma y los archivos de los clientes.

Cada cliente manda su predistribuido con otro formato (Paris trae una fila por bulto, con el
código del cliente y el de Electrolux lado a lado). De todos hace falta lo mismo: producto,
unidades y sucursal. Aquí se ubican esas tres cosas por el nombre de la columna, se elige la
columna de SKU que coincide con la Base de Medidas y se suman las filas repetidas.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from . import domain

EXTENSIONES = (".xlsx", ".xlsm", ".xls", ".csv")

# Los alias van normalizados (sin tildes, minúsculas, sin puntos ni paréntesis).
ALIAS = {
    # nombre de la sucursal: lo que lee una persona
    "sucursal": ["sucursal", "suc", "tienda", "local", "nombre local", "nombre sucursal", "nombre tienda",
                 "punto de venta", "pdv", "sala", "destino", "nombre destino", "bodega destino", "bodega"],
    # código de la sucursal: solo sirve cuando el archivo no trae el nombre
    "sucursal_cod": ["codigo local", "cod local", "n local", "numero local", "cod sucursal",
                     "codigo sucursal", "cod tienda", "codigo tienda", "centro", "cdigo local", "cd local"],
    # varios candidatos: gana el que coincide con la Base de Medidas. En empate, este orden.
    "sku": ["material", "cod prov", "cd prov", "codigo proveedor", "cdigo proveedor", "cod proveedor", "sku", "codigo", "codigo material",
            "cod material", "codigo producto", "cod producto", "articulo", "item",
            "cdigo"],      # "Cód." mal codificado (�) llega sin la vocal
    "unidades": ["unidades", "cantidad", "qty", "unid", "cant", "unidades pedidas", "cantidad pedida",
                 "cantidad a despachar", "unidades a despachar", "un"],
    "bulto": ["unidades por bulto", "por bulto", "bulto", "un por bulto"],
    "descripcion": ["descripcion", "desc", "producto", "nombre producto", "descripcion producto"],
}
NOMBRE_CAMPO = {"sucursal": "Sucursal", "sku": "SKU", "unidades": "Unidades"}
ENCABEZADO_PLANTILLA = ["sku", "unidades", "sucursal"]


@dataclass
class Fila:
    n: int
    sku: str
    qty: float
    sucursal: str = ""
    bulto: int = 0


@dataclass
class Lectura:
    filas: list[Fila] = field(default_factory=list)          # ya sumadas por sucursal y SKU
    errores: list[str] = field(default_factory=list)
    hoja: str = ""
    columnas: dict = field(default_factory=dict)             # campo -> nombre de la columna leída
    leidas: int = 0                                          # filas del archivo con datos
    notas: list[str] = field(default_factory=list)

    @property
    def con_sucursal(self) -> bool:
        return any(f.sucursal for f in self.filas)

    def resumen(self) -> dict:
        return {"hoja": self.hoja, "columnas": self.columnas, "filas_leidas": self.leidas,
                "filas": len(self.filas), "unidades": sum(f.qty for f in self.filas),
                "sucursales": len({f.sucursal for f in self.filas if f.sucursal}), "notas": self.notas}


def norm_col(x) -> str:
    t = unicodedata.normalize("NFKD", str(x if x is not None else "")).encode("ascii", "ignore").decode().lower()
    return " ".join(re.sub(r"\(.*?\)|[.:_/-]", " ", t).split())


def _vacio(x) -> bool:
    return x is None or str(x).strip() == ""


def _hojas(contenido: bytes, nombre: str) -> list[tuple[str, list[list]]]:
    """[(hoja, filas)] de cualquier formato. Las filas son listas de valores."""
    ext = ("." + nombre.rsplit(".", 1)[-1].lower()) if "." in nombre else ""
    if ext not in EXTENSIONES:
        raise ValueError("El archivo debe ser Excel (.xlsx, .xlsm, .xls) o .csv")
    try:
        if ext == ".csv":
            for enc in ("utf-8-sig", "latin-1"):
                try:
                    texto = contenido.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            muestra = texto[:4096]
            delim = max(";,\t|", key=muestra.count)
            return [("CSV", [list(f) for f in csv.reader(io.StringIO(texto), delimiter=delim)])]
        if ext == ".xls":
            try:
                import xlrd
            except ImportError as e:
                raise ValueError("Para leer .xls hace falta el paquete xlrd (pip install xlrd). "
                                 "Mientras tanto, guarda el archivo como .xlsx.") from e
            libro = xlrd.open_workbook(file_contents=contenido)
            return [(h.name, [h.row_values(i) for i in range(h.nrows)]) for h in libro.sheets()]
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        return [(h.title, [list(f) for f in h.iter_rows(values_only=True)]) for h in libro.worksheets]
    except ValueError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"No se pudo leer el archivo: {str(e)[:150]}") from e


def _ubicar(fila: list) -> dict[str, list[int]]:
    """Columnas de una fila de encabezado: campo -> índices (en orden de aparición)."""
    cols: dict[str, list[int]] = {}
    for j, x in enumerate(fila):
        n = norm_col(x)
        if not n:
            continue
        for campo, alias in ALIAS.items():
            if n in alias:
                cols.setdefault(campo, []).append(j)
                break
    return cols


def _puntaje(cols: dict) -> int:
    # unidades y SKU pesan más: una fila con solo "Descripción" no es el encabezado
    return sum(3 if c in ("unidades", "sku") else 1 for c in cols)


def _elegir_hoja(hojas, preferidas=("carga", "predistribuido")):
    for pref in preferidas:
        for h in hojas:
            if norm_col(h[0]) == pref and h[1]:
                return h
    mejor, puntaje = None, -1
    for h in hojas:
        p = max((_puntaje(_ubicar(f)) for f in h[1][:20]), default=0)
        if p > puntaje and h[1]:
            mejor, puntaje = h, p
    return mejor


def _cantidad(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):         # 1.250 / 1.250,5
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _texto_sku(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)                                           # .xls entrega 240097509.0
    return domain.norm_sku(v)


def leer_carga(contenido: bytes, nombre: str, conocidos: set[str] | None = None,
               obligatorios: tuple = ("sku", "unidades"), hojas_preferidas=("carga", "predistribuido"),
               orden_plantilla: tuple = tuple(ENCABEZADO_PLANTILLA)) -> Lectura:
    """Lee una carga desde un archivo de cualquier cliente. `conocidos`: SKU (en minúsculas) de la
    Base de Medidas, para elegir entre varias columnas de código. Lanza ValueError si no se puede."""
    hojas = _hojas(contenido, nombre)
    hoja = _elegir_hoja(hojas, hojas_preferidas)
    if hoja is None or not any(any(not _vacio(x) for x in f) for f in hoja[1]):
        raise ValueError("El archivo está vacío.")
    nombre_hoja, filas = hoja
    lec = Lectura(hoja=nombre_hoja)

    # encabezado: la fila con más columnas reconocidas entre las primeras 20 (los clientes ponen títulos arriba)
    mejor, fila_enc = {}, -1
    for i, f in enumerate(filas[:20]):
        cols = _ubicar(f)
        if _puntaje(cols) > _puntaje(mejor):
            mejor, fila_enc = cols, i
    if mejor:
        cols, nombres = mejor, filas[fila_enc]
        datos = list(enumerate(filas[fila_enc + 1:], start=fila_enc + 2))
    else:                                                    # sin encabezado: orden de la plantilla
        cols = {c: [j] for j, c in enumerate(orden_plantilla)}
        nombres = [NOMBRE_CAMPO.get(c, c.capitalize()) for c in orden_plantilla]
        datos = list(enumerate(filas, start=1))
        lec.notas.append("El archivo no trae encabezado: se leyó en el orden " + ", ".join(NOMBRE_CAMPO.get(c, c) for c in orden_plantilla) + ".")

    faltan = [NOMBRE_CAMPO[c] for c in obligatorios if c not in cols]
    if faltan:
        raise ValueError("Falta la columna " + " y ".join(faltan) + " en el encabezado.")

    def etiqueta(j):
        if j >= len(nombres) or _vacio(nombres[j]):
            return f"columna {j + 1}"
        return str(nombres[j]).strip().replace("C\ufffdd", "Cód").replace("\ufffd", "")   # tildes que el cliente exportó mal

    # SKU: de varios candidatos, el que más productos reconoce en la Base de Medidas
    j_sku = cols["sku"][0]
    if len(cols["sku"]) > 1:
        def aciertos(j):
            vistos = {_texto_sku(f[j]).lower() for _, f in datos if j < len(f) and not _vacio(f[j])}
            return len(vistos & conocidos) if conocidos else 0
        puntajes = {j: aciertos(j) for j in cols["sku"]}
        def prioridad(j):                                    # en empate, el orden de ALIAS["sku"]
            n = norm_col(nombres[j])
            return len(ALIAS["sku"]) - ALIAS["sku"].index(n) if n in ALIAS["sku"] else 0
        j_sku = max(cols["sku"], key=lambda j: (puntajes[j], prioridad(j)))
        otros = [etiqueta(j) for j in cols["sku"] if j != j_sku]
        if conocidos and puntajes[j_sku]:
            lec.notas.append(f"Código de producto: «{etiqueta(j_sku)}» (coincide con tu Base de Medidas; "
                             f"se ignoró {', '.join('«' + o + '»' for o in otros)}).")
        else:
            lec.notas.append(f"El archivo trae varias columnas de código ({etiqueta(j_sku)}, {', '.join(otros)}) y "
                             "ninguna coincide con tu Base de Medidas: se usó «" + etiqueta(j_sku) + "». Revisa que sea el código Electrolux.")
    j_suc = (cols.get("sucursal") or [None])[0]
    j_suc_cod = (cols.get("sucursal_cod") or [None])[0]
    j_qty = cols["unidades"][0] if "unidades" in cols else 1
    j_bulto = (cols.get("bulto") or [None])[0]

    lec.columnas = {"sku": etiqueta(j_sku), "unidades": etiqueta(j_qty)}
    if j_suc is not None:
        lec.columnas["sucursal"] = etiqueta(j_suc)
    elif j_suc_cod is not None:
        lec.columnas["sucursal"] = etiqueta(j_suc_cod)
        lec.notas.append("El archivo trae el código de la sucursal pero no su nombre.")

    def celda(f, j):
        return f[j] if j is not None and j < len(f) else None

    suma: dict[tuple, Fila] = {}
    for n, f in datos:
        sku_v, qty_v = celda(f, j_sku), celda(f, j_qty)
        suc = celda(f, j_suc)
        if _vacio(suc):
            suc = celda(f, j_suc_cod)
        if all(_vacio(x) for x in (sku_v, qty_v, suc)):
            continue
        lec.leidas += 1
        sku = _texto_sku(sku_v)
        if not sku:
            lec.errores.append(f"fila {n}: falta el SKU")
            continue
        qty = _cantidad(qty_v) if not _vacio(qty_v) else 0.0
        if qty is None:
            lec.errores.append(f"fila {n}: '{qty_v}' no es una cantidad")
            continue
        if qty <= 0:
            continue
        suc = str(suc if not _vacio(suc) else "").strip().upper()
        if isinstance(celda(f, j_suc), float) and not _vacio(suc):
            suc = re.sub(r"\.0$", "", suc)
        if "sucursal" in obligatorios and not suc:
            lec.errores.append(f"fila {n}: faltan datos (sucursal, SKU y unidades)")
            continue
        bulto = 0
        if j_bulto is not None and not _vacio(celda(f, j_bulto)):
            b = _cantidad(celda(f, j_bulto))
            if b is None:
                lec.errores.append(f"fila {n}: '{celda(f, j_bulto)}' no es una cantidad por bulto")
                continue
            bulto = max(0, int(b))
        clave = (suc, sku.lower())
        if clave in suma:
            suma[clave].qty += qty
        else:
            suma[clave] = Fila(n=n, sku=sku, qty=qty, sucursal=suc, bulto=bulto)
    lec.filas = list(suma.values())
    if lec.leidas > len(lec.filas) and lec.filas and not lec.errores:
        lec.notas.append(f"{lec.leidas} filas del archivo se sumaron en {len(lec.filas)} (misma sucursal y producto).")
    return lec
