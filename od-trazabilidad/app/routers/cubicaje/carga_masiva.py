"""Plantilla de carga masiva del cubicador y su importación."""

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from ...db import get_session
from ...integrations import medidas as med_mod
from .libre import cubicaje_libre
from .predistribuido import _conocidos

router = APIRouter()


PLANTILLA_COLS = ["SKU", "Unidades", "Grupo (opcional)", "Sucursal (opcional)"]


@router.get("/cubicaje-libre/plantilla")
def plantilla_carga(s: Session = Depends(get_session)):
    """Excel para armar una carga fuera de la plataforma y después importarla (con ejemplo de grupos)."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Carga"
    ws.append(PLANTILLA_COLS)
    ejemplos = [f[0] for f in med_mod.filas_para_cubicaje(s)[:6]]
    ejemplos += [x for x in ("900081624", "900276671", "900276672") if x not in ejemplos][:max(0, 3 - len(ejemplos))]
    grupos_ejemplo = [1, 1, 2, 2, 3, 3]                       # el ejemplo ya muestra cómo se arman las tandas
    for i, sku in enumerate(ejemplos):
        ws.append([sku, 10, grupos_ejemplo[i % len(grupos_ejemplo)], ""])
    obligatoria, opcional = PatternFill("solid", fgColor="1A2B4A"), PatternFill("solid", fgColor="2442D6")
    for i, celda in enumerate(ws[1]):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = obligatoria if i < 2 else opcional           # azul oscuro: obligatorias; azul: opcionales
        celda.alignment = Alignment(horizontal="center", vertical="center")
    ws["C1"].comment = Comment("Opcional. 1, 2, 3… para cargar por tandas, en orden.\nEl grupo 1 va al fondo del camión y cada grupo "
                               "empieza donde terminó el anterior.\nSi lo dejas vacío en todas las filas, el cubicador busca la mejor carga.", "Order Desk")
    ws["D1"].comment = Comment("Opcional. Solo para clientes con reparto por sucursal.\nSi usas sucursal, el grupo no se aplica.", "Order Desk")
    for col, ancho in zip("ABCD", (18, 12, 18, 22)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for celda in fila:
            celda.border = Border(bottom=Side(style="thin", color="CBD1D7"))
        fila[2].alignment = Alignment(horizontal="center")
    val_g = DataValidation(type="whole", operator="between", formula1="1", formula2="99", allow_blank=True,
                           errorTitle="Grupo", error="El grupo es un número entero: 1, 2, 3…", showErrorMessage=True)
    val_u = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                           errorTitle="Unidades", error="Escribe un número de unidades (0 o más).", showErrorMessage=True)
    ws.add_data_validation(val_g)
    ws.add_data_validation(val_u)
    val_g.add("C2:C2000")
    val_u.add("B2:B2000")
    ws.freeze_panes = "A2"

    ayuda = wb.create_sheet("Cómo se usa")
    filas_ayuda = [
        ["Carga masiva para el cubicador"],
        [""],
        ["Lo básico"],
        ["1. Un producto por fila en la hoja «Carga»: SKU y Unidades. Las filas en 0 o vacías se ignoran."],
        ["2. SKU: el código del producto. Si tiene caja master, usa C delante (C900081624)."],
        ["3. Los productos tienen que estar en la Base de Medidas (Configuración). Los que no estén no se cubican:"],
        ["   se avisa su SKU y sus unidades para que los agregues."],
        [""],
        ["Cargar por grupos (cuando hay un orden que respetar)"],
        ["Por defecto el cubicador busca la forma más óptima de cargar. Si un pedido se tiene que cargar por tandas,"],
        ["pon un número de grupo en cada producto:"],
        ["   • El grupo 1 va al fondo del camión."],
        ["   • El grupo 2 empieza donde terminó el 1, y así sucesivamente."],
        ["   • Lo que no tenga grupo se carga al final."],
        ["   • En pallets, cada grupo arma sus propios pallets."],
        ["   • Se acepta 1, 2, 3 o también «Grupo 1», «G2». Con un solo grupo para todo no hay nada que ordenar: se ignora."],
        ["Ejemplo (las primeras filas de la hoja «Carga» ya lo muestran):"],
        ["   Grupo 1 → productos A, B (van al fondo)  |  Grupo 2 → C, D (siguen)  |  Grupo 3 → E, F (al final)"],
        [""],
        ["Reparto por sucursal (solo clientes predistribuidos)"],
        ["Repite el producto en una fila por sucursal. Con sucursales, el cubicador pasa solo a MDA o SDA predistribuido,"],
        ["y si usas sucursal, el grupo no se aplica (manda la sucursal)."],
        [""],
        ["Con el archivo del cliente, sin tocarlo"],
        ["Puedes subir directo el predistribuido que manda el cliente (.xlsx, .xls o .csv)."],
        ["La plataforma busca sola las columnas de sucursal (Local, Sucursal, Tienda), producto (SKU, Cód. Prov., Material),"],
        ["unidades (Cantidad, Unidades) y grupo (Grupo, Bloque, Orden de carga). Las filas repetidas se suman."],
        ["Al terminar te muestra qué columnas leyó, para que lo confirmes."],
    ]
    for linea in filas_ayuda:
        ayuda.append(linea)
    ayuda.column_dimensions["A"].width = 112
    for fila in (1, 3, 9, 20, 24):
        ayuda[f"A{fila}"].font = Font(bold=True, size=13 if fila == 1 else 11)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_cubicaje.xlsx"'})


@router.post("/cubicaje-libre/importar")
def importar_carga(file: UploadFile = File(...), reemplazar: str = Form("si"),
                   modo: str = Form(""), caja_master: str = Form(""),
                   piso_pallet: str = Form(""), destino: str = Form(""),
                   s: Session = Depends(get_session)):
    """Carga masiva: lee el Excel y deja los productos listos para cubicar."""
    import json as _json

    from ... import cargas
    from ...models import Config
    try:
        lec = cargas.leer_carga(file.file.read(), file.filename or "", _conocidos(s))
    except ValueError as e:
        raise HTTPException(422, str(e)) from e

    lineas, sin_medidas, reparto = [], [], []
    errores = list(lec.errores)
    conocidos = _conocidos(s)
    leidas = []                                      # (fila, sku, qty, sucursal, grupo)
    for f in lec.filas:
        if f.sku.lower() not in conocidos:
            sin_medidas.append(f.sku)              # se queda en la carga: se avisa con sus unidades y se cubica cuando tenga medidas
        leidas.append((f.n, f.sku, f.qty, f.sucursal, f.grupo))

    # Con sucursales, el archivo es un reparto: una fila por sucursal y SKU. Las unidades
    # por SKU se suman (es el tope de carga) y el detalle va al motor como predistribuido.
    con_reparto = any(x[3] for x in leidas)
    por_sku: dict = {}
    grupos: list = []
    if con_reparto and any(x[4] for x in leidas):
        lec.notas.append("El archivo trae sucursal y grupo: se usó la sucursal y se ignoró el grupo.")
    for i, sku, qty, suc, grupo in leidas:
        if con_reparto and not suc:
            errores.append(f"fila {i}: falta la sucursal (el archivo trae reparto por sucursal)")
            continue
        if con_reparto:
            reparto.append({"sucursal": suc, "sku": sku, "qty": qty})
        elif grupo or any(x[4] for x in leidas):
            grupos.append({"grupo": grupo, "sku": sku, "qty": qty})
        l = por_sku.setdefault(sku, {"sku": sku, "qty": 0})
        l["qty"] += qty
    lineas = list(por_sku.values())

    if not lineas:
        detalle = "; ".join(errores[:3]) or ("SKU sin medidas: " + ", ".join(sin_medidas[:5])
                                             if sin_medidas else "")
        raise HTTPException(422, "El archivo no trae productos cubicables. " + detalle)

    c_prev = s.get(Config, "cubicaje_libre")
    anterior = _json.loads(c_prev.valor) if c_prev else {}
    previas = []
    if str(reemplazar).lower() not in ("si", "sí", "true", "1"):
        previas = anterior.get("lineas") or []
        reparto = (anterior.get("predistribuido") or []) + reparto
        grupos = (anterior.get("grupos") or []) + grupos
    juntas = {l["sku"]: dict(l) for l in previas}
    for l in lineas:
        if l["sku"] in juntas:
            juntas[l["sku"]]["qty"] = (juntas[l["sku"]].get("qty") or 0) + l["qty"]
            juntas[l["sku"]].update({k: v for k, v in l.items() if k != "qty"})
        else:
            juntas[l["sku"]] = l

    # el modo que se ve en pantalla manda: el guardado puede ir un cálculo atrás
    modo = (modo or str(anterior.get("modo") or "MDA")).strip().upper()
    caja_master = caja_master or anterior.get("caja_master") or ""
    aviso_modo = ""
    if reparto and "PREDISTRIBUIDO" not in modo:
        modo = "SDA PREDISTRIBUIDO" if "SDA" in modo else "MDA PREDISTRIBUIDO"
        if "SDA" in modo:
            caja_master = caja_master or "SIN CAJA MASTER"     # el modo exige indicarla
        aviso_modo = (f"El archivo trae reparto por sucursal: se cubicó en {modo}"
                      + (f", {caja_master.lower()}" if "SDA" in modo else "")
                      + ". Los modos de stock no separan por sucursal.")
    doc = cubicaje_libre({"lineas": list(juntas.values()), "predistribuido": reparto, "grupos": grupos,
                          "cliente": anterior.get("cliente") or "", "modo": modo,
                          "caja_master": caja_master, "vista": anterior.get("vista") or "rampla",
                          "piso_pallet": piso_pallet or anterior.get("piso_pallet") or "",
                          "destino": destino or anterior.get("destino") or ""}, s)
    if aviso_modo:
        doc["avisos"] = [aviso_modo, *(doc.get("avisos") or [])]
    doc["importadas"] = len(lineas)
    doc["lectura"] = lec.resumen()
    doc["sin_medidas_archivo"] = sin_medidas
    doc["errores_archivo"] = errores
    return doc
