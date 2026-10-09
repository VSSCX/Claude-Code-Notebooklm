"""Descarga del cubicaje en Excel."""

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session

from ...db import get_session
from ..comun import _clave_cubicaje

router = APIRouter()


@router.get("/cubicaje-libre/excel")
def excel_cubicaje_libre(s: Session = Depends(get_session)):
    """Descarga el cubicaje armado a mano como Excel (resumen por camión y detalle)."""
    import json as _json

    from ...models import Config
    c = s.get(Config, "cubicaje_libre")
    if c is None:
        raise HTTPException(404, "Todavía no hay un cubicaje armado.")
    return _excel_cubicaje(_json.loads(c.valor), "cubicaje")


@router.get("/cubicaje/{numero}/excel")
def excel_cubicaje_pedido(numero: str, s: Session = Depends(get_session)):
    import json as _json

    from ...models import Config
    c = s.get(Config, _clave_cubicaje(numero))
    if c is None:
        raise HTTPException(404, "Este pedido todavía no está cubicado.")
    return _excel_cubicaje(_json.loads(c.valor), f"cubicaje_{numero}")


def _excel_cubicaje(doc: dict, nombre: str):
    """Arma el archivo con las mismas columnas de las hojas 03 y 04 del Excel."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    cab = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1A2B4A")

    ws = wb.active
    ws.title = "Camiones"
    ws.append(["Camión", "Tipo", "Capacidad m3", "SKU", "Descripción", "Unidades",
               "Sucursal", "Pedido", "Tipo de carga"])
    for f in doc.get("filas", []):
        ws.append([f["camion"], f["tipo_camion"], f["cap_m3"], f["sku"], f["descripcion"],
                   f["unidades"], f.get("sucursal", ""), f.get("pedido", ""), f["tipo_carga"]])

    if doc.get("filas04"):
        ws2 = wb.create_sheet("Pallets")
        ws2.append(["Vehículo", "Tipo vehículo", "Pallet", "Sucursal", "SKU", "Descripción",
                    "Cajas master", "Bultos", "Unidades", "Tipo"])
        for f in doc["filas04"]:
            ws2.append([f["vehiculo"], f["tipo_vehiculo"], f["pallet"] or "", f.get("sucursal", ""),
                        f["sku"], f["descripcion"], f["cajas"], f["bultos"], f["unidades"],
                        f["tipo"]])

    for hoja in wb.worksheets:
        for celda in hoja[1]:
            celda.font = cab
            celda.fill = fondo
        for col in hoja.columns:
            ancho = max(len(str(x.value or "")) for x in col) + 2
            hoja.column_dimensions[col[0].column_letter].width = min(ancho, 42)
        hoja.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})
