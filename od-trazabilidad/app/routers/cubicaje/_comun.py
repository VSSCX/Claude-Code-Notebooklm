"""Piezas que comparten los endpoints de cubicaje: guardado y faltantes de medidas."""

from pathlib import Path


from ... import domain
from ...integrations import maestra
from ...integrations import medidas as med_mod
from ..comun import _commit  # noqa: F401  (los demás módulos del paquete lo toman de aquí)




def _faltantes_de(lineas: list[dict], conoce) -> list[dict]:
    """Productos de la carga que no están en la Base de Medidas: SKU, unidades (se suman si se repite)
    y el nombre que tienen en la maestra, para que el analista sepa cuáles son y los complete."""
    nombres = maestra.descripciones()
    out: dict[str, dict] = {}
    for l in lineas:
        sku = domain.norm_sku(l.get("sku"))
        if not sku or conoce(sku):
            continue
        f = out.setdefault(sku, {"sku": sku, "unidades": 0, "descripcion": nombres.get(sku, "")})
        f["unidades"] += float(l.get("qty") or 0)
    return sorted(out.values(), key=lambda f: (-f["unidades"], f["sku"]))


def _xlsx_faltantes(faltantes: list[dict], nombre: str):
    """Excel con el formato de la Base de Medidas, ya con los SKU que faltan: se completan las medidas
    y se importa en Configuración → Base de Medidas."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Base para carga"
    ws.append([*med_mod.CABECERA, "Unidades en la carga"])
    for f in faltantes:
        ws.append([f["sku"], f.get("descripcion") or "", 1, None, None, None, None, "Y", "N", "N", None, None, f["unidades"]])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A2B4A")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    for col, ancho in zip("ABCDEFGHIJKLM", (16, 38, 8, 11, 11, 11, 11, 8, 9, 8, 11, 11, 14)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ayuda = wb.create_sheet("Cómo se usa")
    for linea in ["Productos que no están en la Base de Medidas",
                  "",
                  "1. Completa Longitud, Anchura, Altura (cm) y Peso total (kg) de cada producto.",
                  "2. Piezas: unidades por caja (1 si es unitario). Apilar / Inclinar / Rotar: Y o N.",
                  "3. Máx Camión y Máx Pallet son opcionales.",
                  "4. Importa el archivo en Configuración → Base de Medidas. Agrega los productos nuevos y no toca los demás.",
                  "5. Vuelve al cubicador: se recalcula solo. La columna 'Unidades en la carga' es solo informativa."]:
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 100
    ayuda["A1"].font = Font(bold=True, size=13)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


def _lineas_sin_kits(lineas: list[tuple], kits: dict) -> list[dict]:
    """Lineas (sku, cantidad) con cada kit reemplazado por sus componentes: lo que necesita medidas son las cajas."""
    out = []
    for sku, qty in lineas:
        k = kits.get(domain.norm_sku(sku))
        if k is None:
            out.append({"sku": sku, "qty": qty})
        else:
            out += [{"sku": c, "qty": float(qty) * q} for c, q in k.componentes]
    return out


def _resumen_kits(kits: dict, lineas: list[tuple]) -> list[dict]:
    """Los kits de la carga, con sus componentes, para mostrarlos en el cubicaje."""
    out = {}
    for sku, qty in lineas:
        k = kits.get(domain.norm_sku(sku))
        if k is not None:
            f = out.setdefault(k.sku, {"sku": k.sku, "descripcion": k.desc, "unidades": 0,
                                       "componentes": [{"sku": c, "por_kit": q} for c, q in k.componentes]})
            f["unidades"] += int(qty)
    return list(out.values())


def _datos_visor(nombre: str) -> str:
    from ...config import settings
    from ...cubicaje.visor import leer_datos_visor
    return leer_datos_visor(settings.visores_dir, nombre)


def _visor_vivo() -> str:
    """Deja listo el visor de dirección fija (plantilla y librerías) y devuelve su URL."""
    from pathlib import Path as _Path

    from ...config import settings as _st
    from ...cubicaje.visor import asegurar_visor_vivo, preparar_carpeta
    carpeta = Path(_st.visores_dir)
    try:
        preparar_carpeta(carpeta, _st.visor_assets)
        return asegurar_visor_vivo(carpeta, _Path(_st.plantilla_visor).read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001 - sin visor la carga igual se calcula
        return ""
