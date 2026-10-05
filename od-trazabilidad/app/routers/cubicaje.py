"""Cubicaje, medidas y cubicador libre."""
import shutil
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import domain
from ..db import get_session
from ..integrations import acciones, bases, clientes as cli_mod, maestra, medidas as med_mod
from sqlalchemy import select

from ..models import Archivo, Entrega
from ..schemas import EntregaIn, PaqueteIn, PedidoIn


def _commit(s: Session):
    try:
        s.commit()
    except IntegrityError as e:
        s.rollback()
        raise HTTPException(409, "El registro choca con otro existente.") from e

from .comun import (AJUSTES_DEFECTO, CAMIONES_DEFECTO, CAMIONES_VISTA, _ajustes_cubicaje,
                    _aplicar_ajustes, _calefones_de, _clave_analisis, _clave_cubicaje)

router = APIRouter()

@router.get("/medidas")
def get_medidas(buscar: str = "", limite: int = 50, s: Session = Depends(get_session)):
    from sqlalchemy import or_, select as _sel
    from ..models import Medida
    q = _sel(Medida)
    if buscar:
        p = f"%{buscar.strip()}%"
        q = q.where(or_(Medida.sku.like(p), Medida.descripcion.like(p)))
    filas = s.scalars(q.order_by(Medida.sku).limit(max(1, min(limite, 500)))).all()
    return {**med_mod.estado(s), "filas": [
        {"sku": m.sku, "descripcion": m.descripcion, "piezas": m.piezas, "largo": m.largo,
         "ancho": m.ancho, "alto": m.alto, "peso": m.peso, "apilar": m.apilar,
         "inclinar": m.inclinar, "rotar": m.rotar, "max_camion": m.max_camion,
         "max_pallet": m.max_pallet, "actualizado": domain._iso(m.actualizado)} for m in filas]}


@router.post("/medidas/importar")
def importar_medidas(file: UploadFile = File(...), s: Session = Depends(get_session)):
    """Carga masiva desde Base de Medidas.xlsm: agrega los nuevos y actualiza los cambiados."""
    if Path(file.filename or "").suffix.lower() not in (".xlsm", ".xlsx"):
        raise HTTPException(422, "El archivo debe ser .xlsm o .xlsx")
    try:
        filas = med_mod.leer_archivo(file.file)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el archivo: {str(e)[:200]}") from e
    if not filas:
        raise HTTPException(422, "No se encontró la hoja 'Base para carga' con sus encabezados.")
    r = med_mod.importar(s, filas)
    _commit(s)
    return {**r.__dict__, **med_mod.estado(s)}


@router.put("/medidas/{sku}")
def put_medida(sku: str, body: dict, s: Session = Depends(get_session)):
    """Corrección puntual de un producto (por ejemplo, uno que falta y frena el cubicaje)."""
    fila = {"Grupo": sku, "Descripción": body.get("descripcion"), "Piezas": body.get("piezas"),
            "Longitud": body.get("largo"), "Anchura": body.get("ancho"), "Altura": body.get("alto"),
            "Peso total": body.get("peso"), "Apilar": body.get("apilar"),
            "Inclinar": body.get("inclinar"), "Rotar": body.get("rotar"),
            "Máx Camión": body.get("max_camion"), "Máx Pallet": body.get("max_pallet")}
    r = med_mod.importar(s, [fila])
    if r.ignorados:
        raise HTTPException(422, "Largo, ancho y alto tienen que ser mayores que cero.")
    _commit(s)
    return get_medidas(buscar=sku, s=s)


@router.get("/ajustes-cubicaje")
def get_ajustes_cubicaje(s: Session = Depends(get_session)):
    return _ajustes_cubicaje(s)


@router.put("/ajustes-cubicaje")
def put_ajustes_cubicaje(body: dict, s: Session = Depends(get_session)):
    orient = str(body.get("orientacion_pallet", "largo")).lower()
    if orient not in ("excel", "largo"):
        raise HTTPException(422, "Orientación no válida (largo o excel).")
    celda = int(body.get("celda_cm") or 1)
    if celda not in (1, 2):
        raise HTTPException(422, "La precisión tiene que ser 1 o 2 cm.")
    cap = str(body.get("capacidad_pallet", "geometria")).lower()
    if cap not in ("geometria", "tabla"):
        raise HTTPException(422, "Capacidad no válida (geometria o tabla).")
    datos = {"orientacion_pallet": orient, "celda_cm": celda, "capacidad_pallet": cap}
    domain.guardar_config(s, "ajustes_cubicaje", datos)
    _commit(s)
    return datos


def _clave_predist(pedido: str) -> str:
    return f"predist:{pedido}"


PREDIST_COLS = ["Sucursal", "SKU", "Unidades", "Unidades por bulto (opcional)"]


def _conocidos(s: Session) -> set[str]:
    return {m[0].lower() for m in med_mod.filas_para_cubicaje(s)}


def _fila_predist(n: int, suc, sku, cant, bulto=None) -> tuple[dict | None, str | None]:
    """Valida una fila del reparto. Mismas reglas para el texto pegado y para el Excel.
    Devuelve (fila, None), (None, error) o (None, None) si la fila se ignora."""
    suc = str(suc if suc is not None else "").strip().upper()
    sku = domain.norm_sku(sku)
    if not suc or not sku:
        return None, f"fila {n}: faltan datos (sucursal, SKU y unidades)"
    if isinstance(cant, (int, float)):
        unidades = float(cant)
    else:
        try:
            unidades = float(str(cant).strip().replace(".", "").replace(",", "."))
        except ValueError:
            return None, f"fila {n}: '{cant}' no es una cantidad"
    if unidades <= 0:
        return None, None
    por_bulto = 0
    if bulto not in (None, ""):
        try:
            por_bulto = int(float(str(bulto).replace(",", ".")))
        except ValueError:
            return None, f"fila {n}: '{bulto}' no es una cantidad por bulto"
    return {"sucursal": suc, "sku": sku, "unidades": unidades, "por_bulto": max(0, por_bulto)}, None


def _guardar_predist(s: Session, numero: str, filas: list, errores: list) -> dict:
    doc = {"pedido": numero, "filas": filas, "errores": errores,
           "sucursales": sorted({f["sucursal"] for f in filas}),
           "unidades": sum(f["unidades"] for f in filas)}
    domain.guardar_config(s, _clave_predist(numero), doc)
    _commit(s)
    return doc


# Va antes de /predistribuido/{numero}: si no, "plantilla" se toma como número de pedido.
@router.get("/predistribuido/plantilla")
def plantilla_predistribuido():
    """Excel para cargar el reparto por sucursal de un pedido predistribuido."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Predistribuido"
    ws.append(PREDIST_COLS)
    ws.append(["SUC-01", "900081624", 20, None])
    ws.append(["SUC-02", "900081624", 10, None])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A2B4A")
        celda.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCD", (18, 16, 12, 30)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ayuda = wb.create_sheet("Cómo se usa")
    for linea in [
        ["Reparto por sucursal (pedidos predistribuidos)"],
        [""],
        ["1. Una fila por cada sucursal y producto, en la hoja Predistribuido."],
        ["2. Si un producto va a varias sucursales, repítelo en una fila por sucursal."],
        ["3. Unidades: cuántas van a esa sucursal. Las filas en 0 o vacías se ignoran."],
        ["4. Unidades por bulto: solo si el cliente arma bultos propios (SDA Predistribuido)."],
        ["   Déjala vacía para usar la caja del producto."],
        [""],
        ["Al importar, este reparto reemplaza al que tenía el pedido."],
        ["Si la suma de una sucursal supera la carga del análisis, se carga hasta ese tope."],
        ["Las filas de ejemplo se pueden borrar."],
    ]:
        ayuda.append(linea)
    ayuda.column_dimensions["A"].width = 90
    ayuda["A1"].font = Font(bold=True, size=13)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_predistribuido.xlsx"'})


@router.post("/predistribuido/{numero}/importar")
def importar_predistribuido(numero: str, file: UploadFile = File(...),
                            s: Session = Depends(get_session)):
    """Carga masiva del reparto: la plantilla o el archivo del cliente. Reemplaza el reparto del pedido."""
    from .. import cargas
    try:
        lec = cargas.leer_carga(file.file.read(), file.filename or "", _conocidos(s),
                                obligatorios=("sucursal", "sku", "unidades"), hojas_preferidas=("predistribuido", "carga"),
                                orden_plantilla=("sucursal", "sku", "unidades", "bulto"))
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    filas = [{"sucursal": f.sucursal, "sku": f.sku, "unidades": f.qty, "por_bulto": f.bulto} for f in lec.filas]
    if not filas:
        raise HTTPException(422, "El archivo no trae filas de reparto válidas. " + "; ".join(lec.errores[:3]))
    doc = _guardar_predist(s, numero, filas, lec.errores)
    doc["lectura"] = lec.resumen()
    return doc


@router.get("/predistribuido/{numero}")
def get_predistribuido(numero: str, s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, _clave_predist(numero))
    return _json.loads(c.valor) if c else {"filas": []}


@router.put("/predistribuido/{numero}")
def put_predistribuido(numero: str, body: dict, s: Session = Depends(get_session)):
    """Tabla de reparto por sucursal: se pega desde el archivo del cliente."""
    import re as _re
    filas = []
    errores = []
    for i, linea in enumerate(str(body.get("texto", "")).splitlines(), 1):
        if not linea.strip():
            continue
        partes = [x.strip() for x in _re.split(r"\t|;|,", linea)]
        partes = [x for x in partes if x != ""]
        if len(partes) < 3:
            errores.append(f"línea {i}: faltan datos (sucursal, SKU y unidades)")
            continue
        if partes[0].lower() in ("sucursal", "suc"):    # encabezado pegado por error
            continue
        f, error = _fila_predist(i, partes[0], partes[1], partes[-1])
        if error:
            errores.append(error.replace("fila", "línea", 1))
        elif f:
            filas.append(f)
    return _guardar_predist(s, numero, filas, errores)


@router.get("/cubicaje/{numero}")
def get_cubicaje(numero: str, s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, _clave_cubicaje(numero))
    if c is None:
        raise HTTPException(404, "Este pedido todavía no está cubicado.")
    doc = _json.loads(c.valor)
    if doc.get("visor_vivo"):
        doc["visor_json"] = _datos_visor(f"pedido_{numero}_datos")
    return doc


def _faltantes_de(lineas: list[dict], conoce) -> list[dict]:
    """Productos de la carga que no están en la Base de Medidas: SKU, unidades (se suman si se repite)
    y el nombre que tienen en la maestra, para que el analista sepa cuáles son y los complete."""
    from ..integrations import maestra
    nombres = maestra.descripciones()
    out: dict[str, dict] = {}
    for l in lineas:
        sku = domain.norm_sku(l.get("sku"))
        if not sku or conoce(sku):
            continue
        f = out.setdefault(sku, {"sku": sku, "unidades": 0, "descripcion": nombres.get(sku, "")})
        f["unidades"] += float(l.get("qty") or 0)
    return sorted(out.values(), key=lambda f: (-f["unidades"], f["sku"]))


def _xlsx_faltantes(faltantes: list[dict], nombre: str):
    """Excel con el formato de la Base de Medidas, ya con los SKU que faltan: se completan las medidas
    y se importa en Configuración → Base de Medidas."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Base para carga"
    ws.append([*med_mod.CABECERA, "Unidades en la carga"])
    for f in faltantes:
        ws.append([f["sku"], f.get("descripcion") or "", 1, None, None, None, None, "Y", "N", "N", None, None, f["unidades"]])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A2B4A")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
    for col, ancho in zip("ABCDEFGHIJKLM", (16, 38, 8, 11, 11, 11, 11, 8, 9, 8, 11, 11, 14)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ayuda = wb.create_sheet("Cómo se usa")
    for linea in ["Productos que no están en la Base de Medidas",
                  "",
                  "1. Completa Longitud, Anchura, Altura (cm) y Peso total (kg) de cada producto.",
                  "2. Piezas: unidades por caja (1 si es unitario). Apilar / Inclinar / Rotar: Y o N.",
                  "3. Máx Camión y Máx Pallet son opcionales.",
                  "4. Importa el archivo en Configuración → Base de Medidas. Agrega los productos nuevos y no toca los demás.",
                  "5. Vuelve al cubicador: se recalcula solo. La columna 'Unidades en la carga' es solo informativa."]:
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 100
    ayuda["A1"].font = Font(bold=True, size=13)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


@router.get("/cubicaje-libre/faltantes.xlsx")
def faltantes_libre(s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, "cubicaje_libre")
    lineas = (_json.loads(c.valor).get("lineas") if c else None) or []
    conocidos = {m[0].lower() for m in med_mod.filas_para_cubicaje(s)}
    faltantes = _faltantes_de(lineas, lambda sku: sku.lower() in conocidos)
    if not faltantes:
        raise HTTPException(404, "Todos los productos de la carga tienen medidas.")
    return _xlsx_faltantes(faltantes, "medidas_faltantes.xlsx")


@router.get("/cubicaje/{numero}/faltantes.xlsx")
def faltantes_pedido(numero: str, s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    ca = s.get(Config, _clave_analisis(numero))
    if ca is None:
        raise HTTPException(404, "Este pedido todavía no tiene análisis.")
    filas = _json.loads(ca.valor)["resultado"]["filas"]
    conocidos = {m[0].lower() for m in med_mod.filas_para_cubicaje(s)}
    faltantes = _faltantes_de([{"sku": f["sku"], "qty": f["carga"]} for f in filas if f["carga"] > 0],
                              lambda sku: sku.lower() in conocidos)
    if not faltantes:
        raise HTTPException(404, "Todos los productos del pedido tienen medidas.")
    return _xlsx_faltantes(faltantes, f"medidas_faltantes_{numero}.xlsx")


@router.post("/cubicaje/{numero}")
def cubicar(numero: str, body: dict, s: Session = Depends(get_session)):
    return _cubicar(numero, body or {}, s)


def _cubicar(numero: str, body: dict, s: Session):
    """Cubica con la carga del análisis (incluidos los ajustes manuales)."""
    import copy as _copy
    import json as _json
    from datetime import date as _date
    from ..cubicaje.datos import cargar_cache_dims, leer_camiones
    from ..cubicaje.mda import Posicion
    from ..cubicaje.mda_predist import FilaPredist
    from ..cubicaje.motor import Entrada, ModoNoPortado, segmentar
    from ..cubicaje.visor import construir_json
    from ..integrations import base_medidas
    from ..models import Config

    ca = s.get(Config, _clave_analisis(numero))
    if ca is None:
        raise HTTPException(422, "Primero hay que analizar el pedido: la carga sale del análisis.")
    an = _json.loads(ca.valor)
    filas = an["resultado"]["filas"]
    if not filas:
        raise HTTPException(422, "El análisis no tiene productos.")

    cfg = s.get(Config, "app")
    cfg_val = _json.loads(cfg.valor) if cfg else {}
    camiones_cfg = body.get("camiones") or cfg_val.get("camiones") or CAMIONES_DEFECTO
    modo = str(body.get("modo") or an.get("modo_cubicaje") or "MDA").strip().upper()
    cliente = an.get("cliente", "")

    med_filas = med_mod.filas_para_cubicaje(s)          # la cargada en la plataforma
    origen_medidas = "plataforma"
    if not med_filas:                                    # si no hay, se lee el archivo de red
        med_filas = base_medidas.filas()
        origen_medidas = "archivo de red"
    if not med_filas:
        raise HTTPException(422, "No hay Base de Medidas: impórtala en la pestaña SAP y archivos.")

    ajustes = _aplicar_ajustes(s)
    regla = cli_mod.buscar(s, cliente)
    entrada = Entrada(
        cliente=cliente, modo=modo,
        posiciones=[Posicion(sku=f["sku"], desc=f["descripcion"], carga=f["carga"],
                             pedido=numero, fila=i + 5) for i, f in enumerate(filas)],
        pedidos=[numero], medidas=cargar_cache_dims(med_filas),
        camiones=leer_camiones(camiones_cfg),
        caja_master=str(body.get("caja_master") or (regla.caja_master if regla else "")),
        pallet=tuple(cli_mod.doc(regla)["pallet"]) if regla else None,
        hibrido=regla.hibrido if regla else None,
        orientacion_pallet=ajustes["orientacion_pallet"],
        capacidad_pallet=ajustes.get("capacidad_pallet", "geometria"),
        piso_pallet=str(body.get("piso_pallet") or ""),
        calefones=set(body.get("calefones") or _calefones_de(cliente, s)),
        predistribuido=[FilaPredist(sucursal=f["sucursal"], sku=f["sku"], unidades=f["unidades"],
                                    por_bulto=int(f.get("por_bulto") or 0))
                        for f in _json.loads(pre.valor)["filas"]] if (pre := s.get(Config, _clave_predist(numero))) else [])
    try:
        r = segmentar(entrada)
    except ModoNoPortado as e:
        raise HTTPException(422, str(e)) from e
    except ValueError as e:
        raise HTTPException(422, str(e)) from e

    doc = {"pedido": numero, "cliente": cliente, "modo": r.modo, "generado": _date.today().isoformat(),
           "caja_master": entrada.caja_master, "piso_pallet": entrada.piso_pallet,
           "pallet": list(entrada.pallet or (120.0, 100.0, 140.0)),
           "camiones": [{"numero": c.numero, "tipo": c.tipo, "L": c.L, "w": c.w, "h": c.h,
                         "vol_m3": c.vol_m3} for c in r.camiones],
           "filas": [f.__dict__ for f in r.filas03],
           "filas04": [f.__dict__ for f in r.filas04],
           "pallets_detalle": r.pallets,
           "avisos": r.avisos, "sin_medidas": r.sin_medidas, "no_encontrados": r.no_encontrados,
           "sin_ubicar": r.sin_ubicar, "unidades": r.unidades, "origen_medidas": origen_medidas,
           "ajustes": ajustes,
           "faltantes": _faltantes_de([{"sku": f["sku"], "qty": f["carga"]} for f in filas if f["carga"] > 0],
                                      lambda sku, _c=entrada.medidas: _c.get(sku.lower()) is not None)}
    domain.guardar_config(s, _clave_cubicaje(numero), doc)

    # Visor 3D: misma plantilla del Excel, servida desde la plataforma con sus librerías
    doc["visor"] = ""
    try:
        from pathlib import Path as _Path
        from ..config import settings as _st
        from ..cubicaje.visor import (asegurar_visor_vivo, guardar_datos_visor, html_visor,
                                      preparar_carpeta)
        carpeta = Path(_st.visores_dir)
        faltan = preparar_carpeta(carpeta, _st.visor_assets)
        if faltan:
            doc["avisos"] = list(doc["avisos"]) + [
                "Al visor le faltan librerías (" + ", ".join(faltan) + "). "
                "Revisa VISOR_ASSETS en el archivo .env."]
        plantilla = _Path(_st.plantilla_visor).read_text(encoding="utf-8")
        from ..cubicaje.mda import Camion as _Cam
        # Sin carga no se dibuja una plantilla en blanco: se muestra la rampla vacía
        cams = r.camiones or [_Cam(numero=1, tipo="Rampla 53", L=1540, w=245, h=230)]
        datos_visor = construir_json(r.placed, cams, es_sda=bool(r.pallets), pallets=r.pallets,
                                      cliente=cliente, modo=r.modo)
        html = html_visor(plantilla, datos_visor)
        for viejo in carpeta.glob(f"pedido_{numero}_*.html"):        # deja solo el último
            viejo.unlink(missing_ok=True)
        nombre_fs = f"pedido_{numero}_{uuid.uuid4().hex[:8]}.html"
        (carpeta / nombre_fs).write_text(html, encoding="utf-8")
        doc["visor"] = f"/visor/{nombre_fs}"
        # para la página: visor de dirección fija + los datos de este cálculo
        doc["visor_vivo"] = asegurar_visor_vivo(carpeta, plantilla)
        guardar_datos_visor(carpeta, f"pedido_{numero}_datos", datos_visor)
    except Exception as e:  # noqa: BLE001 - el cubicaje vale aunque el visor falle
        doc["avisos"] = list(doc["avisos"]) + [f"No se pudo generar el visor 3D: {str(e)[:150]}"]
    domain.guardar_config(s, _clave_cubicaje(numero), doc)      # sin visor_json: va en su archivo
    domain.anotar_flujo(s, numero, cubicaje=domain.resumen_cubicaje(doc))
    _commit(s)
    if doc.get("visor_vivo"):
        doc["visor_json"] = datos_visor
    return doc


@router.get("/medidas/sugerir")
def sugerir_medidas(q: str = "", s: Session = Depends(get_session)):
    """Autocompletar de SKU: devuelve el producto y, si existe, su caja master."""
    from sqlalchemy import or_, select as _sel
    from ..models import Medida
    t = str(q or "").strip()
    if len(t) < 2:
        return {"sugerencias": []}
    p = f"%{t}%"
    filas = s.scalars(_sel(Medida).where(or_(Medida.sku.like(p), Medida.descripcion.like(p)))
                      .order_by(Medida.sku).limit(15)).all()
    por_sku = {m.sku: m for m in filas}
    # se agregan las cajas master de los SKU encontrados aunque no coincidan con el texto
    cajas = s.scalars(_sel(Medida).where(Medida.sku.in_([f"C{m.sku}" for m in filas
                                                         if not m.sku.startswith("C")]))).all()
    out = []
    for m in list(filas) + list(cajas):
        if m.sku in [x["sku"] for x in out]:
            continue
        es_caja = m.sku.startswith("C") and m.sku[1:] in por_sku
        out.append({"sku": m.sku, "descripcion": m.descripcion, "piezas": m.piezas,
                    "caja_master": es_caja or m.piezas > 1,
                    "medidas": f"{m.largo:g} × {m.ancho:g} × {m.alto:g} cm",
                    "max_camion": m.max_camion, "max_pallet": m.max_pallet})
    return {"sugerencias": out[:20]}


@router.get("/cubicaje-libre")
def get_cubicaje_libre(s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, "cubicaje_libre")
    doc = _json.loads(c.valor) if c else {"lineas": [], "cliente": "", "modo": "MDA",
                                          "vista": "rampla"}
    doc["visor_vivo"] = _visor_vivo()        # la página abre el visor antes del primer cálculo
    if doc["visor_vivo"] and doc.get("lineas"):
        doc["visor_json"] = _datos_visor("libre_datos")
    return doc


def _datos_visor(nombre: str) -> str:
    from ..config import settings
    from ..cubicaje.visor import leer_datos_visor
    return leer_datos_visor(settings.visores_dir, nombre)


def _visor_vivo() -> str:
    """Deja listo el visor de dirección fija (plantilla y librerías) y devuelve su URL."""
    from pathlib import Path as _Path
    from ..config import settings as _st
    from ..cubicaje.visor import asegurar_visor_vivo, preparar_carpeta
    carpeta = Path(_st.visores_dir)
    try:
        preparar_carpeta(carpeta, _st.visor_assets)
        return asegurar_visor_vivo(carpeta, _Path(_st.plantilla_visor).read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001 - sin visor la carga igual se calcula
        return ""


@router.post("/cubicaje-libre/desde-pedido")
def cubicaje_desde_pedido(body: dict, s: Session = Depends(get_session)):
    """Trae al cubicador las líneas de un pedido ya analizado, para editarlas a mano."""
    import json as _json
    from ..models import Config
    numero = str(body.get("pedido", "")).strip()
    if not numero:
        raise HTTPException(422, "Falta el número de pedido.")
    ca = s.get(Config, _clave_analisis(numero))
    if ca is None:
        raise HTTPException(422, f"El pedido {numero} no está analizado. Analízalo en la pestaña "
                                 f"Pedidos y vuelve acá.")
    an = _json.loads(ca.valor)
    lineas = [{"sku": f["sku"], "qty": f["carga"]} for f in an["resultado"]["filas"]
              if (f.get("carga") or 0) > 0]
    if not lineas:
        raise HTTPException(422, "El análisis de ese pedido no tiene carga.")
    return cubicaje_libre({"lineas": lineas, "cliente": an.get("cliente", ""),
                           "modo": body.get("modo") or "MDA", "vista": body.get("vista") or "rampla",
                           "caja_master": body.get("caja_master") or "", "pedido": numero}, s)


def _reubicar(p, dx: float, dy: float, dz: float):
    """Copia una colocación moviéndola al origen del pallet."""
    from dataclasses import replace
    return replace(p, container=1, x=p.x - dx, y=p.y - dy, z=max(p.z - dz, 0.0))


def _merge_resultados(partes: list, pallet: list) -> dict:
    """Une varios cubicajes (los fijados a un camión y el resto) renumerando camiones."""
    camiones, filas, filas04, pallets, placed, avisos = [], [], [], [], [], []
    sin_medidas, sin_ubicar = [], {}
    for r in partes:
        off_cam = len(camiones)
        off_pal = len(pallets)
        for c in r.camiones:
            camiones.append({"numero": c.numero + off_cam, "tipo": c.tipo, "L": c.L, "w": c.w,
                             "h": c.h, "vol_m3": c.vol_m3})
        for f in r.filas03:
            d = dict(f.__dict__)
            d["camion"] += off_cam
            filas.append(d)
        for f in r.filas04:
            d = dict(f.__dict__)
            d["vehiculo"] += off_cam
            d["pallet"] = d["pallet"] + off_pal if d["pallet"] else 0
            filas04.append(d)
        for pl in r.pallets:
            d = dict(pl)
            d["numero"] += off_pal
            d["vehiculo"] += off_cam
            pallets.append(d)
        for p in r.placed:
            p.container += off_cam
            placed.append(p)
        avisos += r.avisos
        sin_medidas += r.sin_medidas
        sin_ubicar.update(r.sin_ubicar)
    return {"camiones": camiones, "filas": filas, "filas04": filas04, "pallets_detalle": pallets,
            "placed": placed, "avisos": avisos, "sin_medidas": sin_medidas,
            "sin_ubicar": sin_ubicar, "pallet": list(pallet),
            "unidades": sum(f["unidades"] for f in (filas04 or filas))}


def _reparto_por_grupos(grupos: list) -> list:
    """Cada grupo es un bloque, en orden numérico; lo que no trae grupo va al final."""
    from ..cubicaje.mda_predist import FilaPredist
    numeros = sorted({int(x.get("grupo") or 0) for x in grupos if int(x.get("grupo") or 0) > 0})
    ancho = len(str(numeros[-1])) if numeros else 1       # con ceros a la izquierda "GRUPO 02" ordena antes que "GRUPO 10"
    def etiqueta(g):
        return f"GRUPO {g:0{max(ancho, 2)}d}" if g else "SIN GRUPO"
    orden = sorted(grupos, key=lambda x: (int(x.get("grupo") or 0) or 10 ** 9))
    return [FilaPredist(sucursal=etiqueta(int(x.get("grupo") or 0)), sku=domain.norm_sku(x.get("sku")),
                        unidades=float(x.get("qty") or 0)) for x in orden]


def _avisos_reparto(lineas: list, pre: list) -> list[str]:
    """Reglas del predistribuido: la carga de cada SKU es el tope del reparto, y lo que no tiene sucursal no se carga."""
    carga: dict[str, float] = {}
    for l in lineas:
        sku = domain.norm_sku(l.get("sku"))
        carga[sku] = carga.get(sku, 0) + float(l.get("qty") or 0)
    repartido: dict[str, float] = {}
    for f in pre:
        repartido[f.sku] = repartido.get(f.sku, 0) + f.unidades
    sobran = {k: carga[k] - v for k, v in repartido.items() if k in carga and v < carga[k]}
    pasan = {k: v - carga.get(k, 0) for k, v in repartido.items() if v > carga.get(k, 0)}
    sin_reparto = {k: v for k, v in carga.items() if k not in repartido and v > 0}
    av = []
    if pasan:
        av.append("El reparto supera la carga en " + ", ".join(f"{k} (+{v:g})" for k, v in list(pasan.items())[:6])
                  + ": se recorta en el orden del archivo (las últimas sucursales quedan cortas).")
    if sobran or sin_reparto:
        resto = {**sobran, **sin_reparto}
        av.append("Sin sucursal asignada, no se cargan: " + ", ".join(f"{k} ({v:g} un.)" for k, v in list(resto.items())[:6])
                  + (" …" if len(resto) > 6 else ""))
    return av


@router.post("/cubicaje-libre")
def cubicaje_libre(body: dict, s: Session = Depends(get_session)):
    """Cubica una carga armada a mano: sin pedido, sin SAP.

    Con reparto por sucursal (predistribuido), las unidades por SKU son el tope de carga.
    """
    import copy as _copy
    import json as _json
    from datetime import date as _date
    from ..cubicaje.datos import cargar_cache_dims, leer_camiones
    from ..cubicaje.mda import Posicion
    from ..cubicaje.mda_predist import FilaPredist
    from ..cubicaje.motor import Entrada, ModoNoPortado, segmentar
    from ..cubicaje.visor import construir_json

    cliente = str(body.get("cliente") or "").strip().upper()
    from ..cubicaje.motor import resolver_destino
    piso_elegido = str(body.get("piso_pallet") or "")
    modo_elegido = str(body.get("modo") or "MDA").strip().upper()        # lo que eligió el usuario, antes de Destino y de la vista
    modo = resolver_destino(str(body.get("modo") or "MDA").strip().upper(), body.get("destino"))
    vista = str(body.get("vista") or "rampla").strip().lower()
    lineas = [x for x in (body.get("lineas") or []) if str(x.get("sku", "")).strip()]

    ajustes = _aplicar_ajustes(s)
    regla = cli_mod.buscar(s, cliente) if cliente else None
    pallet = tuple(cli_mod.doc(regla)["pallet"]) if regla else (120.0, 100.0, 140.0)
    med_filas = med_mod.filas_para_cubicaje(s)
    if not med_filas:
        raise HTTPException(422, "No hay Base de Medidas cargada: impórtala en SAP y archivos.")
    cache = cargar_cache_dims(med_filas)

    if vista == "pallet":
        # Un pallet solo: se cubica normal (en pallets) y después se muestra uno.
        # En MDA no existen pallets, así que para esta vista se usa el motor de pallets.
        body = {**body, "piso_pallet": ""}          # ver un pallet exige pallets: "a piso" no aplica
        if "SDA" not in modo:
            modo = "SDA PREDISTRIBUIDO" if "PREDISTRIBUIDO" in modo else "SDA STOCK"
            # los modos SDA exigen indicar caja master: si no viene, se asume sin caja
            if not str(body.get("caja_master") or "").strip():
                body = {**body, "caja_master": "SIN CAJA MASTER"}
        camiones = leer_camiones([list(CAMIONES_VISTA["rampla"])])
    elif vista == "camion50":
        camiones = leer_camiones([list(CAMIONES_VISTA["camion50"])])
    else:
        camiones = leer_camiones([list(CAMIONES_VISTA["rampla"])])

    posiciones = []
    desconocidos = []
    for i, l in enumerate(lineas):
        sku = domain.norm_sku(l.get("sku"))
        d = cache.get(sku.lower())
        if d is None:
            desconocidos.append(sku)
            continue
        posiciones.append(
            Posicion(sku=sku, desc=d.desc or sku, carga=float(l.get("qty") or 0),
                     pedido=str(body.get("pedido") or "LIBRE"), fila=i + 1))

    pre = [FilaPredist(sucursal=str(x.get("sucursal", "")).strip().upper(),
                       sku=domain.norm_sku(x.get("sku")), unidades=float(x.get("qty") or 0))
           for x in (body.get("predistribuido") or [])
           if cache.get(domain.norm_sku(x.get("sku")).lower()) is not None]      # sin medidas no se reparte

    # Grupos de carga: el 1 va al fondo y cada grupo empieza donde terminó el anterior. Es el mismo mecanismo
    # de bloques de un predistribuido, con el grupo en lugar de la sucursal.
    grupos = [x for x in (body.get("grupos") or []) if cache.get(domain.norm_sku(x.get("sku")).lower()) is not None]
    por_grupos = bool(grupos) and not pre and str(body.get("destino") or "").upper() != "STOCK"
    if por_grupos:
        pre = _reparto_por_grupos(grupos)
        if "PREDISTRIBUIDO" not in modo:
            modo = "SDA PREDISTRIBUIDO" if "SDA" in modo else "MDA PREDISTRIBUIDO"

    partes = []
    if posiciones:
        entrada = Entrada(cliente=cliente or "SIN CLIENTE", modo=modo, posiciones=posiciones,
                          pedidos=[str(body.get("pedido") or "LIBRE")], medidas=cache,
                          camiones=camiones,
                          # los modos que no son MDA exigen indicarla: sin dato se parte en "sin caja master"
                          caja_master=str(body.get("caja_master") or (regla.caja_master if regla else "")
                                          or ("SIN CAJA MASTER" if modo != "MDA" else "")),
                          piso_pallet=str(body.get("piso_pallet") or ""),
                          calefones=_calefones_de(cliente, s), predistribuido=pre, pallet=pallet,
                          hibrido=regla.hibrido if regla else None,
                          orientacion_pallet=ajustes["orientacion_pallet"],
                          capacidad_pallet=ajustes.get("capacidad_pallet", "geometria"))
        try:
            partes.append(segmentar(entrada))
        except (ModoNoPortado, ValueError) as e:
            raise HTTPException(422, str(e)) from e

    base = _merge_resultados(partes, list(pallet)) if partes else {
        "camiones": [], "filas": [], "filas04": [], "pallets_detalle": [], "placed": [],
        "avisos": [], "sin_medidas": [], "sin_ubicar": {}, "pallet": list(pallet), "unidades": 0}
    placed = base.pop("placed")
    from ..cubicaje.visor import PALETA, letra as _letra
    # La letra y el color son los mismos que dibuja el visor: se asignan por orden de carga
    orden, n = {}, 0
    skus_carga = {domain.norm_sku(x.get("sku")) for x in lineas}
    for pl in placed:
        # la caja master (C + sku) se muestra con la letra de su producto
        cod = pl.cod[1:] if pl.cod.startswith("C") and pl.cod[1:] in skus_carga else pl.cod
        if cod not in orden:
            orden[cod] = (_letra(n), PALETA[n % len(PALETA)])
            n += 1
    detalle = {}
    siguiente = len(orden)          # los productos sin colocar siguen la serie: no repiten letra ni color
    for l in lineas:
        sku = domain.norm_sku(l.get("sku"))
        d = cache.get(sku.lower())
        if sku not in orden:
            orden[sku] = (_letra(siguiente), PALETA[siguiente % len(PALETA)])
            siguiente += 1
        letra, color = orden[sku]
        if d is not None:
            dc = cache.get(("c" + sku).lower())
            # caja master: C + el SKU del producto suelto (o una "C…" con varias piezas)
            es_master = sku[:1] in ("C", "c") and (cache.get(sku[1:].lower()) is not None or int(d.piezas or 1) > 1)
            detalle[sku] = {"piezas": int(d.piezas or 1), "master": es_master,
                            "caja": int(dc.piezas or 1) if dc is not None else 0,
                            "descripcion": d.desc or sku,
                            "medidas": f"{d.L:g} × {d.w:g} × {d.h:g} cm",
                            "apilable": d.apilable, "peso": d.peso,
                            "letra": letra, "color": color}
    doc = {**base, "cliente": cliente, "modo": modo_elegido, "vista": vista, "lineas": lineas,
           "detalle_lineas": detalle,
           "caja_master": str(body.get("caja_master") or ""), "piso_pallet": piso_elegido,
           "predistribuido": body.get("predistribuido") or [], "pedido": body.get("pedido") or "",
           "destino": str(body.get("destino") or ""),
           "generado": _date.today().isoformat(), "ajustes": ajustes,
           "desconocidos": desconocidos, "modo_usado": modo,
           "faltantes": _faltantes_de(lineas, lambda sku: cache.get(sku.lower()) is not None)}

    # Las cajas master cuentan como un bulto con varias unidades: 10 cajas de 4 = 40 unidades en 10 bultos.
    # El motor coloca cajas; aquí las unidades de cada fila y el total se llevan a unidades de producto.
    por_caja = {k: v["piezas"] for k, v in detalle.items() if v.get("master") and v["piezas"] > 1}
    if por_caja:
        for lista in (doc.get("filas") or [], doc.get("filas04") or []):
            for f in lista:
                m = por_caja.get(str(f.get("sku")))
                if m:
                    f["bultos"] = f.get("bultos") or f["unidades"]
                    f["unidades"] = f["unidades"] * m
        doc["unidades"] = sum(f["unidades"] for f in (doc.get("filas04") or doc.get("filas") or []))
    doc["por_caja_master"] = por_caja
    if "PREDISTRIBUIDO" in modo and pre:
        doc["avisos"] = list(doc.get("avisos") or []) + _avisos_reparto(lineas, pre)
    if por_grupos:
        n_grupos = len({x.sucursal for x in pre})
        doc["avisos"] = [f"Cubicado en {n_grupos} grupos: el primero va al fondo y cada grupo empieza donde terminó el anterior."
                         + (" En pallets, cada grupo arma sus pallets." if "SDA" in modo and not piso_elegido == "PISO" else ""),
                         *(doc.get("avisos") or [])]
    doc["grupos"] = body.get("grupos") or []

    try:
        from pathlib import Path as _Path
        from ..config import settings as _st
        from ..cubicaje.mda import Camion as _Cam
        from ..cubicaje.sda import TARIMA
        from ..cubicaje.visor import (asegurar_visor_vivo, guardar_datos_visor, html_visor,
                                      preparar_carpeta)
        carpeta = Path(_st.visores_dir)
        preparar_carpeta(carpeta, _st.visor_assets)
        pallets_visor = doc["pallets_detalle"]
        if vista == "pallet":
            todos = sorted({p.pallet for p in placed if p.pallet})
            n_pal = int(body.get("pallet_n") or (todos[0] if todos else 0))
            if n_pal not in todos:
                n_pal = todos[0] if todos else 0
            geo = next((x for x in doc["pallets_detalle"] if x["numero"] == n_pal), None)
            doc["pallet_visto"] = n_pal
            doc["pallets_disponibles"] = todos
            doc["pallet_camion"] = bool(body.get("pallet_camion"))
            doc["pallet_vehiculo"] = (geo or {}).get("vehiculo", 0)       # camión donde va el pallet elegido
        if vista == "pallet" and not body.get("pallet_camion"):
            # Se aísla el pallet elegido: sus cajas se mueven al origen y se dibuja su tarima
            dx = geo["x"] if geo else 0.0
            dy = geo["y"] if geo else 0.0
            # Se mueven al origen en X e Y, pero la altura se conserva: las cajas van
            # apoyadas sobre la tarima, no atravesándola.
            placed = [_reubicar(p, dx, dy, 0.0) for p in placed if p.pallet == n_pal]
            cams = [_Cam(numero=1, tipo=f"Pallet {n_pal}", L=pallet[1], w=pallet[0],
                         h=pallet[2] + TARIMA)]
            pallets_visor = [{"numero": n_pal, "vehiculo": 1, "x": 0.0, "y": 0.0,
                              "dl": pallet[1], "dw": pallet[0],
                              "tipo": (geo or {}).get("tipo", ""),
                              "sucursal": (geo or {}).get("sucursal", "")}]
            if n_pal:
                doc["camiones"] = [{"numero": n_pal, "tipo": "Pallet", "L": pallet[1],
                                    "w": pallet[0], "h": pallet[2],
                                    "vol_m3": pallet[0] * pallet[1] * pallet[2] / 1_000_000}]
                dentro = [f for f in doc["filas04"] if f.get("pallet") == n_pal]
                vol = pallet[0] * pallet[1] * pallet[2] / 1_000_000 or 1
                usado = sum(p.volM3 for p in placed)      # lo que ocupan sus cajas
                doc["filas"] = [{**f, "camion": n_pal, "tipo_camion": "Pallet",
                                 "cap_m3": round(vol, 2), "ocup_linea": usado / vol,
                                 "ocup_acum": usado / vol,
                                 "libre_m3": 0, "pedido": "", "pedidos_camion": "",
                                 "tipo_carga": f.get("tipo", ""), "descripcion": f["descripcion"],
                                 "sku": f["sku"], "unidades": f["unidades"],
                                 "fila_origen": 0, "sucursal": f.get("sucursal", "")}
                                for f in dentro]
        else:
            cams = ([_Cam(numero=c["numero"], tipo=c["tipo"], L=c["L"], w=c["w"], h=c["h"])
                     for c in doc["camiones"]] or
                    [_Cam(numero=1, tipo=camiones[0].tipo, L=camiones[0].L, w=camiones[0].w,
                          h=camiones[0].h)])
        plantilla = _Path(_st.plantilla_visor).read_text(encoding="utf-8")
        datos_visor = construir_json(placed, cams, es_sda=bool(pallets_visor), pallets=pallets_visor, modo=modo,
                                    por_caja=por_caja)
        html = html_visor(plantilla, datos_visor)
        for viejo in carpeta.glob("libre_*.html"):
            viejo.unlink(missing_ok=True)
        nombre_fs = f"libre_{uuid.uuid4().hex[:8]}.html"
        (carpeta / nombre_fs).write_text(html, encoding="utf-8")
        doc["visor"] = f"/visor/{nombre_fs}"
        doc["visor_vivo"] = asegurar_visor_vivo(carpeta, plantilla)
        guardar_datos_visor(carpeta, "libre_datos", datos_visor)
    except Exception as e:  # noqa: BLE001
        doc["visor"] = ""
        doc["avisos"] = list(doc["avisos"]) + [f"No se pudo generar el visor: {str(e)[:150]}"]

    domain.guardar_config(s, "cubicaje_libre", doc)             # sin visor_json: va en su archivo
    _commit(s)
    if doc.get("visor_vivo"):
        doc["visor_json"] = datos_visor
    return doc


PLANTILLA_COLS = ["SKU", "Unidades", "Sucursal (opcional)", "Grupo (opcional)"]


@router.get("/cubicaje-libre/plantilla")
def plantilla_carga(s: Session = Depends(get_session)):
    """Excel para armar una carga fuera de la plataforma y después importarla."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Carga"
    ws.append(PLANTILLA_COLS)
    ejemplos = med_mod.filas_para_cubicaje(s)[:3]
    for fila in ejemplos:
        ws.append([fila[0], 10, "", ""])
    if not ejemplos:
        ws.append(["900081624", 10, ""])
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A2B4A")
        celda.alignment = Alignment(horizontal="center")
    for col, ancho in zip("ABCD", (18, 12, 22, 18)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ayuda = wb.create_sheet("Cómo se usa")
    for linea in [
        ["Carga masiva para el cubicador"],
        [""],
        ["Con esta plantilla"],
        ["1. Un producto por fila en la hoja Carga: SKU y Unidades. Las filas en 0 o vacías se ignoran."],
        ["2. SKU: el código del producto. Si tiene caja master, usa C delante (C900081624)."],
        ["3. Sucursal: solo si el cliente es predistribuido. Repite el producto en una fila por sucursal."],
        ["   Con sucursales, el cubicador pasa solo a MDA o SDA predistribuido."],
        ["4. Grupo: 1, 2, 3… para cargar por tandas. El grupo 1 va al fondo del camión y cada grupo"],
        ["   empieza donde terminó el anterior. Lo que no tenga grupo se carga al final."],
        ["   Si usas sucursal, el grupo no se aplica (manda la sucursal)."],
        [""],
        ["Con el archivo del cliente, sin tocarlo"],
        ["Puedes subir directo el predistribuido que manda el cliente (.xlsx, .xls o .csv)."],
        ["La plataforma busca sola las columnas de sucursal (Local, Sucursal, Tienda), producto"],
        ["(SKU, Cód. Prov., Material) y unidades (Cantidad, Unidades). Si el archivo trae dos códigos,"],
        ["usa el que coincide con la Base de Medidas. Las filas repetidas se suman por sucursal y producto."],
        ["Al terminar te muestra qué columnas leyó, para que lo confirmes."],
        [""],
        ["Los productos tienen que estar en la Base de Medidas (Configuración). Los que no estén"],
        ["no se cubican: se avisa su SKU y sus unidades para que los agregues."],
    ]:
        ayuda.append(linea)
    ayuda.column_dimensions["A"].width = 95
    ayuda["A1"].font = Font(bold=True, size=13)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_cubicaje.xlsx"'})


@router.post("/cubicaje-libre/importar")
def importar_carga(file: UploadFile = File(...), reemplazar: str = Form("si"),
                   modo: str = Form(""), caja_master: str = Form(""),
                   piso_pallet: str = Form(""), destino: str = Form(""),
                   s: Session = Depends(get_session)):
    """Carga masiva: lee el Excel y deja los productos listos para cubicar."""
    import json as _json

    from .. import cargas
    from ..models import Config
    try:
        lec = cargas.leer_carga(file.file.read(), file.filename or "", _conocidos(s))
    except ValueError as e:
        raise HTTPException(422, str(e)) from e

    lineas, sin_medidas, reparto = [], [], []
    errores = list(lec.errores)
    conocidos = _conocidos(s)
    leidas = []                                      # (fila, sku, qty, sucursal, grupo)
    for f in lec.filas:
        if f.sku.lower() not in conocidos:
            sin_medidas.append(f.sku)              # se queda en la carga: se avisa con sus unidades y se cubica cuando tenga medidas
        leidas.append((f.n, f.sku, f.qty, f.sucursal, f.grupo))

    # Con sucursales, el archivo es un reparto: una fila por sucursal y SKU. Las unidades
    # por SKU se suman (es el tope de carga) y el detalle va al motor como predistribuido.
    con_reparto = any(x[3] for x in leidas)
    por_sku: dict = {}
    grupos: list = []
    if con_reparto and any(x[4] for x in leidas):
        lec.notas.append("El archivo trae sucursal y grupo: se usó la sucursal y se ignoró el grupo.")
    for i, sku, qty, suc, grupo in leidas:
        if con_reparto and not suc:
            errores.append(f"fila {i}: falta la sucursal (el archivo trae reparto por sucursal)")
            continue
        if con_reparto:
            reparto.append({"sucursal": suc, "sku": sku, "qty": qty})
        elif grupo or any(x[4] for x in leidas):
            grupos.append({"grupo": grupo, "sku": sku, "qty": qty})
        l = por_sku.setdefault(sku, {"sku": sku, "qty": 0})
        l["qty"] += qty
    lineas = list(por_sku.values())

    if not lineas:
        detalle = "; ".join(errores[:3]) or ("SKU sin medidas: " + ", ".join(sin_medidas[:5])
                                             if sin_medidas else "")
        raise HTTPException(422, "El archivo no trae productos cubicables. " + detalle)

    c_prev = s.get(Config, "cubicaje_libre")
    anterior = _json.loads(c_prev.valor) if c_prev else {}
    previas = []
    if str(reemplazar).lower() not in ("si", "sí", "true", "1"):
        previas = anterior.get("lineas") or []
        reparto = (anterior.get("predistribuido") or []) + reparto
        grupos = (anterior.get("grupos") or []) + grupos
    juntas = {l["sku"]: dict(l) for l in previas}
    for l in lineas:
        if l["sku"] in juntas:
            juntas[l["sku"]]["qty"] = (juntas[l["sku"]].get("qty") or 0) + l["qty"]
            juntas[l["sku"]].update({k: v for k, v in l.items() if k != "qty"})
        else:
            juntas[l["sku"]] = l

    # el modo que se ve en pantalla manda: el guardado puede ir un cálculo atrás
    modo = (modo or str(anterior.get("modo") or "MDA")).strip().upper()
    caja_master = caja_master or anterior.get("caja_master") or ""
    aviso_modo = ""
    if reparto and "PREDISTRIBUIDO" not in modo:
        modo = "SDA PREDISTRIBUIDO" if "SDA" in modo else "MDA PREDISTRIBUIDO"
        if "SDA" in modo:
            caja_master = caja_master or "SIN CAJA MASTER"     # el modo exige indicarla
        aviso_modo = (f"El archivo trae reparto por sucursal: se cubicó en {modo}"
                      + (f", {caja_master.lower()}" if "SDA" in modo else "")
                      + ". Los modos de stock no separan por sucursal.")
    doc = cubicaje_libre({"lineas": list(juntas.values()), "predistribuido": reparto, "grupos": grupos,
                          "cliente": anterior.get("cliente") or "", "modo": modo,
                          "caja_master": caja_master, "vista": anterior.get("vista") or "rampla",
                          "piso_pallet": piso_pallet or anterior.get("piso_pallet") or "",
                          "destino": destino or anterior.get("destino") or ""}, s)
    if aviso_modo:
        doc["avisos"] = [aviso_modo, *(doc.get("avisos") or [])]
    doc["importadas"] = len(lineas)
    doc["lectura"] = lec.resumen()
    doc["sin_medidas_archivo"] = sin_medidas
    doc["errores_archivo"] = errores
    return doc


@router.get("/cubicaje-libre/excel")
def excel_cubicaje_libre(s: Session = Depends(get_session)):
    """Descarga el cubicaje armado a mano como Excel (resumen por camión y detalle)."""
    import json as _json
    from ..models import Config
    c = s.get(Config, "cubicaje_libre")
    if c is None:
        raise HTTPException(404, "Todavía no hay un cubicaje armado.")
    return _excel_cubicaje(_json.loads(c.valor), "cubicaje")


@router.get("/cubicaje/{numero}/excel")
def excel_cubicaje_pedido(numero: str, s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, _clave_cubicaje(numero))
    if c is None:
        raise HTTPException(404, "Este pedido todavía no está cubicado.")
    return _excel_cubicaje(_json.loads(c.valor), f"cubicaje_{numero}")


def _excel_cubicaje(doc: dict, nombre: str):
    """Arma el archivo con las mismas columnas de las hojas 03 y 04 del Excel."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    cab = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1A2B4A")

    ws = wb.active
    ws.title = "Camiones"
    ws.append(["Camión", "Tipo", "Capacidad m3", "SKU", "Descripción", "Unidades",
               "Sucursal", "Pedido", "Tipo de carga"])
    for f in doc.get("filas", []):
        ws.append([f["camion"], f["tipo_camion"], f["cap_m3"], f["sku"], f["descripcion"],
                   f["unidades"], f.get("sucursal", ""), f.get("pedido", ""), f["tipo_carga"]])

    if doc.get("filas04"):
        ws2 = wb.create_sheet("Pallets")
        ws2.append(["Vehículo", "Tipo vehículo", "Pallet", "Sucursal", "SKU", "Descripción",
                    "Cajas master", "Bultos", "Unidades", "Tipo"])
        for f in doc["filas04"]:
            ws2.append([f["vehiculo"], f["tipo_vehiculo"], f["pallet"] or "", f.get("sucursal", ""),
                        f["sku"], f["descripcion"], f["cajas"], f["bultos"], f["unidades"],
                        f["tipo"]])

    for hoja in wb.worksheets:
        for celda in hoja[1]:
            celda.font = cab
            celda.fill = fondo
        for col in hoja.columns:
            ancho = max(len(str(x.value or "")) for x in col) + 2
            hoja.column_dimensions[col[0].column_letter].width = min(ancho, 42)
        hoja.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})
