"""SAP: leer, analizar, crear y borrar."""
import math

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session

from .. import domain
from ..db import get_session
from ..integrations import acciones, bases
from sqlalchemy import select

from ..models import Entrega
from ..schemas import PaqueteIn
from .comun import _clave_analisis, _clave_cubicaje, _commit, _sop_de, solicitantes_del_grupo






router = APIRouter()

@router.post("/paquetes")
def post_paquete(body: PaqueteIn, s: Session = Depends(get_session)):
    """Recibe el paquete que envía TrazWeb.bas al terminar cada macro."""
    res = domain.cargar_paquete(s, body)
    _commit(s)
    return res


@router.post("/sap/leer_pedido", status_code=202)
def sap_leer_pedido(body: dict):
    """Lee un pedido de VL01N con Python (sin Excel). Solo lectura."""
    from datetime import date as _date
    import re as _re
    from ..integrations import sap
    pedido = str(body.get("pedido", "")).strip()
    puesto = str(body.get("puesto", "")).strip().upper()
    fecha = str(body.get("fecha", "")).strip()
    cliente = str(body.get("cliente", "")).strip()
    if not _re.fullmatch(r"\d{4,12}", pedido):
        raise HTTPException(422, "N° de pedido inválido.")
    if not _re.fullmatch(r"[A-Z0-9]{2,6}", puesto):
        raise HTTPException(422, "Puesto de expedición inválido (ejemplo: PN01).")
    try:
        f = _date.fromisoformat(fecha) if fecha else _date.today()
    except ValueError as e:
        raise HTTPException(422, "Fecha inválida.") from e

    def leer(avance):
        avance("Leyendo VL01N")
        lectura = sap.leer_pedido(pedido, puesto, f.strftime("%d.%m.%Y"))
        if not lectura.posiciones:
            raise RuntimeError(lectura.aviso or "El pedido no devolvió posiciones.")
        posiciones = [{"sku": x.sku, "qty_entrega": x.qty_entrega, "qty_pendiente": x.qty_pendiente}
                      for x in lectura.posiciones]
        from ..db import SessionLocal
        with SessionLocal() as ses:
            res = domain.cargar_lectura_sap(ses, pedido, cliente, posiciones)
            ses.commit()
        res["aviso"] = lectura.aviso
        return res

    try:
        return acciones.lanzar_python("sap_leer_pedido", "Leer pedido desde SAP (directo)",
                                      [pedido], leer)
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e


def _causa(e: Exception) -> str:
    """El motivo de un error técnico, en una línea y sin el ruido de los controladores ODBC."""
    txt = " ".join(str(e).split())
    return txt[:220] + ("…" if len(txt) > 220 else "")


PASOS_OPCIONALES = {"zsd": "Qty en entrega (ZSD001_03)", "mmbe": "Stock (MMBE)"}
CLAVE_AJUSTES_ANALISIS = "ajustes_analisis"


def _ajustes_analisis(s: Session) -> dict:
    """Qué pasos opcionales del análisis están activos en este equipo. Por defecto, todos."""
    import json as _json

    from ..models import Config
    c = s.get(Config, CLAVE_AJUSTES_ANALISIS)
    guardado = _json.loads(c.valor) if c else {}
    return {k: bool(guardado.get(k, True)) for k in PASOS_OPCIONALES}


@router.get("/ajustes-analisis")
def get_ajustes_analisis(s: Session = Depends(get_session)):
    return _ajustes_analisis(s)


@router.put("/ajustes-analisis")
def put_ajustes_analisis(body: dict, s: Session = Depends(get_session)):
    """Activa o desactiva pasos opcionales del análisis (p. ej. si ZSD001_03 falla en el equipo de un analista).
    Leer el pedido en VL01N no se puede desactivar: sin él no hay análisis."""
    actual = _ajustes_analisis(s)
    for k in PASOS_OPCIONALES:
        if k in body:
            actual[k] = bool(body[k])
    domain.guardar_config(s, CLAVE_AJUSTES_ANALISIS, actual)
    _commit(s)
    return actual


def _firma_analisis(doc: dict) -> str:
    """Lo que de verdad cambia un análisis al actualizar (sin las horas): sirve para escribir solo cuando hay algo nuevo."""
    import json as _json
    return _json.dumps([doc.get(k) for k in ("plan", "disponibilidad", "control_calidad", "posiciones", "completos", "resultado", "sin_sql")],
                       sort_keys=True, default=str)


def refrescar_analisis(s: Session, numero: str):
    """Actualiza lo que se mueve solo y no necesita SAP: plan SOP, disponibilidad y control de calidad (SQL Server), y lo que
    ya tiene entrega en la plataforma. El stock (MMBE) y la Qty en entrega de SAP quedan como estaban en el análisis.
    Devuelve (doc, ok, motivo) o None si el pedido no tiene análisis. Si SQL no responde, se conservan los datos anteriores."""
    import json as _json
    from datetime import datetime

    from ..analisis import calcular
    from ..models import Config
    c = s.get(Config, _clave_analisis(numero))
    if c is None:
        return None
    doc = _json.loads(c.valor)
    crudas = doc.get("posiciones_sap") or doc["posiciones"]
    antes = _firma_analisis(doc)
    skus = {domain.norm_sku(p["sku"]) for p in crudas}
    sin_sql = dict(doc.get("sin_sql") or {})
    ok, motivo = True, ""
    try:
        plan = bases.plan_sop(doc["grupo_sop"], fresco=True)
        doc["plan"] = {k: plan[k] for k in skus if k in plan}
        sin_sql.pop("plan", None)
        doc["sql_actualizado"] = datetime.now().astimezone().isoformat(timespec="minutes")
        try:
            doc["disponibilidad"] = {k: v for k, v in bases.disponibilidad(fresco=True).items() if k in skus}
            sin_sql.pop("disponibilidad", None)
        except Exception as e:  # noqa: BLE001
            ok, motivo = False, _causa(e)
        try:
            cc = bases.control_calidad(fresco=True)
            doc["control_calidad"] = {k: cc[k] for k in skus if k in cc}
            sin_sql.pop("control_calidad", None)
        except Exception as e:  # noqa: BLE001
            ok, motivo = False, motivo or _causa(e)
    except Exception as e:  # noqa: BLE001 - sin conexión: se siguen mostrando los datos de antes
        ok, motivo = False, _causa(e)
    # lo que falta de verdad según las entregas de hoy (el pedido pudo tener entregas nuevas desde el análisis)
    doc["posiciones_sap"] = crudas
    doc["posiciones"], doc["completos"] = domain.netear_posiciones(s, numero, crudas)
    doc["sin_sql"] = sin_sql
    doc["sql_estado"] = {"ok": ok, "motivo": motivo, "intento": datetime.now().astimezone().isoformat(timespec="minutes")}
    doc["resultado"] = calcular(doc["posiciones"], doc["plan"], doc["en_entrega"], doc["medidas"],
                                doc["disponibilidad"], doc["stock"], doc["ajustes"],
                                sin_plan=bool(sin_sql.get("plan")), control_calidad=doc.get("control_calidad"))
    if _firma_analisis(doc) != antes:                  # sin cambios no se escribe: no se mueve la versión ni se recargan los demás
        domain.guardar_config(s, _clave_analisis(numero), doc)
        domain.anotar_flujo(s, numero, analisis=domain.resumen_analisis(doc))
        _commit(s)
    return doc, ok, motivo


@router.post("/analisis/{numero}/actualizar")
def actualizar_analisis(numero: str, s: Session = Depends(get_session)):
    """Pone al día el plan SOP, la disponibilidad, el control de calidad y las entregas sin volver a entrar a SAP."""
    r = refrescar_analisis(s, numero)
    if r is None:
        raise HTTPException(404, "Este pedido todavía no tiene análisis.")
    return r[0]


@router.post("/analisis/{numero}", status_code=202)
def analizar(numero: str, body: dict):
    import json as _json
    import re as _re
    from datetime import date as _date, datetime
    from ..analisis import calcular, en_entrega_por_modelo, oc_de_zsd
    from ..db import SessionLocal
    from ..integrations import base_medidas, kits as kits_mod, sap
    from ..models import Config, Medida

    puesto = str(body.get("puesto", "")).strip().upper()
    cliente = str(body.get("cliente", "")).strip().upper()
    if not _re.fullmatch(r"\d{4,12}", numero):
        raise HTTPException(422, "N° de pedido inválido.")
    if not _re.fullmatch(r"[A-Z0-9]{2,6}", puesto):
        raise HTTPException(422, "Puesto de expedición inválido (ejemplo: PN01).")
    if not cliente:
        raise HTTPException(422, "Falta el cliente.")
    try:
        f = _date.fromisoformat(str(body.get("fecha") or "")) if body.get("fecha") else _date.today()
    except ValueError as e:
        raise HTTPException(422, "Fecha inválida.") from e

    def correr(avance):
        # Primero lo que puede fallar sin tocar SAP (cliente, SQL Server, Base de Medidas): antes todo eso se
        # descubría DESPUÉS de leer VL01N y exportar ZSD001_03, y el análisis moría justo al terminar la exportación.
        avance("0/4 Revisando conexiones")
        with SessionLocal() as ses:
            pasos = _ajustes_analisis(ses)
        # la carpeta de exportación y la ruta de publicación solo importan si se va a exportar ZSD001_03
        carpeta = sap.carpeta_export() if pasos["zsd"] else ""     # la misma de siempre, y comprobada antes de abrir SAP
        aviso_pub = sap.revisar_publicacion() if pasos["zsd"] else ""
        with SessionLocal() as ses:
            grupo, codigo = _sop_de(cliente, ses)
            solicitantes = solicitantes_del_grupo(cliente, ses)
            kits_def = kits_mod.definiciones(ses)
            # la Base de Medidas cargada en la plataforma manda; el archivo de red solo si no hay ninguna
            filas_med = ses.execute(select(Medida.sku, Medida.descripcion, Medida.max_camion)).all()
        if not codigo:
            raise RuntimeError(f"El cliente {cliente} no tiene código de solicitante: agrégalo en Configuración → Clientes.")
        # Sin conexión a SQL Server el análisis sigue con SAP: lo que se pierde es el cruce con el plan SOP (y la
        # disponibilidad), y queda avisado en el resultado en vez de detener todo.
        sin_sql: dict = {}
        try:
            plan = bases.plan_sop(grupo, fresco=True)           # el análisis siempre pregunta de nuevo
        except Exception as e:  # noqa: BLE001
            plan = None
            sin_sql = {"plan": True, "motivo": _causa(e)}
        try:
            disp = bases.disponibilidad(fresco=True)
        except Exception as e:  # noqa: BLE001
            disp = {}
            sin_sql["disponibilidad"] = True
            sin_sql.setdefault("motivo", _causa(e))
        calidad: dict = {}
        if sin_sql.get("plan") and sin_sql.get("disponibilidad"):
            sin_sql["control_calidad"] = True                    # el servidor no responde: no se espera otra vez
        else:
            try:
                calidad = bases.control_calidad(fresco=True)
            except Exception as e:  # noqa: BLE001
                sin_sql["control_calidad"] = True
                sin_sql.setdefault("motivo", _causa(e))
        if sin_sql.get("plan") or sin_sql.get("disponibilidad"):
            avance("Sin conexión con SQL Server: se sigue con SAP, sin el plan SOP")
        if filas_med:
            med = {sku: {"desc": d, "max_camion": mc} for sku, d, mc in filas_med}
        else:
            try:
                med = base_medidas.medidas()
            except Exception as e:  # noqa: BLE001
                raise RuntimeError("No hay Base de Medidas cargada y no se pudo abrir la del servidor de archivos. "
                                   "Cárgala en Configuración → Base de Medidas. Detalle: " + _causa(e)) from e

        for k in kits_def.values():                  # un kit no está en la Base de Medidas: se conoce por su listado
            med.setdefault(k.sku, {"desc": f"{k.sku} {k.desc}".strip(), "max_camion": 0})

        avance("1/4 Leyendo el pedido en VL01N")
        lectura = sap.leer_pedido(numero, puesto, f.strftime("%d.%m.%Y"))
        if not lectura.posiciones:
            raise RuntimeError(lectura.aviso or "El pedido no devolvió posiciones.")
        posiciones_sap = [{"sku": domain.norm_sku(x.sku), "qty_entrega": x.qty_entrega,
                           "qty_pendiente": x.qty_pendiente} for x in lectura.posiciones]
        # lo que ya tiene entrega en la plataforma se resta: el análisis marca solo lo que falta
        with SessionLocal() as ses:
            posiciones, completos = domain.netear_posiciones(ses, numero, posiciones_sap)

        if lectura.aviso:
            avance("Aviso de SAP: " + lectura.aviso[:80])
        avisos_zsd: list[str] = []
        publicado: dict = {}
        if pasos["zsd"]:
            avance("2/4 Consultando Qty en entrega en ZSD001_03" + (
                f" ({len(solicitantes)} solicitantes del grupo {grupo})" if len(solicitantes) > 1 else ""))
            # un archivo distinto por análisis: no choca con uno abierto en Excel ni con otro analista
            nombre_zsd = f"Qty En Entrega {numero} {datetime.now():%Y%m%d-%H%M%S}.xlsx"
            filas_zsd = sap.zsd001_03(solicitantes if len(solicitantes) > 1 else codigo, [p["sku"] for p in posiciones_sap],
                                      str(carpeta), nombre_zsd, avisar=avisos_zsd.append, publicado=publicado)
        else:
            avance("2/4 Qty en entrega omitida (desactivada en Configuración)")
            filas_zsd = []

        avance("3/4 Calculando saldos y alertas")
        # OC del pedido: Pedidos Ingresados (SQL) y, si no aparece o no hay conexión, el mismo reporte ZSD001_03
        oc, oc_origen = ("" if sin_sql.get("plan") and sin_sql.get("disponibilidad") else bases.oc_de(numero)), "Pedidos Ingresados"
        if not oc:
            oc, oc_origen = oc_de_zsd(filas_zsd, numero), "ZSD001_03"
        if not oc:
            oc_origen = ""
        en_ent = en_entrega_por_modelo(filas_zsd)
        sin_plan = plan is None
        plan = plan or {}
        res = calcular(posiciones, plan, en_ent, med, disp, sin_plan=sin_plan, control_calidad=calidad)

        stock = {}
        if not pasos["mmbe"]:
            avance("4/4 Stock omitido (MMBE desactivado en Configuración)")
        elif res["alertadas"]:
            stock = sap.mmbe(res["alertadas"], lambda t: avance("4/4 " + t))
            res = calcular(posiciones, plan, en_ent, med, disp, stock, sin_plan=sin_plan, control_calidad=calidad)
        else:
            avance("4/4 Sin alertas: no se consulta MMBE")

        skus = [p["sku"] for p in posiciones_sap]
        # Se guardan los datos usados, para recalcular ajustes sin volver a SAP
        doc = {"pedido": numero, "cliente": cliente, "grupo_sop": grupo, "puesto": puesto,
               "fecha": f.isoformat(), "posiciones": posiciones, "posiciones_sap": posiciones_sap, "completos": completos,
               "sql_actualizado": "" if sin_sql.get("plan") else datetime.now().astimezone().isoformat(timespec="minutes"),
               "en_entrega": en_ent, "solicitantes": solicitantes,
               "plan": {k: plan[k] for k in skus if k in plan},
               "medidas": {k: med[k] for k in skus if k in med},
               "disponibilidad": {k: disp[k] for k in skus if k in disp},
               "control_calidad": {k: calidad[k] for k in skus if k in calidad},
               "stock": stock, "ajustes": {}, "resultado": res, "oc_sap": oc, "oc_origen": oc_origen,
               "sin_sql": sin_sql, "omitidos": [k for k, activo in pasos.items() if not activo],
               "generado": _date.today().isoformat(),
               "generado_hora": datetime.now().astimezone().isoformat(timespec="minutes")}
        with SessionLocal() as ses:
            nombres = {sku: _re.sub(r"^\S+\s+", "", m["desc"]).strip() or m["desc"] for sku, m in med.items()
                       if sku in {x["sku"] for x in posiciones_sap}}
            domain.cargar_lectura_sap(ses, numero, cliente, posiciones_sap, nombres)       # el total del pedido, no lo que falta
            p_ped = domain._get_pedido(ses, numero)
            if p_ped is not None and p_ped.cliente != cliente:
                p_ped.cliente = cliente                  # el análisis manda: el pedido queda con el cliente analizado
            if p_ped is not None:
                domain.completar_oc(p_ped, oc)           # solo si no tiene: no pisa una OC escrita a mano
            # ¿el pedido ya tuvo un suministro con sus entregas y esto es lo que se liberó después?
            previo = ses.get(Config, _clave_cubicaje(numero))
            previo_valor = previo.valor if previo is not None else None
            doc["suministro"], suministro_nuevo = domain.abrir_suministro(ses, numero, sum(f["carga"] for f in res["filas"]))
            domain.guardar_config(ses, _clave_analisis(numero), doc)
            domain.anotar_flujo(ses, numero, analisis=domain.resumen_analisis(doc))
            ses.commit()
            # si el pedido ya estaba cubicado, el cubicaje se rehace con la carga nueva (y el cliente nuevo)
            recubicado = None
            if previo_valor is not None:
                from .cubicaje import _cubicar
                anterior = _json.loads(previo_valor)
                try:
                    _cubicar(numero, {"modo": anterior.get("modo"), "caja_master": anterior.get("caja_master", ""),
                                      "piso_pallet": anterior.get("piso_pallet", "")}, ses)
                    recubicado = True
                except HTTPException as e:
                    recubicado = str(e.detail)
        return {"pedido": numero, "posiciones": len(posiciones), "completos": len(completos), "alertadas": len(res["alertadas"]),
                "limitadas": domain.resumen_analisis(doc)["limitadas"],
                "aviso_sap": " ".join(x for x in [lectura.aviso, *avisos_zsd] if x),
                "sin_cruce": res.get("zsd_sin_cruce", []), "oc": oc, "oc_origen": oc_origen, "sin_sql": sin_sql,
                "omitidos": doc["omitidos"], "recubicado": recubicado, "sesion_sap": sap.ultima_sesion,
                "suministro": doc["suministro"], "suministro_nuevo": suministro_nuevo,
                "export_ruta": publicado.get("ruta", "") if publicado.get("ok") else "",
                "export_aviso": aviso_pub or ((publicado.get("copia") or {}).get("motivo", ""))}

    try:
        return acciones.lanzar_python("analizar_pedido", "Analizar pedido", [numero], correr)
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e


@router.get("/analisis/{numero}")
def get_analisis(numero: str, s: Session = Depends(get_session)):
    import json as _json
    from ..models import Config
    c = s.get(Config, _clave_analisis(numero))
    if c is None:
        raise HTTPException(404, "Este pedido todavía no tiene análisis.")
    return _json.loads(c.valor)


@router.get("/analisis/{numero}/excel")
def excel_analisis(numero: str, s: Session = Depends(get_session)):
    """Excel de estatus para responder a los KAM: qué pidió el cliente, cuánto de eso ya está en entrega, cuánto falta
    y, de lo que falta, si hay stock y cuándo es la próxima liberación.

    Hojas: Estatus (la que se comparte), Entregas (qué lleva cada entrega) y Análisis SAP (el detalle interno)."""
    import json as _json
    import re
    from datetime import datetime as _dt
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..models import Config, Pedido

    ped = s.scalar(select(Pedido).where(Pedido.pedido == numero))
    if ped is None:
        raise HTTPException(404, "El pedido no existe.")
    c = s.get(Config, _clave_analisis(numero))
    doc = _json.loads(c.valor) if c else None
    refrescado = None                                  # ¿el plan SOP, la liberación y el control de calidad son de ahora?
    if doc:
        r = refrescar_analisis(s, numero)              # lo que se mueve (SQL) y las entregas de hoy: el Excel sale al día
        if r:
            doc, refrescado = r[0], r[1]
    oc = (ped.oc or (doc or {}).get("oc_sap") or "").strip()
    activas = [e for e in ped.entregas if not e.anulada]
    filas_an = {f["sku"]: f for f in (doc["resultado"]["filas"] if doc else [])}
    descr = {l.sku: (l.descripcion or (filas_an.get(l.sku) or {}).get("descripcion", "")) for l in ped.lineas}

    # En entrega = lo de las entregas de ESTE pedido (activas) + lo que SAP ya tenía en entregas no registradas aquí
    en_ent: dict[str, int] = {}
    por_sum: dict[int, dict[str, int]] = {}           # suministro -> SKU -> unidades (lo liberado en distintos momentos)
    for e in activas:
        for l in e.lineas:
            en_ent[l.sku] = en_ent.get(l.sku, 0) + l.qty
            d_sum = por_sum.setdefault(e.suministro or 1, {})
            d_sum[l.sku] = d_sum.get(l.sku, 0) + l.qty
    sums = sorted(por_sum) if len(por_sum) > 1 else []
    filas = []
    for l in ped.lineas:
        entrega = en_ent.get(l.sku, 0) + l.externa
        falta = max(0, l.qty - entrega)
        filas.append({"sku": l.sku, "desc": descr.get(l.sku, ""), "pedido": l.qty, "en_entrega": entrega, "falta": falta})
    for sku, q in en_ent.items():                       # entregas con productos que el pedido no trae: no se pierden
        if sku not in {f["sku"] for f in filas}:
            filas.append({"sku": sku, "desc": descr.get(sku, ""), "pedido": 0, "en_entrega": q, "falta": 0})

    # Próxima liberación: se consulta de nuevo AHORA (lo del análisis puede tener horas); si no hay conexión, el del análisis
    pendientes = {f["sku"] for f in filas if f["falta"] > 0}
    disp, disp_fresca = {}, False
    if pendientes:
        try:
            disp, disp_fresca = bases.disponibilidad(fresco=not refrescado), True
        except Exception:  # noqa: BLE001 - el Excel sale igual, con lo último que se sabía
            disp = {}

    # Control de calidad: también se pregunta de nuevo al exportar (si SQL respondió para la liberación)
    calidad, calidad_fresca = {}, False
    if disp_fresca:
        try:
            calidad, calidad_fresca = bases.control_calidad(fresco=not refrescado), True
        except Exception:  # noqa: BLE001
            calidad = {}

    def en_calidad(sku: str):
        if calidad_fresca:
            return calidad.get(sku, 0)
        return (filas_an.get(sku) or {}).get("control_calidad")

    def liberacion(sku: str):
        d = disp.get(sku) if disp_fresca else None
        if d and d.get("cantidad"):
            return d.get("fecha", ""), d["cantidad"]
        cant, _, fecha = str((filas_an.get(sku) or {}).get("disponibilidad") or "").partition(" - ")
        return fecha.strip(), (int(cant) if cant.strip().isdigit() else "")

    def fecha_corta(iso: str) -> str:
        try:
            return _dt.fromisoformat(iso[:10]).strftime("%d-%m-%Y")
        except ValueError:
            return iso or ""

    orden = {"Pendiente": 0, "Parcial": 1, "En entrega": 2}

    def estado(f: dict) -> str:
        return "En entrega" if f["falta"] == 0 else ("Parcial" if f["en_entrega"] > 0 else "Pendiente")
    filas.sort(key=lambda f: (orden[estado(f)], f["sku"]))

    tot_ped, tot_ent = sum(f["pedido"] for f in filas), sum(min(f["en_entrega"], f["pedido"]) for f in filas)
    tot_falta = sum(f["falta"] for f in filas)
    ahora = _dt.now()
    hora_an = (doc or {}).get("generado_hora") or (doc or {}).get("generado") or ""
    try:
        hora_an_txt = _dt.fromisoformat(hora_an).strftime("%d-%m-%Y %H:%M") if "T" in hora_an else fecha_corta(hora_an)
    except ValueError:
        hora_an_txt = hora_an

    wb = Workbook()
    ws = wb.active
    ws.title = "Estatus"
    ws.append([f"Estatus del pedido {numero}" + (f" · OC {oc}" if oc else "") + (f" · {ped.cliente}" if ped.cliente else "")])
    ws.append([f"Pedido: {tot_ped:,} un. · En entrega: {tot_ent:,} un. · Falta: {tot_falta:,} un.".replace(",", ".")
               + (f" · {round(100 * tot_ent / tot_ped)} % cubierto" if tot_ped else "")])
    ws.append([f"Generado el {ahora:%d-%m-%Y %H:%M}. Entregas de la plataforma al momento de exportar"
               + ("; próxima liberación consultada ahora." if disp_fresca else
                  ("." if not pendientes else "; próxima liberación del análisis (no hubo conexión para actualizarla)."))
               + (f" Stock y cantidades de SAP: análisis del {hora_an_txt}." if doc else " Sin análisis de SAP: stock no disponible.")])
    sin_plan = bool(((doc or {}).get("sin_sql") or {}).get("plan"))
    ws.append(["Alerta (del análisis de SAP): Sin stock · Stock parcial · Limitado SOP = el plan SOP deja menos de lo pedido · Completo."
               + (" Este análisis se hizo sin conexión a SQL Server: no se cruzó con el plan SOP." if sin_plan else "")]
              if doc else [])
    CAB = (["SKU", "Descripción", "Pedido", "En entrega", "Falta", "Estado", "Alerta"] + [f"Suministro {k}" for k in sums]
           + ["Stock CD30", "Reserva CD30", "Stock EC01", "Stock TP01", "Control de calidad", "Próx. liberación (fecha)",
              "Próx. liberación (cant.)"])
    ws.append(CAB)
    fila_cab = ws.max_row
    for f in filas:
        st = (filas_an.get(f["sku"]) or {}).get("stock") or {}
        fecha, cant = liberacion(f["sku"]) if f["falta"] > 0 else ("", "")
        fa = filas_an.get(f["sku"]) or {}
        # solo de lo que aún falta; una fila puede tener dos alertas (p. ej. Sin stock y Limitado SOP)
        alerta = " + ".join(fa.get("alertas") or ([fa["alerta"]] if fa.get("alerta") else [])) if f["falta"] > 0 else ""
        ws.append([f["sku"], f["desc"], f["pedido"], f["en_entrega"], f["falta"], estado(f), alerta]
                  + [por_sum[k].get(f["sku"], 0) for k in sums]
                  + ([st.get("cd30"), st.get("reserva_cd30"), st.get("ec01"), st.get("tp01"), en_calidad(f["sku"])]
                     if f["falta"] > 0 else [None] * 5)
                  + [fecha_corta(fecha) if fecha else "", cant])

    we = wb.create_sheet("Entregas")
    we.append(["N° entrega", "SKU", "Descripción", "Qty", "Suministro"])
    for e in sorted(activas, key=lambda x: (x.suministro or 1, x.entrega)):
        for l in e.lineas:
            we.append([e.entrega, l.sku, descr.get(l.sku, ""), l.qty, e.suministro or 1])

    wa = wb.create_sheet("Análisis SAP")
    wa.append(["OC", "SKU", "Descripción", "Alerta", "Qty pendiente", "Qty entrega", "Stock CD30", "Reserva CD30",
               "Stock EC01", "Stock TP01", "Control de calidad", "Próx. liberación (fecha)", "Próx. liberación (cant.)"])
    for f in (doc["resultado"]["filas"] if doc else []):
        st = f.get("stock") or {}
        cant, _, fecha = str(f.get("disponibilidad") or "").partition(" - ")
        wa.append([oc, f["sku"], f["descripcion"], " + ".join(f.get("alertas") or [f.get("alerta", "")]), f["pendiente"],
                   f["qty_entrega"], st.get("cd30"), st.get("reserva_cd30"), st.get("ec01"), st.get("tp01"),
                   f.get("control_calidad"), fecha.strip(), cant.strip()])

    cab, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1A2B4A")
    tonos = {"Pendiente": "FDE2E2", "Parcial": "FFF1CC", "En entrega": "E3F4E8"}
    tonos_alerta = {"Sin stock": "FADDD9", "Stock parcial": "FBEBCD", "Limitado SOP": "DDE8F3", "Completo": "DDF0E5"}
    tinta_alerta = {"Sin stock": "B3261E", "Stock parcial": "8F5200", "Limitado SOP": "3F6A99", "Completo": "17724A"}
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(bold=True)
    ws["A3"].font = Font(italic=True, color="555555")
    if doc:
        ws["A4"].font = Font(italic=True, color="555555")
    for hoja in wb.worksheets:
        fc = fila_cab if hoja is ws else 1
        for celda in hoja[fc]:
            celda.font, celda.fill = cab, fondo
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in hoja.columns:
            ancho = max(len(str(x.value if x.value is not None else "")) for x in col[(fc - 1 if hoja is ws else 0):]) + 2
            hoja.column_dimensions[col[0].column_letter].width = min(max(ancho, 10), 46)
        hoja.freeze_panes = f"A{fc + 1}"
        hoja.auto_filter.ref = f"A{fc}:{hoja.cell(row=hoja.max_row, column=hoja.max_column).coordinate}"
    for r in range(fila_cab + 1, ws.max_row + 1):                   # el estado se lee de un vistazo
        ws.cell(row=r, column=6).fill = PatternFill("solid", fgColor=tonos[ws.cell(row=r, column=6).value])
        al = ws.cell(row=r, column=7)                                   # la alerta, con los mismos tonos de la plataforma
        primera = str(al.value or "").split(" + ")[0]                  # con dos alertas manda el color de la primera
        if primera in tonos_alerta:
            al.fill = PatternFill("solid", fgColor=tonos_alerta[primera])
            al.font = Font(bold=True, color=tinta_alerta[primera])
    ws.column_dimensions["A"].width = 16
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nombre = re.sub(r"[^\w.-]", "_", f"Estatus_{numero}" + (f"_OC_{oc}" if oc else ""))
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})


@router.put("/analisis/{numero}/carga")
def ajustar_carga(numero: str, body: dict, s: Session = Depends(get_session)):
    """Ajuste manual de la CARGA (como editar la columna F del Excel).

    Un SKU:    {"sku": "...", "carga": 40}      (carga vacía = volver al valor calculado)
    Varios:    {"ajustes": {"sku": 40, "otro": null}}
    Subir la carga por sobre lo que deja el plan SOP es una autorización: queda anotado quién la dio."""
    import json as _json
    from datetime import datetime as _dt
    from ..analisis import calcular
    from ..models import Config
    c = s.get(Config, _clave_analisis(numero))
    if c is None:
        raise HTTPException(404, "Este pedido todavía no tiene análisis.")
    doc = _json.loads(c.valor)
    cambios = body.get("ajustes")
    if cambios is None:
        cambios = {body.get("sku", ""): body.get("carga")}
    if not isinstance(cambios, dict) or not cambios:
        raise HTTPException(422, "No hay ajustes que aplicar.")
    en_analisis = {p["sku"] for p in doc["posiciones"]}
    calculada = {f["sku"]: f["carga_calculada"] for f in (doc.get("resultado") or {}).get("filas", [])}
    doc.setdefault("autorizaciones", {})
    for sku_raw, valor in cambios.items():
        sku = domain.norm_sku(sku_raw)
        if sku not in en_analisis:
            raise HTTPException(422, f"El SKU {sku or sku_raw} no está en el análisis.")
        if valor in (None, ""):
            doc["ajustes"].pop(sku, None)                    # volver al valor calculado
            doc["autorizaciones"].pop(sku, None)
            continue
        try:
            carga = float(valor)
        except (TypeError, ValueError) as e:
            raise HTTPException(422, "Cantidad inválida.") from e
        if not math.isfinite(carga) or carga < 0:
            raise HTTPException(422, "La cantidad no puede ser negativa ni infinita.")
        doc["ajustes"][sku] = carga
        if carga > calculada.get(sku, 0):
            doc["autorizaciones"][sku] = {"por": domain.usuario_actual(), "at": _dt.now().isoformat(timespec="seconds"),
                                          "calculada": calculada.get(sku, 0), "carga": carga}
        else:
            doc["autorizaciones"].pop(sku, None)
    # Recalcula con los mismos datos guardados (no vuelve a consultar SAP)
    doc["resultado"] = calcular(doc["posiciones"], doc["plan"], doc["en_entrega"], doc["medidas"],
                                doc["disponibilidad"], doc["stock"], doc["ajustes"],
                                sin_plan=bool((doc.get("sin_sql") or {}).get("plan")),
                                control_calidad=doc.get("control_calidad"))
    domain.guardar_config(s, _clave_analisis(numero), doc)
    domain.anotar_flujo(s, numero, analisis=domain.resumen_analisis(doc))
    _commit(s)
    # Cubicaje en vivo: si el pedido ya estaba cubicado, se rehace con la carga nueva
    cubicaje = None
    from ..models import Config as _Cfg
    prev = s.get(_Cfg, _clave_cubicaje(numero))
    if prev is not None:
        from .cubicaje import _cubicar
        anterior = _json.loads(prev.valor)
        try:
            cubicaje = _cubicar(numero, {"modo": anterior.get("modo"),
                                         "caja_master": anterior.get("caja_master", ""),
                                         "piso_pallet": anterior.get("piso_pallet", "")}, s)
        except HTTPException as e:
            cubicaje = {"error": e.detail}
    return {**doc, "cubicaje": cubicaje}


@router.post("/sap/crear_entregas", status_code=202)
def sap_crear_entregas(body: dict, s: Session = Depends(get_session)):
    """Crea en SAP una entrega por camión (y por pedido), con las cantidades del cubicaje.
    Con `pedidos` (varios, cubicados juntos) se crea una entrega por cada par (camión, pedido): los camiones se comparten
    y cada entrega queda marcada con su camión para poder agruparlas después."""
    import json as _json
    from ..integrations import sap_crear
    from ..models import Config
    numeros = [str(x).strip() for x in (body.get("pedidos") or []) if str(x).strip()] or [str(body.get("pedido", "")).strip()]
    ensayo = bool(body.get("ensayo", True))
    puesto_body = str(body.get("puesto", "")).strip().upper()
    camiones_pedidos = body.get("camiones") or []
    tareas = []                                   # (pedido, puesto, {camión: {sku: unidades}}, id del conjunto)
    for numero in dict.fromkeys(numeros):
        c = s.get(Config, _clave_cubicaje(numero))
        if c is None:
            raise HTTPException(422, f"Primero hay que cubicar el pedido {numero}.")
        cub = _json.loads(c.valor)
        puesto = puesto_body
        if not puesto:
            an = s.get(Config, _clave_analisis(numero))
            puesto = (_json.loads(an.valor).get("puesto", "") if an else "").upper()
        if not puesto:
            raise HTTPException(422, "Falta el puesto de expedición (PN01, PN02…).")
        # Un camión = una entrega: {camión: {sku: unidades}}
        por_camion: dict[int, dict] = {}
        for f in cub["filas"]:
            if camiones_pedidos and f["camion"] not in camiones_pedidos:
                continue
            por_camion.setdefault(f["camion"], {})
            por_camion[f["camion"]][f["sku"]] = por_camion[f["camion"]].get(f["sku"], 0) + f["unidades"]
        if not por_camion:
            raise HTTPException(422, f"El cubicaje del pedido {numero} no tiene camiones para crear.")
        tareas.append((numero, puesto, por_camion, ((cub.get("conjunto") or {}).get("id") or ""),
                       _fecha_vl01n(s, numero, body.get("fecha"))))
    numero = tareas[0][0]
    total = sum(len(x[2]) for x in tareas)

    def correr(avance):
        salidas = []
        hecho = 0
        for numero, puesto, por_camion, conj, fecha in tareas:
            for cam, materiales in sorted(por_camion.items()):
                hecho += 1
                avance(f"{'Ensayo' if ensayo else 'Creando'} entrega {hecho} de {total}: pedido {numero}, camión {cam} (fecha de selección {fecha})")
                r = sap_crear.crear_entrega(numero, puesto, fecha, materiales,
                                            fecha_cita=str(body.get("fecha_cita") or ""),
                                            hora_cita=str(body.get("hora_cita") or ""), ensayo=ensayo)
                salidas.append({"pedido": numero, "camion": cam, "ok": r.ok, "entrega": r.entrega, "mensaje": r.mensaje,
                                "pasos": r.pasos, "borradas": r.borradas, "ajustadas": r.ajustadas,
                                "incidencias": r.incidencias})
                if r.ok and r.entrega and not ensayo:
                    # La entrega ya existe en SAP y no se puede deshacer: si no se alcanza a registrar aquí,
                    # se avisa con su número para cargarla a mano y se sigue con los demás camiones.
                    try:
                        from ..db import SessionLocal
                        with SessionLocal() as ses:
                            from ..schemas import EntregaIn
                            domain.guardar_entrega(ses, EntregaIn(
                                entrega=r.entrega, pedido=numero, grupo="",
                                camion_ref=f"{conj}-{cam}" if conj else "",
                                lineas=[{"sku": k, "qty": v} for k, v in materiales.items()]))
                            ses.commit()
                    except Exception as err:  # noqa: BLE001
                        salidas[-1]["incidencias"] = list(r.incidencias or []) + [
                            f"Creada en SAP, pero no se pudo registrar en la plataforma: {str(err)[:150]}"]
                if not r.ok:
                    break                      # ante el primer problema se detiene
            if salidas and not salidas[-1]["ok"]:
                break
        errores = [x for x in salidas if not x["ok"]]
        if errores:
            raise RuntimeError(errores[0]["mensaje"])
        return {"pedido": numero, "pedidos": [x[0] for x in tareas], "ensayo": ensayo, "resultados": salidas}

    try:
        return acciones.lanzar_python("crear_entregas",
                                      f"{'Ensayo de entregas' if ensayo else 'Crear entregas'} {numero}" + (f" y otros {len(tareas) - 1}" if len(tareas) > 1 else ""),
                                      [numero], correr)
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e


def _fecha_vl01n(s: Session, numero: str, fecha_body) -> str:
    """Fecha de selección de VL01N (dd.mm.aaaa) para crear las entregas de un pedido.

    Manda la que se indique al crear. Si no, la fecha con la que se analizó el pedido: hay clientes (PARIS) cuyas
    órdenes de compra rigen desde la semana siguiente, y con la fecha de hoy VL01N no encuentra las líneas. Nunca
    anterior a hoy, porque una fecha de selección más tardía solo suma líneas."""
    import json as _json
    from datetime import date as _date
    from ..models import Config
    hoy = _date.today()
    if str(fecha_body or "").strip():
        txt = str(fecha_body).strip()
        try:                                    # acepta dd.mm.aaaa o aaaa-mm-dd
            d = _date.fromisoformat(txt) if "-" in txt else _date(int(txt[6:10]), int(txt[3:5]), int(txt[0:2]))
        except (ValueError, IndexError):
            raise HTTPException(422, "La fecha de selección debe ser dd.mm.aaaa.") from None
        return d.strftime("%d.%m.%Y")
    an = s.get(Config, _clave_analisis(numero))
    try:
        d = _date.fromisoformat(_json.loads(an.valor).get("fecha", "")) if an else hoy
    except (ValueError, TypeError):
        d = hoy
    return max(d, hoy).strftime("%d.%m.%Y")


@router.post("/sap/crear_grupo", status_code=202)
def sap_crear_grupo(body: dict, s: Session = Depends(get_session)):
    """Crea el grupo de transporte de un camión (VL06), igual que el paso del Excel."""
    from ..integrations import sap_crear
    entregas = [str(x).strip() for x in (body.get("entregas") or []) if str(x).strip()]
    ensayo = bool(body.get("ensayo", True))
    camion = str(body.get("camion") or "1")
    cliente = str(body.get("cliente") or "")
    referencia = str(body.get("referencia") or "")
    if not entregas:
        raise HTTPException(422, "No hay entregas para agrupar.")

    def correr(avance):
        avance(f"{'Ensayo' if ensayo else 'Creando grupo'} del camión {camion} "
               f"({len(entregas)} entregas)")
        r = sap_crear.crear_grupo(entregas, referencia=referencia, ensayo=ensayo,
                                  cliente=cliente, camion=camion)
        if not r.ok:
            raise RuntimeError(r.mensaje)
        if r.grupo:
            from ..db import SessionLocal
            with SessionLocal() as ses:
                for n in entregas:
                    e = ses.scalar(select(Entrega).where(Entrega.entrega == n))
                    if e is not None:
                        e.grupo = r.grupo
                        domain._log(e, f"Agrupada en {r.grupo} desde la plataforma", "web")
                        e.actualizado = domain.ahora()
                ses.commit()
        return {"grupo": r.grupo, "ensayo": ensayo, "mensaje": r.mensaje, "pasos": r.pasos,
                "incidencias": r.incidencias}

    try:
        return acciones.lanzar_python("crear_grupo", f"Grupo del camión {camion}", entregas, correr)
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e


@router.post("/sap/fecha_grupo", status_code=202)
def sap_fecha_grupo(body: dict, s: Session = Depends(get_session)):
    """Fecha, hora y referencia de la cita de un grupo (VG02), como el paso del Excel."""
    from datetime import date as _date
    from ..integrations import sap_crear
    grupo = str(body.get("grupo", "")).strip()
    fecha = str(body.get("fecha", "")).strip()
    hora = str(body.get("hora", "")).strip()
    referencia = str(body.get("referencia", "")).strip()
    ensayo = bool(body.get("ensayo", True))
    if not grupo:
        raise HTTPException(422, "Falta el número de grupo.")
    if not (fecha or hora or referencia):
        raise HTTPException(422, "Indica al menos fecha, hora o referencia.")

    def correr(avance):
        avance(f"{'Ensayo' if ensayo else 'Actualizando'} la cita del grupo {grupo}")
        r = sap_crear.actualizar_grupo(grupo, fecha, hora, referencia, ensayo=ensayo)
        if not r.ok:
            raise RuntimeError(r.mensaje)
        if not ensayo:
            from ..db import SessionLocal
            with SessionLocal() as ses:
                for e in ses.scalars(select(Entrega).where(Entrega.grupo == grupo)).all():
                    if fecha:
                        try:
                            d, m, a = fecha.split(".")
                            e.cita_fecha = _date(int(a), int(m), int(d))
                        except ValueError:
                            pass
                    if hora:
                        e.cita_hora = hora[:5]
                    if referencia:
                        e.cita_numero = referencia
                    domain._log(e, f"Cita en SAP: {fecha} {hora} {referencia}".strip(), "web")
                    e.actualizado = domain.ahora()
                ses.commit()
        return {"grupo": grupo, "ensayo": ensayo, "mensaje": r.mensaje, "pasos": r.pasos,
                "incidencias": r.incidencias}

    try:
        return acciones.lanzar_python("fecha_grupo", f"Cita del grupo {grupo}", [grupo], correr)
    except RuntimeError as e:
        raise HTTPException(409, str(e)) from e


@router.post("/sap/borrar_entrega", status_code=202)
def sap_borrar_entrega(body: dict, s: Session = Depends(get_session)):
    """Borra la entrega en SAP (VL06) y recién entonces la quita de la plataforma."""
    import re as _re
    from ..integrations import sap
    entrega = str(body.get("entrega", "")).strip()
    if not _re.fullmatch(r"\d{4,12}", entrega):
        raise HTTPException(422, "N° de entrega inválido.")
    if str(body.get("confirmar", "")).strip() != entrega:
        raise HTTPException(422, "Falta confirmar escribiendo el número de la entrega.")
    e = s.scalar(select(Entrega).where(Entrega.entrega == entrega))
    if e is None:
        raise HTTPException(404, "Esa entrega no está en la plataforma.")
    if domain.paso_ok(e, "facturada"):
        raise HTTPException(422, "La entrega está facturada: no se puede borrar.")

    def correr(avance):
        avance("Borrando en SAP (VL06)")
        ok, mensaje = sap.borrar_entrega(entrega)
        from ..db import SessionLocal
        with SessionLocal() as ses:
            if ok:
                domain.borrar_entrega(ses, entrega)
                ses.commit()
        if not ok:
            raise RuntimeError(mensaje)
        return {"entrega": entrega, "mensaje": mensaje, "borrada": True}

    try:
        return acciones.lanzar_python("borrar_entrega", f"Borrar entrega {entrega} en SAP",
                                      [entrega], correr)
    except RuntimeError as err:
        raise HTTPException(409, str(err)) from err


@router.post("/sap/borrar_grupo", status_code=202)
def sap_borrar_grupo(body: dict, s: Session = Depends(get_session)):
    """Borra el grupo en SAP (VG02). Las entregas quedan sin grupo, no se borran."""
    import re as _re
    from ..integrations import sap
    grupo = str(body.get("grupo", "")).strip()
    if not _re.fullmatch(r"\d{3,12}", grupo):
        raise HTTPException(422, "N° de grupo inválido.")
    if str(body.get("confirmar", "")).strip() != grupo:
        raise HTTPException(422, "Falta confirmar escribiendo el número del grupo.")

    def correr(avance):
        avance("Borrando el grupo en SAP (VG02)")
        ok, mensaje = sap.borrar_grupo(grupo)
        from ..db import SessionLocal
        with SessionLocal() as ses:
            if ok:
                for e in ses.scalars(select(Entrega).where(Entrega.grupo == grupo)).all():
                    e.grupo = ""
                    domain._log(e, f"Grupo {grupo} borrado en SAP", "web")
                    e.actualizado = domain.ahora()
                ses.commit()
        if not ok:
            raise RuntimeError(mensaje)
        return {"grupo": grupo, "mensaje": mensaje, "borrado": True}

    try:
        return acciones.lanzar_python("borrar_grupo", f"Borrar grupo {grupo} en SAP",
                                      [grupo], correr)
    except RuntimeError as err:
        raise HTTPException(409, str(err)) from err


@router.get("/acciones/trabajos")
def lista_trabajos():
    """Los últimos trabajos (el que está en curso primero). Si se recarga la página mientras SAP trabaja,
    la plataforma retoma el seguimiento en vez de dar el trabajo por perdido."""
    return acciones.recientes()


@router.get("/acciones/trabajos/{tid}")
def get_trabajo(tid: str):
    t = acciones.trabajo(tid)
    if t is None:
        raise HTTPException(404, "Trabajo no encontrado.")
    return t
