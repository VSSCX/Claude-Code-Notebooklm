"""Base de Medidas (xlsm en la red): descripción y "Máx Camión" por SKU.

Es la misma fuente que usa el Excel (consulta "Base de Medidas", hoja
"Base para carga"). La descripción de aquí es la que se cruza con el reporte
ZSD001_03 para calcular la Qty en entrega, igual que la fórmula de la columna O.
"""
import logging
import re
import threading
import time
from pathlib import Path

from ..config import settings

log = logging.getLogger(__name__)
HOJA = "Base para carga"
_estado = {"mtime": None, "revisado": 0.0, "datos": {}, "error": "",
           "filas": None, "revisado_filas": 0.0}
_lock = threading.Lock()


def _sku(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"^0+(?=\d)", "", str(v).strip())


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _leer_filas(ruta: Path) -> list[list]:
    """Filas crudas (columnas A..L) desde la fila de encabezados, para el cubicaje."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]
        filas, empezo = [], False
        for f in ws.iter_rows(values_only=True):
            valores = [c for c in f[:12]]
            textos = [str(c).strip() if c is not None else "" for c in valores]
            if not empezo:
                empezo = "Grupo" in textos and "Máx Camión" in textos
                continue
            filas.append(valores)
        return filas
    finally:
        wb.close()


def filas() -> list[list]:
    ruta = Path(settings.base_medidas) if settings.base_medidas else None
    if ruta is None:
        return []
    if not _lock.acquire(blocking=False):                  # otra petición ya la lee: se usa lo último
        return _estado.get("filas") or []
    try:
        if _estado.get("filas") is not None and time.time() - _estado["revisado_filas"] < 60:
            return _estado["filas"]
        _estado["revisado_filas"] = time.time()
        try:
            _estado["filas"] = _leer_filas(ruta)
        except Exception as e:  # noqa: BLE001
            _estado["error"] = str(e)[:200]
            log.warning("No se pudieron leer las filas de la Base de Medidas: %s", e)
            _estado["filas"] = _estado.get("filas") or []
        return _estado["filas"]
    finally:
        _lock.release()


def _leer(ruta: Path) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb[HOJA] if HOJA in wb.sheetnames else wb.worksheets[0]
        filas = ws.iter_rows(values_only=True)
        cab = None
        for f in filas:                       # la cabecera real está en la 2a fila
            valores = [str(c).strip() if c is not None else "" for c in f]
            if "Grupo" in valores and "Máx Camión" in valores:
                cab = valores
                break
        if cab is None:
            raise ValueError("No se encontró la fila de encabezados (Grupo / Máx Camión).")
        i_g, i_d, i_m = cab.index("Grupo"), cab.index("Descripción"), cab.index("Máx Camión")
        datos = {}
        for f in filas:
            if len(f) <= max(i_g, i_d, i_m):
                continue
            sku = _sku(f[i_g])
            if sku and sku not in datos:
                datos[sku] = {"desc": str(f[i_d] or "").strip(), "max_camion": _num(f[i_m])}
        return datos
    finally:
        wb.close()


def medidas() -> dict:
    ruta = Path(settings.base_medidas) if settings.base_medidas else None
    if ruta is None:
        return {}
    if not _lock.acquire(blocking=False):
        return _estado["datos"]
    try:
        if time.time() - _estado["revisado"] < 60:          # también tras un fallo: no se reintenta a cada petición
            return _estado["datos"]
        _estado["revisado"] = time.time()
        try:
            mtime = ruta.stat().st_mtime
            if mtime != _estado["mtime"]:
                _estado["datos"], _estado["mtime"], _estado["error"] = _leer(ruta), mtime, ""
        except Exception as e:  # noqa: BLE001
            _estado["error"] = str(e)[:200]
            log.warning("No se pudo leer la Base de Medidas: %s", e)
        return _estado["datos"]
    finally:
        _lock.release()


def estado() -> dict:
    return {"ruta": settings.base_medidas, "productos": len(_estado["datos"]), "error": _estado["error"]}
