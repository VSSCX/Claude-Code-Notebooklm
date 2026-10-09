"""Kits: un SKU propio que reúne cajas separadas (horno + encimera + campana...) que viajan juntas.

El listado se carga desde un Excel (una fila por componente) y queda guardado en la plataforma, igual que la
Base de Medidas. Cada importación agrega los kits nuevos y reemplaza la composición de los que vienen en el archivo;
el resto no se toca.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..cubicaje.kits import KitDef
from ..models import Kit, KitComponente, ahora

CABECERA = ("Kit código SAP", "Descripción del kit", "Componente código SAP", "Cantidad por kit")


def norm_sku(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"^0+(?=\d)", "", str(v).strip())


@dataclass
class Resumen:
    nuevos: int = 0
    actualizados: int = 0
    sin_cambios: int = 0
    ignoradas: int = 0
    problemas: list = field(default_factory=list)
    total_filas: int = 0


def leer_archivo(origen) -> list[dict]:
    """Lee la primera hoja con las 4 columnas (acepta .xlsx; los encabezados se reconocen aunque cambie la redacción)."""
    from openpyxl import load_workbook
    wb = load_workbook(origen, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            cols, filas = None, []
            for f in ws.iter_rows(values_only=True):
                textos = [str(c).strip().lower() if c is not None else "" for c in f]
                if cols is None:
                    cols = _columnas(textos)
                    continue
                if not any(c is not None and str(c).strip() for c in f):
                    continue
                filas.append({"kit": f[cols[0]] if cols[0] < len(f) else None,
                              "descripcion": f[cols[1]] if cols[1] is not None and cols[1] < len(f) else "",
                              "componente": f[cols[2]] if cols[2] < len(f) else None,
                              "cantidad": f[cols[3]] if cols[3] is not None and cols[3] < len(f) else 1})
            if cols is not None and filas:
                return filas
        return []
    finally:
        wb.close()


def _columnas(textos: list[str]):
    """Posición de (kit, descripción, componente, cantidad) en una fila de encabezados; None si no es el encabezado."""
    def buscar(*claves, excluir=()):
        for i, t in enumerate(textos):
            if any(c in t for c in claves) and not any(x in t for x in excluir):
                return i
        return None
    kit = buscar("kit", excluir=("componente", "descrip"))
    comp = buscar("componente")
    if kit is None or comp is None:
        return None
    return kit, buscar("descrip"), comp, buscar("cantidad", "cant")


def _entero(v) -> int:
    try:
        return int(round(float(str(v).replace(",", "."))))
    except (TypeError, ValueError):
        return 0


def importar(s: Session, filas: list[dict]) -> Resumen:
    r = Resumen(total_filas=len(filas))
    grupos: dict[str, dict] = {}
    for n, f in enumerate(filas, start=2):
        kit, comp = norm_sku(f.get("kit")), norm_sku(f.get("componente"))
        cant = _entero(f.get("cantidad") if f.get("cantidad") not in (None, "") else 1)
        if not kit or not comp:
            r.ignoradas += 1
            r.problemas.append(f"Fila {n}: falta el código del kit o del componente.")
            continue
        if cant <= 0:
            r.ignoradas += 1
            r.problemas.append(f"Fila {n}: la cantidad por kit de {comp} debe ser mayor que cero.")
            continue
        if comp == kit:
            r.ignoradas += 1
            r.problemas.append(f"Fila {n}: el kit {kit} no puede ser componente de sí mismo.")
            continue
        g = grupos.setdefault(kit, {"desc": "", "comp": {}})
        g["desc"] = g["desc"] or str(f.get("descripcion") or "").strip()
        g["comp"][comp] = g["comp"].get(comp, 0) + cant           # una fila repetida suma
    actuales = {k.sku: k for k in s.scalars(select(Kit)).all()}
    for sku, g in grupos.items():
        anidados = [c for c in g["comp"] if c in grupos]
        if anidados:
            r.ignoradas += len(g["comp"])
            r.problemas.append(f"Kit {sku}: lleva dentro el kit {anidados[0]}; los kits dentro de kits no se admiten, "
                               "descompón sus componentes.")
            continue
        k = actuales.get(sku)
        nuevos = sorted(g["comp"].items())
        if k is None:
            s.add(Kit(sku=sku, descripcion=g["desc"][:200],
                      componentes=[KitComponente(sku=c, cantidad=q) for c, q in nuevos]))
            r.nuevos += 1
            continue
        antes = sorted((c.sku, c.cantidad) for c in k.componentes)
        if antes == nuevos and (not g["desc"] or g["desc"] == k.descripcion):
            r.sin_cambios += 1
            continue
        if g["desc"]:
            k.descripcion = g["desc"][:200]
        if antes != nuevos:
            k.componentes.clear()
            s.flush()
            k.componentes.extend(KitComponente(sku=c, cantidad=q) for c, q in nuevos)
        k.actualizado = ahora()
        r.actualizados += 1
    s.flush()
    return r


def definiciones(s: Session) -> dict[str, KitDef]:
    """Kits listos para el motor, por SKU."""
    return {k.sku: KitDef(k.sku, k.descripcion, [(c.sku, c.cantidad) for c in k.componentes])
            for k in s.scalars(select(Kit)).all() if k.componentes}


def listar(s: Session, buscar: str = "", limite: int = 100) -> dict:
    q = select(Kit).order_by(Kit.sku)
    t = (buscar or "").strip()
    if t:
        q = q.where(Kit.sku.like(f"%{t}%") | Kit.descripcion.like(f"%{t}%"))
    kits = s.scalars(q.limit(max(1, min(limite, 500)))).all()
    total = len(s.scalars(select(Kit.sku)).all())
    return {"total": total, "filas": [
        {"sku": k.sku, "descripcion": k.descripcion,
         "componentes": [{"sku": c.sku, "cantidad": c.cantidad} for c in k.componentes],
         "actualizado": k.actualizado.isoformat() if k.actualizado else ""} for k in kits]}


def borrar(s: Session, sku: str) -> bool:
    k = s.get(Kit, norm_sku(sku))
    if k is None:
        return False
    s.delete(k)
    s.flush()
    return True


def plantilla() -> bytes:
    """Excel para cargar kits: una fila por componente."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Kits"
    ws.append(list(CABECERA))
    ejemplo = [("KIT-0001", "Kit cocina Horno + Encimera + Campana", "900100001", 1),
               ("KIT-0001", "", "900100002", 1), ("KIT-0001", "", "900100003", 1),
               ("KIT-0002", "Kit Encimera + Campana", "900100002", 1), ("KIT-0002", "", "900100003", 1)]
    for f in ejemplo:
        ws.append(list(f))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center")
    for col, ancho in zip("ABCD", (20, 44, 24, 18)):
        ws.column_dimensions[col].width = ancho
    ws["A1"].comment = Comment("El SKU del kit, tal como viene en el pedido. Repítelo en cada fila de sus componentes.", "Order Desk")
    ws["D1"].comment = Comment("Cuántas unidades de este componente lleva UN kit.", "Order Desk")
    ws.freeze_panes = "A2"
    ayuda = wb.create_sheet("Cómo se usa")
    for linea in (
        "Una fila por cada componente de cada kit.",
        "Repite el código del kit en todas sus filas (la descripción se puede escribir solo en la primera).",
        "Los componentes deben existir en la Base de Medidas: el kit se cubica con las medidas de cada caja, no con una medida única.",
        "Cargar de nuevo un kit reemplaza su composición; los kits que no vienen en el archivo no se tocan.",
        "Los ejemplos de esta plantilla son solo de muestra: bórralos antes de cargar tu listado."):
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
